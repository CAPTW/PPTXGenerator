"""Text object creation for E06.2 contract compilation."""

from __future__ import annotations

from typing import Any

from pptx.dml.color import RGBColor
from pptx.util import Pt

from src.presentation_agent.magic_layer.e06_2_contract_object_factory import contract_shape_name


def add_contract_text(slide: Any, obj: dict[str, Any]) -> Any:
    bbox = obj["bbox_emu"]
    shape = slide.shapes.add_textbox(int(bbox["x"]), int(bbox["y"]), max(1, int(bbox["w"])), max(1, int(bbox["h"])))
    shape.name = contract_shape_name(obj)
    text = str(obj.get("text_excerpt") or "")
    shape.text = text
    frame = shape.text_frame
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    for paragraph in frame.paragraphs:
        for run in paragraph.runs:
            run.font.name = "Aptos"
            run.font.size = Pt(_font_size_for(obj))
            run.font.color.rgb = RGBColor(226, 232, 240)
    return shape


def _font_size_for(obj: dict[str, Any]) -> float:
    h_in = float(obj.get("bbox_in", {}).get("h", 0.2))
    role = str(obj.get("semantic_role") or "")
    if role == "title" or "title" in str(obj.get("name", "")).lower():
        return min(26, max(12, h_in * 28))
    if obj.get("object_type") == "source_footer":
        return 7
    return min(14, max(6, h_in * 18))
