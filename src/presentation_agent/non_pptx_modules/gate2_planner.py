"""Deterministic Gate 2 blueprint and visual-system planning."""

from __future__ import annotations

import json
import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import yaml
from pydantic import Field, field_validator

from ..compat.legacy_non_pptx import WorkflowGate
from .authoring_compiler import compile_authoring_layer
from .generation_mode_router import build_canonical_generation_profile
from .gate2_context import load_gate2_context
from .lecture_synthesis import plan_concept_driven_lecture
from .shared_proof_registry import SharedProofCompatMode, shared_proof_consumer_policy
from ..compat.legacy_non_pptx import (
    AssetKind,
    AssetPriority,
    AssetRequest,
    AssetRequests,
    AuthoringPreview,
    Blueprint,
    BlueprintPreview,
    BlueprintSlide,
    BriefMaterialType,
    CanonicalGenerationProfile,
    ColorToken,
    CommunicationCore,
    ConceptEdge,
    ConceptEdgeType,
    ConceptGraph,
    ConceptNode,
    ConceptType,
    ContractModel,
    CountRange,
    ContentTier,
    CropReviewAction,
    DEFAULT_STATE_FILENAMES,
    DeckConstitution,
    DeckMode,
    DesignSystem,
    EvidencePlanItem,
    InfographicPlanItem,
    LayoutLibrary,
    LayoutPattern,
    LectureFamily,
    PresentationBrief,
    PresentationType,
    ProductionBridge,
    ProductionMode,
    ProjectMaterial,
    ReferenceDNA,
    ScaleMode,
    SlideIntent,
    SlideFunctionOutline,
    SlideEvidenceClass,
    SlideDensityBudget,
    SlideLedger,
    SlideLedgerEntry,
    ProofCoverageClass,
    ProofEvidenceOrigin,
    ProofUnit,
    ProofUnitRegistry,
    ProofModule,
    ProofModuleManifest,
    ProofModuleStatus,
    SlideRole,
    StoryArchitectureSummary,
    SourceMaterialRef,
    StageStatus,
    StorySection,
    TeachingPlan,
    TeachingPlanSlide,
    TypographyToken,
    VerificationPoint,
    VisualReferenceSummary,
    VisualRoute,
    VisualSourcePreference,
    VisualType,
    WorkflowDeltaDetail,
    WorkflowOption,
    WorkflowPhase,
    WorkflowPlan,
    proof_module_manifest_from_proof_unit_registry,
    save_state_file,
)
from .workflow_planner import WorkflowBriefInput, load_workflow_brief


HEADING_RE = re.compile(r"^(?P<number>\d+(?:\.\d+)+)\s+(?P<title>.+)$")
LECTURE_MAX_BULLETS = 3
LECTURE_MAX_BULLET_CHARS = 92
LECTURE_MAIN_TARGET = 60
LECTURE_APPENDIX_TARGET = 18


@dataclass(slots=True)
class LectureHeadingBlock:
    number: str
    title: str
    level: int
    body: list[str] = field(default_factory=list)
    page_hint: int | None = None

    @property
    def parent_key(self) -> str | None:
        parts = self.number.split(".")
        if len(parts) >= 2:
            return ".".join(parts[:2])
        return None

    @property
    def chapter_key(self) -> str:
        return self.number.split(".", 1)[0]


class BrandInputs(ContractModel):
    brand_name: str | None = None
    primary_color: str | None = None
    accent_color: str | None = None
    font_family: str | None = None
    tone_keywords: list[str] = Field(default_factory=list)
    chart_preferences: list[str] = Field(default_factory=list)
    icon_style: str | None = None
    section_divider_style: str | None = None
    prohibited_elements: list[str] = Field(default_factory=list)

    @field_validator("tone_keywords", "chart_preferences", "prohibited_elements", mode="before")
    @classmethod
    def _coerce_list(cls, value: object) -> object:
        if value is None:
            return []
        if isinstance(value, str):
            return [value]
        return value


class Gate2Outputs(ContractModel):
    blueprint: Blueprint
    proof_unit_registry: ProofUnitRegistry | None = None
    proof_module_manifest: ProofModuleManifest | None = None
    presentation_brief: PresentationBrief | None = None
    canonical_generation_profile: CanonicalGenerationProfile | None = None
    slide_function_outline: SlideFunctionOutline | None = None
    concept_graph: ConceptGraph
    teaching_plan: TeachingPlan
    blueprint_preview: BlueprintPreview
    authoring_preview: AuthoringPreview
    design_system: DesignSystem
    deck_constitution: DeckConstitution
    layout_library: LayoutLibrary
    slide_ledger: SlideLedger
    asset_requests: AssetRequests


def load_brand_inputs(path: str | Path) -> BrandInputs:
    brand_path = Path(path)
    text = brand_path.read_text(encoding="utf-8")
    if brand_path.suffix.lower() in {".yaml", ".yml"}:
        payload = yaml.safe_load(text)
    else:
        payload = json.loads(text)
    if payload is None:
        payload = {}
    if not isinstance(payload, dict):
        raise ValueError(f"brand inputs must contain a top-level object: {brand_path}")
    return BrandInputs.model_validate(payload)


def _slugify(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-") or "item"


def _dedupe(items: list[str]) -> list[str]:
    seen: set[str] = set()
    ordered: list[str] = []
    for item in items:
        cleaned = item.strip()
        if cleaned and cleaned not in seen:
            seen.add(cleaned)
            ordered.append(cleaned)
    return ordered


def _supports_continuity_controls(scale_mode: ScaleMode) -> bool:
    return scale_mode in {ScaleMode.EXTENDED, ScaleMode.LARGE_DECK, ScaleMode.MEGA_DECK}


def _is_lecture_mode(workflow_plan: WorkflowPlan, brief: WorkflowBriefInput | None = None) -> bool:
    if workflow_plan.workflow_option == "graduate-lecture-clustered":
        return True
    if brief is None:
        return False
    audience_text = " ".join(brief.audience).lower()
    note_text = " ".join([brief.topic, brief.purpose, *brief.constraints, *brief.notes]).lower()
    duration = brief.expected_duration_minutes or 0
    has_documents = any(material.material_type == BriefMaterialType.DOCUMENT for material in brief.current_materials)
    lecture_signal = any(
        token in audience_text or token in note_text
        for token in (
            "lecture",
            "graduate",
            "students",
            "\uac15\uc758",
            "\ub300\ud559\uc6d0",
            "\ub300\ud559\uc6d0\uc0dd",
            "\uad50\uc7ac",
        )
    )
    return has_documents and duration >= 90 and lecture_signal


def _contains_hangul(text: str) -> bool:
    return any("\uac00" <= character <= "\ud7a3" for character in text)


def _normalize_text(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def _shorten_line(value: str, limit: int = LECTURE_MAX_BULLET_CHARS) -> str:
    cleaned = _normalize_text(value)
    if len(cleaned) <= limit:
        return cleaned
    trimmed = cleaned[: limit - 1].rstrip(" ,;:/-")
    return f"{trimmed}\u2026"


def _resolve_material_path(path_text: str | None) -> Path | None:
    if not path_text:
        return None
    candidate = Path(path_text)
    if candidate.is_absolute():
        return candidate if candidate.exists() else None
    resolved = candidate.resolve()
    return resolved if resolved.exists() else None


def _material_path_by_suffix(materials: list[ProjectMaterial], suffixes: tuple[str, ...]) -> Path | None:
    for material in materials:
        if material.material_type != BriefMaterialType.DOCUMENT or not material.path:
            continue
        if Path(material.path).suffix.lower() not in suffixes:
            continue
        resolved = _resolve_material_path(material.path)
        if resolved is not None:
            return resolved
    return None


def _parse_docx_blocks(docx_path: Path | None) -> list[LectureHeadingBlock]:
    if docx_path is None:
        return []
    try:
        from docx import Document
    except ImportError:
        return []
    document = Document(str(docx_path))
    blocks: list[LectureHeadingBlock] = []
    current: LectureHeadingBlock | None = None
    heading_counters = [0] * 6
    for paragraph in document.paragraphs:
        text = _normalize_text(paragraph.text)
        if not text:
            continue
        match = HEADING_RE.match(text)
        style_name = _normalize_text(getattr(getattr(paragraph, "style", None), "name", ""))
        style_match = re.match(r"heading\s+([1-6])$", style_name, re.IGNORECASE)
        if match or style_match:
            if current is not None:
                blocks.append(current)
            if match:
                number = match.group("number")
                title = match.group("title")
                level = number.count(".") + 1
                if 1 <= level <= len(heading_counters):
                    heading_counters[level - 1] += 1
                    for index in range(level, len(heading_counters)):
                        heading_counters[index] = 0
            else:
                level = int(style_match.group(1))
                heading_counters[level - 1] += 1
                for index in range(level, len(heading_counters)):
                    heading_counters[index] = 0
                number = ".".join(str(value) for value in heading_counters[:level] if value > 0)
                title = text
            current = LectureHeadingBlock(number=number, title=title, level=level)
            continue
        if current is not None:
            current.body.append(text)
    if current is not None:
        blocks.append(current)
    return blocks


def _build_pdf_page_index(pdf_path: Path | None) -> dict[int, str]:
    if pdf_path is None:
        return {}
    try:
        import fitz
    except ImportError:
        return {}
    pages: dict[int, str] = {}
    with fitz.open(pdf_path) as pdf:
        for index, page in enumerate(pdf, start=1):
            pages[index] = _normalize_text(page.get_text("text")).lower()
    return pages


def _find_heading_page(block: LectureHeadingBlock, pdf_pages: dict[int, str]) -> int | None:
    if not pdf_pages:
        return None
    probes = [_normalize_text(f"{block.number} {block.title}").lower(), _normalize_text(block.title).lower()]
    probes.extend(_normalize_text(line).lower() for line in block.body[:2])
    for page_number, text in pdf_pages.items():
        for probe in probes:
            if probe and probe[:24] in text:
                return page_number
    return None


def _lecture_language(section_title: str, child_titles: list[str]) -> str:
    joined = " ".join([section_title, *child_titles])
    return "ko" if _contains_hangul(joined) else "en"


def _lecture_copy(section_title: str, child_titles: list[str], *, korean: str, english: str) -> str:
    return korean.format(section=section_title, topics=", ".join(child_titles[:3])) if _lecture_language(section_title, child_titles) == "ko" else english.format(section=section_title, topics=", ".join(child_titles[:3]))


def _bullet_candidates_from_body(lines: list[str]) -> list[str]:
    bullets: list[str] = []
    for raw in lines:
        text = _normalize_text(raw)
        if not text or len(text) < 8 or HEADING_RE.match(text):
            continue
        text = re.sub(r"^[\-\u2022+\s]*", "", text)
        if text not in bullets:
            bullets.append(text)
    return bullets


def _compressed_bullets(lines: list[str], *, fallback_title: str, limit: int = LECTURE_MAX_BULLETS) -> list[str]:
    bullets: list[str] = []
    for line in lines:
        cleaned = _shorten_line(line)
        if cleaned and cleaned not in bullets:
            bullets.append(cleaned)
        if len(bullets) >= limit:
            break
    if bullets:
        return bullets
    if _contains_hangul(fallback_title):
        return [
            _shorten_line(f"{fallback_title}을 먼저 정의한 뒤 세부 설명으로 넘어간다."),
            _shorten_line(f"{fallback_title}이 강의 전체 구조에서 어떤 역할을 하는지 연결한다."),
            _shorten_line(f"{fallback_title}이 방법 선택에 어떤 영향을 주는지 정리한다."),
        ][:limit]
    return [
        _shorten_line(f"Define {fallback_title} before moving into detail."),
        _shorten_line(f"Connect {fallback_title} to the surrounding problem structure."),
        _shorten_line(f"State what {fallback_title} changes in method selection."),
    ][:limit]


def _section_visual_family(section_title: str, child_titles: list[str]) -> VisualType:
    title_text = " ".join([section_title, *child_titles]).lower()
    if any(
        token in title_text
        for token in (
            "algorithm",
            "flow",
            "process",
            "search",
            "gradient",
            "newton",
            "descent",
            "procedure",
            "알고리즘",
            "흐름",
            "과정",
            "절차",
            "탐색",
        )
    ):
        return VisualType.PROCESS
    if any(token in title_text for token in ("compare", "comparison", "trade-off", "pros", "cons", "vs", "비교", "장점", "단점", "트레이드오프")):
        return VisualType.COMPARISON
    if any(token in title_text for token in ("equation", "example", "kkt", "hessian", "taylor", "방정식", "예제", "헤시안", "테일러", "수식")):
        return VisualType.DOCUMENT_CROP
    return VisualType.FRAMEWORK


def _merge_materials(
    workflow_plan: WorkflowPlan,
    brief: WorkflowBriefInput | None,
) -> list[ProjectMaterial]:
    materials = list(workflow_plan.project_snapshot.current_materials)
    if brief is not None:
        materials.extend(brief.current_materials)
    unique: dict[tuple[str, str | None], ProjectMaterial] = {}
    for material in materials:
        key = (material.label, material.path)
        if key not in unique:
            unique[key] = material
    return list(unique.values())


def _material_refs(materials: list[ProjectMaterial]) -> list[SourceMaterialRef]:
    refs: list[SourceMaterialRef] = []
    for index, material in enumerate(materials, start=1):
        if material.path is None:
            continue
        refs.append(
            SourceMaterialRef(
                source_id=f"material-{index}-{_slugify(material.label)}",
                label=material.label,
                path=material.path,
                notes=material.notes,
            )
        )
    return refs


def _material_groups(materials: list[ProjectMaterial]) -> dict[BriefMaterialType, list[ProjectMaterial]]:
    grouped: dict[BriefMaterialType, list[ProjectMaterial]] = {}
    for material in materials:
        grouped.setdefault(material.material_type, []).append(material)
    return grouped


def _select_source_refs(
    grouped_materials: dict[BriefMaterialType, list[ProjectMaterial]],
    preferred_types: tuple[BriefMaterialType, ...],
    fallback_refs: list[SourceMaterialRef],
) -> list[SourceMaterialRef]:
    refs: list[SourceMaterialRef] = []
    for preferred_type in preferred_types:
        for material in grouped_materials.get(preferred_type, []):
            if material.path is None:
                continue
            refs.append(
                SourceMaterialRef(
                    source_id=f"{preferred_type.value}-{_slugify(material.label)}",
                    label=material.label,
                    path=material.path,
                    notes=material.notes,
                )
            )
    if refs:
        return refs[:2]
    return fallback_refs[:2]


def _communication_core(workflow_plan: WorkflowPlan) -> CommunicationCore:
    types = {workflow_plan.presentation_type_diagnosis.primary_type, *workflow_plan.presentation_type_diagnosis.secondary_types}
    if workflow_plan.workflow_option == "graduate-lecture-clustered":
        single_shift = "Leave able to explain the central concepts, their dependencies, and how they connect into one teaching arc."
        key_question = "What conceptual dependencies must the audience understand before they can explain or apply this material?"
        elevator_pitch_options = [
            "This deck teaches the core concepts first, then bridges them into mechanisms, examples, and limitations.",
            "This deck is organized as a pedagogical sequence rather than as a source-document summary.",
            "This deck keeps the main story focused on what the audience must understand before appendix evidence appears.",
        ]
        recommended_pitch_reason = "Graduate lectures land better when the audience hears the teaching arc before source detail or backup proof."
    elif PresentationType.DECISION in types:
        single_shift = "Leave the review with one decision path and the evidence needed to approve it."
        key_question = "Which path should this audience approve now, and what proof is sufficient to support it?"
        elevator_pitch_options = [
            "This deck recommends one decision path and the smallest proof set needed to approve it.",
            "This deck turns the current evidence into one practical recommendation and a bounded next-step sequence.",
            "This deck keeps the main story compact, separates backup proof into appendix space, and drives a clear approval ask.",
        ]
        recommended_pitch_reason = "Decision decks work best when the audience hears the choice and the minimum proof burden immediately."
    elif PresentationType.TRAINING in types or PresentationType.DEMO in types:
        single_shift = "Leave the session able to understand the new workflow and execute the live path confidently."
        key_question = "How will the audience understand the new workflow quickly enough to execute the live path with confidence?"
        elevator_pitch_options = [
            "This deck teaches the target workflow, the key screens, and the control points that matter in the live path.",
            "This deck explains the operating model first, then uses a small number of screens and diagrams to make the workflow executable.",
            "This deck keeps the walkthrough structured, avoids screenshot clutter, and ends with support and rollout guidance.",
        ]
        recommended_pitch_reason = "Training and demo material lands better when the story explains the model before it shows screens."
    else:
        single_shift = "Leave the audience with one clear narrative and the proof needed to support it."
        key_question = "What should the audience understand, believe, or do differently after this deck?"
        elevator_pitch_options = [
            "This deck tells one clear story, supports it with the strongest local proof, and closes on the next step.",
            "This deck turns the brief into a compact narrative with one main message per slide and appendix overflow kept separate.",
            "This deck uses a consistent visual system to make the main story memorable without overbuilding the evidence.",
        ]
        recommended_pitch_reason = "The default Gate 2 posture should keep the story compact and memorable before production expands it."
    recommended_elevator_pitch = elevator_pitch_options[0]
    return CommunicationCore(
        audience_outcome=f"{', '.join(workflow_plan.audience)} should leave with a clear next step.",
        deck_promise=workflow_plan.objective,
        single_decision_or_shift=single_shift,
        memory_line=workflow_plan.objective,
        key_question=key_question,
        elevator_pitch_options=elevator_pitch_options,
        recommended_elevator_pitch=recommended_elevator_pitch,
        recommended_pitch_reason=recommended_pitch_reason,
        supporting_themes=_dedupe(
            [
                workflow_plan.presentation_type_diagnosis.diagnosis_label,
                "One slide one message",
                "Main story before appendix",
            ]
        ),
    )


def _apply_slide_function_outline(
    slides: list[BlueprintSlide],
    slide_function_outline: SlideFunctionOutline | None,
) -> list[BlueprintSlide]:
    if slide_function_outline is None:
        return slides
    outlined_items = {item.slide_number: item for item in slide_function_outline.slides}
    updated: list[BlueprintSlide] = []
    for slide in slides:
        outline_item = outlined_items.get(slide.slide_number)
        if outline_item is None:
            updated.append(slide)
            continue
        content_budget_summary = dict(slide.content_budget_summary)
        content_budget_summary.update(
            {
                "planner_budget": outline_item.content_budget.model_dump(mode="json"),
                "slide_function": outline_item.slide_function.value,
                "outline_section": outline_item.section,
                "title_hint": outline_item.title_hint,
            }
        )
        authoring_payload = dict(slide.authoring_payload)
        authoring_payload["planner_outline"] = {
            "slide_function": outline_item.slide_function.value,
            "section": outline_item.section,
            "title_hint": outline_item.title_hint,
            "content_budget": outline_item.content_budget.model_dump(mode="json"),
            "rationale": outline_item.rationale,
        }
        updated.append(
            slide.model_copy(
                update={
                    "content_budget_summary": content_budget_summary,
                    "authoring_payload": authoring_payload,
                }
            )
        )
    return updated


def _story_architecture(workflow_plan: WorkflowPlan) -> list[StorySection]:
    types = {workflow_plan.presentation_type_diagnosis.primary_type, *workflow_plan.presentation_type_diagnosis.secondary_types}
    main_story = workflow_plan.main_story_slide_count_range
    appendix = workflow_plan.appendix_candidate_slide_count_range

    if workflow_plan.workflow_option == "evidence-backed-core":
        proof_start = max(2, main_story.start - 3)
        proof_end = max(proof_start, main_story.end - 3)
        sections = [
            StorySection(
                section_id="core-claim",
                title="Core Claim",
                purpose="State the main claim and frame what the audience should believe before the proof stack begins.",
                role_in_story="Anchor the story with one explicit claim and the minimum framing needed to evaluate the proof.",
                slide_count_range=CountRange(start=1 if main_story.start <= 4 else 2, end=2),
                slide_roles=[SlideRole.EXECUTIVE_SUMMARY, SlideRole.ANALYSIS],
            ),
            StorySection(
                section_id="proof-modules",
                title="Proof Modules",
                purpose="Sequence the strongest proof units so each one adds confidence without reopening the whole appendix.",
                role_in_story="Carry the evidentiary burden with multiple proof modules that stay adjacent to the core claim.",
                slide_count_range=CountRange(start=proof_start, end=proof_end),
                slide_roles=[SlideRole.EVIDENCE, SlideRole.COMPARISON, SlideRole.ANALYSIS],
            ),
            StorySection(
                section_id="implications",
                title="Implications",
                purpose="Close on what the proof set means for the audience and what belongs outside the main story.",
                role_in_story="Translate the proof stack into the retained implication or recommended move.",
                slide_count_range=CountRange(start=1, end=2),
                slide_roles=[SlideRole.ANALYSIS, SlideRole.RECOMMENDATION],
            ),
        ]
    elif workflow_plan.workflow_option == "thesis-proof-close":
        proof_start = max(2, main_story.start - 2)
        proof_end = max(proof_start, main_story.end - 2)
        sections = [
            StorySection(
                section_id="thesis",
                title="Thesis",
                purpose="State the investor-facing thesis in one sentence before the supporting proof stack begins.",
                role_in_story="Anchor the audience on the thesis that the proof spine and close must support.",
                slide_count_range=CountRange(start=1, end=1),
                slide_roles=[SlideRole.EXECUTIVE_SUMMARY],
            ),
            StorySection(
                section_id="proof-spine",
                title="Proof Spine",
                purpose="Carry the bounded proof chain that makes the thesis credible without reopening appendix backup.",
                role_in_story="Sequence the strongest proof units so the audience can track how the thesis is substantiated.",
                slide_count_range=CountRange(start=proof_start, end=proof_end),
                slide_roles=[SlideRole.EVIDENCE, SlideRole.COMPARISON, SlideRole.ANALYSIS],
            ),
            StorySection(
                section_id="close",
                title="Close",
                purpose="Convert the proof spine into one explicit ask or conclusion.",
                role_in_story="Close the story on the financing ask or decision that follows naturally from the thesis and proof.",
                slide_count_range=CountRange(start=1, end=1),
                slide_roles=[SlideRole.RECOMMENDATION],
            ),
        ]
    elif PresentationType.TRAINING in types or PresentationType.DEMO in types:
        sections = [
            StorySection(
                section_id="orientation",
                title="Orientation",
                purpose="Frame the learning objective and the workflow change.",
                role_in_story="Set the learning objective and define what will change for the audience.",
                slide_count_range=CountRange(start=2, end=3),
                slide_roles=[SlideRole.EXECUTIVE_SUMMARY, SlideRole.ANALYSIS],
            ),
            StorySection(
                section_id="workflow",
                title="Workflow Model",
                purpose="Explain the target process and key steps.",
                role_in_story="Teach the target workflow in a structured, repeatable way.",
                slide_count_range=CountRange(start=4, end=max(5, main_story.end // 2)),
                slide_roles=[SlideRole.PROCESS, SlideRole.INFOGRAPHIC, SlideRole.EVIDENCE],
            ),
            StorySection(
                section_id="demo",
                title="Demo Sequence",
                purpose="Walk through the live path without turning the deck into raw screenshots.",
                role_in_story="Ground the model in a small number of concrete screen or proof moments.",
                slide_count_range=CountRange(start=3, end=max(4, main_story.end // 3)),
                slide_roles=[SlideRole.SECTION_DIVIDER, SlideRole.EVIDENCE, SlideRole.PROCESS],
            ),
            StorySection(
                section_id="rollout",
                title="Rollout And Support",
                purpose="Define responsibilities, controls, and next steps.",
                role_in_story="Close on operating guardrails, support, and the handoff after training.",
                slide_count_range=CountRange(start=2, end=3),
                slide_roles=[SlideRole.COMPARISON, SlideRole.RECOMMENDATION],
            ),
        ]
    elif PresentationType.REPORT in types and PresentationType.DECISION in types:
        sections = [
            StorySection(
                section_id="executive-summary",
                title="Executive Summary",
                purpose="Lead with the decision and the smallest supporting case.",
                role_in_story="Answer the decision before the audience starts evaluating methods or backup evidence.",
                slide_count_range=CountRange(start=2, end=3),
                slide_roles=[SlideRole.EXECUTIVE_SUMMARY, SlideRole.RECOMMENDATION],
            ),
            StorySection(
                section_id="why-now",
                title="Why Now",
                purpose="Frame the performance context and the need to act.",
                role_in_story="Make the cost of waiting or staying with the current state explicit.",
                slide_count_range=CountRange(start=2, end=3),
                slide_roles=[SlideRole.ANALYSIS, SlideRole.EVIDENCE],
            ),
            StorySection(
                section_id="proof",
                title="Proof And Options",
                purpose="Show the evidence, tradeoffs, and option comparison that support the recommendation.",
                role_in_story="Carry the analytical burden that makes the recommendation credible.",
                slide_count_range=CountRange(start=3, end=max(4, main_story.end - 3)),
                slide_roles=[SlideRole.EVIDENCE, SlideRole.COMPARISON, SlideRole.ANALYSIS],
            ),
            StorySection(
                section_id="next-steps",
                title="Recommendation And Next Steps",
                purpose="Lock the recommendation, rollout logic, and guardrails.",
                role_in_story="Translate the decision into a bounded execution path and approval ask.",
                slide_count_range=CountRange(start=1, end=2),
                slide_roles=[SlideRole.RECOMMENDATION, SlideRole.PROCESS],
            ),
        ]
    elif PresentationType.EXPLAINER in types and PresentationType.PERSUASION in types:
        sections = [
            StorySection(
                section_id="problem",
                title="Problem Framing",
                purpose="Clarify the problem and why the audience should care now.",
                role_in_story="Earn attention by naming the problem in business terms the audience already recognizes.",
                slide_count_range=CountRange(start=2, end=3),
                slide_roles=[SlideRole.EXECUTIVE_SUMMARY, SlideRole.ANALYSIS],
            ),
            StorySection(
                section_id="mechanism",
                title="Mechanism And Proof",
                purpose="Explain the mechanism and support it with evidence.",
                role_in_story="Show how the idea works and why the audience should trust it.",
                slide_count_range=CountRange(start=2, end=3),
                slide_roles=[SlideRole.EVIDENCE, SlideRole.INFOGRAPHIC],
            ),
            StorySection(
                section_id="ask",
                title="Recommendation",
                purpose="State the ask and the expected next step.",
                role_in_story="Convert understanding into one explicit audience commitment.",
                slide_count_range=CountRange(start=1, end=2),
                slide_roles=[SlideRole.RECOMMENDATION, SlideRole.PROCESS],
            ),
        ]
    else:
        sections = [
            StorySection(
                section_id="core-story",
                title="Core Story",
                purpose="Carry the main message and supporting proof.",
                role_in_story="Carry the main narrative without letting supporting detail fracture the message.",
                slide_count_range=main_story,
                slide_roles=[SlideRole.EXECUTIVE_SUMMARY, SlideRole.ANALYSIS, SlideRole.EVIDENCE],
            )
        ]

    if appendix.end > 0:
        sections.append(
            StorySection(
                section_id="appendix",
                title="Appendix",
                purpose="Hold methods, references, backup evidence, and overflow details.",
                role_in_story="Keep methods, references, and backup proof outside the approved main story.",
                deck_mode=DeckMode.APPENDIX,
                slide_count_range=CountRange(start=max(1, appendix.start), end=max(1, appendix.end)),
                slide_roles=[SlideRole.APPENDIX_EVIDENCE, SlideRole.REFERENCES],
            )
        )
    return sections


def _selected_workflow_option(workflow_plan: WorkflowPlan) -> WorkflowOption:
    for option in workflow_plan.workflow_options:
        if option.option_id == workflow_plan.workflow_option:
            return option
    raise ValueError(f"workflow option {workflow_plan.workflow_option!r} was not found in workflow_plan.workflow_options")


def _chosen_workflow_phases(workflow_plan: WorkflowPlan, story_architecture: list[StorySection]) -> list[WorkflowPhase]:
    selected = _selected_workflow_option(workflow_plan)
    if selected.phases:
        return selected.phases
    phases: list[WorkflowPhase] = []
    for index, section in enumerate(story_architecture, start=1):
        phases.append(
            WorkflowPhase(
                phase_id=f"workflow-phase-{index:02d}",
                label=section.title,
                objective=section.role_in_story or section.purpose,
                expected_outputs=[f"Slides aligned to the {section.title} section."],
            )
        )
    return phases


def _story_structure(workflow_plan: WorkflowPlan, story_architecture: list[StorySection]) -> StoryArchitectureSummary:
    section_titles = [section.title for section in story_architecture if section.deck_mode == DeckMode.MAIN_STORY]
    diagnosis = workflow_plan.presentation_type_diagnosis.diagnosis_label
    if workflow_plan.workflow_option == "graduate-lecture-clustered":
        overall_narrative = (
            f"Move the audience through a {diagnosis} teaching sequence that starts with conceptual orientation, "
            "builds prerequisite understanding, then connects mechanisms, examples, and limitations without mirroring source headings."
        )
        message_flow = [f"{section.title}: {section.purpose}" for section in story_architecture]
        audience_journey = [
            "Orient the audience to the teaching question and the conceptual dependency map.",
            "Build understanding by moving from anchor concepts to bridges, mechanisms, and worked examples.",
            "Close with integration while keeping source-heavy support in appendix-only territory.",
        ]
        fit_rationale = (
            f"This architecture fits {', '.join(workflow_plan.audience)} because it keeps the approved `{workflow_plan.workflow_option}` "
            f"workflow concept-led, respects the {workflow_plan.scale_mode.value} scale target, and preserves one-slide-one-message discipline across "
            f"{', '.join(section_titles) or 'the main story'}."
        )
    else:
        overall_narrative = (
            f"Move the audience through a {diagnosis} narrative that starts with the main ask, builds only the proof needed to support it, "
            "and keeps methods or overflow detail outside the core story."
        )
        message_flow = [f"{section.title}: {section.purpose}" for section in story_architecture]
        audience_journey = [
            "Orient the audience to the objective and what decision or shift matters most.",
            "Build confidence with the smallest set of supporting evidence, options, or process structure.",
            "Close on the recommended move and keep overflow detail in appendix-only territory.",
        ]
        fit_rationale = (
            f"This architecture fits {', '.join(workflow_plan.audience)} because it mirrors the approved `{workflow_plan.workflow_option}` workflow, "
            f"respects the {workflow_plan.scale_mode.value} scale target, and preserves one-slide-one-message discipline across {', '.join(section_titles) or 'the main story'}."
        )
    return StoryArchitectureSummary(
        overall_narrative=overall_narrative,
        message_flow=message_flow,
        audience_journey=audience_journey,
        fit_rationale=fit_rationale,
    )


def _visual_routes(
    workflow_plan: WorkflowPlan,
    reference_dna: ReferenceDNA | None,
    grouped_materials: dict[BriefMaterialType, list[ProjectMaterial]],
) -> list[VisualRoute]:
    routes = [
        VisualRoute(
            route_id="evidence-led-editorial",
            label="Evidence-led editorial",
            description="Pair short thesis titles with one dominant proof object and restrained annotations.",
            mood="Calm, analytical, and executive-facing.",
            best_use_case="Best for report, decision, and persuasion decks that already have credible local proof objects.",
            dominant_visual_types=[VisualType.TEXT, VisualType.DOCUMENT_CROP, VisualType.CHART],
            layout_bias=["Title plus evidence", "Right-rail insight callouts", "Sparse captions"],
            key_visual_traits=["Short thesis titles", "One dominant proof object", "Restrained annotation rails"],
            strengths=["Keeps the main story compact.", "Lets local proof do visible credibility work."],
            risks=["Depends on strong local crops or source visuals.", "Can feel underpowered if no concrete proof survives review."],
            tradeoffs=["Relies on high-quality local source material.", "Can underperform if charts need to carry too much explanation."],
            fit_notes=["Strong fit when PDFs, screenshots, or report visuals already exist locally."],
        ),
        VisualRoute(
            route_id="structured-visual-system",
            label="Structured visual system",
            description="Rebuild most evidence as slide-native charts, tables, frameworks, and timelines.",
            mood="Systematic, precise, and slide-native.",
            best_use_case="Best when data structure and repeatable comparison logic matter more than source-document reuse.",
            dominant_visual_types=[VisualType.CHART, VisualType.TABLE, VisualType.FRAMEWORK, VisualType.TIMELINE],
            layout_bias=["Chart-first analysis", "Framework grids", "Consistent data callout patterns"],
            key_visual_traits=["Slide-native charts", "Structured comparison tables", "Repeatable infographic grammar"],
            strengths=["Creates a highly consistent visual system.", "Simplifies complex evidence into repeatable forms."],
            risks=["Requires more production work later.", "Can lose credibility if the audience expects source proof to stay visible."],
            tradeoffs=["Requires more structured rendering later.", "Can feel abstract if no concrete source proof appears."],
            fit_notes=["Best when data and comparison logic matter more than source screenshots."],
        ),
    ]
    if grouped_materials.get(BriefMaterialType.IMAGE) or grouped_materials.get(BriefMaterialType.DECK):
        routes.append(
            VisualRoute(
                route_id="annotated-screen-and-proof",
                label="Annotated screen and proof",
                description="Use cropped screens or slides as evidence anchors, then simplify them with annotations and paired explanatory text.",
                mood="Concrete, guided, and demo-oriented.",
                best_use_case="Best for training, demo, and workflow adoption decks that need a few concrete screens without turning into raw screenshots.",
                dominant_visual_types=[VisualType.DOCUMENT_CROP, VisualType.PROCESS, VisualType.INFOGRAPHIC],
                layout_bias=["Large screenshot frames", "Annotation rail", "Divider-led pacing"],
                key_visual_traits=["Large anchored screens", "Tight annotation zones", "Divider-led pacing resets"],
                strengths=["Grounds abstract process talk in concrete screens.", "Supports demo and enablement use cases cleanly."],
                risks=["Can become screenshot-heavy quickly.", "Needs disciplined crop review to stay readable."],
                tradeoffs=["Needs tight crop discipline.", "Can become screenshot-heavy if not controlled."],
                fit_notes=["Useful for training, demo, and screenshot-rich reference packs."],
            )
        )
    if reference_dna is not None and reference_dna.source_family:
        routes[0].fit_notes.append(f"Reference fit: {reference_dna.source_family}.")
    return routes


def _recommended_visual_route(
    workflow_plan: WorkflowPlan,
    reference_dna: ReferenceDNA | None,
    grouped_materials: dict[BriefMaterialType, list[ProjectMaterial]],
) -> str:
    types = {workflow_plan.presentation_type_diagnosis.primary_type, *workflow_plan.presentation_type_diagnosis.secondary_types}
    if PresentationType.TRAINING in types or PresentationType.DEMO in types:
        return "annotated-screen-and-proof" if grouped_materials.get(BriefMaterialType.IMAGE) or grouped_materials.get(BriefMaterialType.DECK) else "structured-visual-system"
    if reference_dna is not None and "evidence" in " ".join(reference_dna.patterns_worth_borrowing).lower():
        return "evidence-led-editorial"
    if grouped_materials.get(BriefMaterialType.SPREADSHEET) or grouped_materials.get(BriefMaterialType.DATA):
        return "structured-visual-system"
    return "evidence-led-editorial"


def _recommended_route_reason(
    workflow_plan: WorkflowPlan,
    recommended_route: str,
    reference_dna: ReferenceDNA | None,
    grouped_materials: dict[BriefMaterialType, list[ProjectMaterial]],
) -> str:
    types = {workflow_plan.presentation_type_diagnosis.primary_type, *workflow_plan.presentation_type_diagnosis.secondary_types}
    if recommended_route == "annotated-screen-and-proof":
        return "The deck direction benefits from concrete screens or prior-slide references, and the audience likely needs a guided walkthrough more than abstract structure."
    if recommended_route == "structured-visual-system":
        return "The available materials favor structured data and repeatable comparison logic, so a slide-native system will stay clearer than source-heavy proof slides."
    if reference_dna is not None:
        return f"The reference pack points toward calm evidence framing, and the current materials support showing proof directly without copying any one template from {reference_dna.source_family}."
    if PresentationType.DECISION in types or PresentationType.REPORT in types:
        return "Decision-oriented decks usually benefit from short thesis slides plus one visible proof object per slide."
    return "A compact editorial route keeps the main story legible while leaving room for appendix overflow if the deck expands later."


def _visual_reference_summary(
    reference_dna: ReferenceDNA | None,
    visual_routes: list[VisualRoute],
    recommended_route: str,
) -> VisualReferenceSummary:
    chosen_route = next(route for route in visual_routes if route.route_id == recommended_route)
    if reference_dna is None:
        scanned = ["No reference_dna supplied; the planner used local deterministic defaults and available materials only."]
        patterns = ["Default executive-spacing discipline", "One dominant message zone per slide", "Appendix separation for overflow detail"]
    else:
        scanned = [
            f"{source.label} ({source.material_type.value})"
            for source in reference_dna.source_files
        ]
        patterns = _dedupe(
            [
                reference_dna.source_family,
                *reference_dna.patterns_worth_borrowing[:3],
                *reference_dna.layout_logic[:2],
            ]
        )
    rejected = []
    for route in visual_routes:
        if route.route_id == recommended_route:
            continue
        reason = route.risks[0] if route.risks else (route.tradeoffs[0] if route.tradeoffs else "Not selected for the current deck direction.")
        rejected.append(f"{route.label}: {reason}")
    return VisualReferenceSummary(
        reference_sources_scanned=scanned,
        pattern_families_reviewed=patterns,
        chosen_direction=f"{chosen_route.label}: {chosen_route.description}",
        rejected_directions=rejected,
    )


def _target_slide_counts(workflow_plan: WorkflowPlan, materials: list[ProjectMaterial]) -> tuple[int, int]:
    if workflow_plan.workflow_option == "graduate-lecture-clustered":
        main_story_range = workflow_plan.main_story_slide_count_range
        appendix_range = workflow_plan.appendix_candidate_slide_count_range
        main_story_target = min(main_story_range.end, max(main_story_range.start, LECTURE_MAIN_TARGET))
        appendix_target = min(appendix_range.end, max(appendix_range.start, LECTURE_APPENDIX_TARGET))
        return main_story_target, appendix_target

    main_story_target = workflow_plan.smallest_effective_slide_count
    appendix_range = workflow_plan.appendix_candidate_slide_count_range
    if appendix_range.end <= 0:
        appendix_target = 0
    elif workflow_plan.scale_mode == ScaleMode.COMPACT:
        appendix_target = 1 if materials else 0
    elif workflow_plan.scale_mode == ScaleMode.MEGA_DECK:
        appendix_target = max(appendix_range.start, min(appendix_range.end, 8))
    elif workflow_plan.scale_mode == ScaleMode.LARGE_DECK:
        appendix_target = max(appendix_range.start, min(appendix_range.end, 4))
    elif workflow_plan.scale_mode == ScaleMode.EXTENDED:
        appendix_target = max(appendix_range.start, min(appendix_range.end, 3))
    else:
        appendix_target = max(appendix_range.start, min(appendix_range.end, 2))
    return main_story_target, appendix_target


def _lecture_orientation_specs(workflow_plan: WorkflowPlan) -> list[dict[str, Any]]:
    korean = _contains_hangul(f"{workflow_plan.deck_title} {' '.join(workflow_plan.audience)}")
    deck_title = workflow_plan.deck_title
    if korean:
        return [
            {
                "section": "Orientation",
                "role": SlideRole.TITLE,
                "title": deck_title,
                "takeaway": _shorten_line("강의의 큰 흐름을 먼저 맞춘 뒤 세부 내용으로 들어간다.", 120),
                "message": _shorten_line("이번 강의는 문제 정식화, 해석 조건, 방법 선택의 순서로 전개된다.", 140),
                "visual": VisualType.TEXT,
                "core_content": _compressed_bullets(
                    [
                        "알고리즘보다 먼저 최적화 문제를 어떻게 쓸지 정리한다.",
                        "해석 조건을 통해 해의 의미를 읽는다.",
                        "유도 과정과 보조 표는 appendix로 분리한다.",
                    ],
                    fallback_title=deck_title,
                ),
                "required_assets": [deck_title],
                "content_tier": ContentTier.LECTURE_CORE,
            },
            {
                "section": "Orientation",
                "role": SlideRole.PROCESS,
                "title": "강의 로드맵",
                "takeaway": _shorten_line("세 단계의 teaching flow로 이해를 쌓는다.", 120),
                "message": _shorten_line("모형을 먼저 세우고, 조건을 해석한 뒤, 마지막에 방법을 비교한다.", 140),
                "visual": VisualType.PROCESS,
                "core_content": _compressed_bullets(
                    ["문제 정식화", "해석 조건", "방법 비교와 선택"],
                    fallback_title="강의 로드맵",
                ),
                "required_assets": ["강의 단계 맵"],
                "content_tier": ContentTier.LECTURE_CORE,
            },
            {
                "section": "Orientation",
                "role": SlideRole.COMPARISON,
                "title": "모든 슬라이드에서 붙잡을 질문",
                "takeaway": _shorten_line("같은 질문으로 개념, 유도, 알고리즘을 연결한다.", 120),
                "message": _shorten_line("무엇을 최적화하는지, 무엇이 제약하는지, 어떤 방법이 구조에 맞는지 계속 묻는다.", 140),
                "visual": VisualType.COMPARISON,
                "core_content": _compressed_bullets(
                    ["무엇을 최적화하는가", "어떤 변수와 제약이 중요한가", "어떤 방법이 구조에 맞는가"],
                    fallback_title="모든 슬라이드에서 붙잡을 질문",
                ),
                "required_assets": ["핵심 질문 프레임"],
                "content_tier": ContentTier.LECTURE_CORE,
            },
            {
                "section": "Orientation",
                "role": SlideRole.ANALYSIS,
                "title": "강의 제작 원칙",
                "takeaway": _shorten_line("메인 스토리에는 teaching copy만 남기고 세부 자료는 appendix로 보낸다.", 120),
                "message": _shorten_line("긴 유도, 긴 bullet, 원문 crop, slide-native visual의 역할을 분리해 운영한다.", 140),
                "visual": VisualType.FRAMEWORK,
                "core_content": _compressed_bullets(
                    [
                        "메인 스토리의 takeaway는 1-3개의 짧은 bullet로 제한한다.",
                        "원문 표기나 수식이 필요하면 source crop을 사용한다.",
                        "추가 설명과 참고 노트는 appendix에 둔다.",
                    ],
                    fallback_title="강의 제작 원칙",
                ),
                "required_assets": ["강의 제작 원칙"],
                "content_tier": ContentTier.LECTURE_CORE,
            },
        ]
    return [
        {
            "section": "Orientation",
            "role": SlideRole.TITLE,
            "title": deck_title,
            "takeaway": _shorten_line("Align the lecture arc before moving into details.", 120),
            "message": _shorten_line("The lecture moves from problem framing to analytic conditions and then to method selection.", 140),
            "visual": VisualType.TEXT,
            "core_content": _compressed_bullets(
                [
                    "Frame the optimization problem before discussing algorithms.",
                    "Use analytic conditions to interpret what the solution means.",
                    "Move derivations and backup tables into appendix support.",
                ],
                fallback_title=deck_title,
            ),
            "required_assets": [deck_title],
            "content_tier": ContentTier.LECTURE_CORE,
        },
        {
            "section": "Orientation",
            "role": SlideRole.PROCESS,
            "title": "Lecture roadmap",
            "takeaway": _shorten_line("Build understanding in three teaching phases.", 120),
            "message": _shorten_line("Start with the model, interpret the conditions, and only then compare methods.", 140),
            "visual": VisualType.PROCESS,
            "core_content": _compressed_bullets(
                ["Problem formulation", "Analytic conditions", "Method comparison and selection"],
                fallback_title="Lecture roadmap",
            ),
            "required_assets": ["Lecture phase map"],
            "content_tier": ContentTier.LECTURE_CORE,
        },
        {
            "section": "Orientation",
            "role": SlideRole.COMPARISON,
            "title": "Questions to keep on every slide",
            "takeaway": _shorten_line("Use the same questions to connect concepts, derivations, and algorithms.", 120),
            "message": _shorten_line("Ask what is optimized, what constrains it, and which method matches the structure.", 140),
            "visual": VisualType.COMPARISON,
            "core_content": _compressed_bullets(
                ["What is being optimized", "Which variables and constraints matter", "Which method fits the structure"],
                fallback_title="Questions to keep on every slide",
            ),
            "required_assets": ["Lecture question frame"],
            "content_tier": ContentTier.LECTURE_CORE,
        },
        {
            "section": "Orientation",
            "role": SlideRole.ANALYSIS,
            "title": "Reading rules for the deck",
            "takeaway": _shorten_line("Keep slide copy short and move dense support elsewhere.", 120),
            "message": _shorten_line("Visible slides carry teaching copy only; notes and appendix hold the overflow detail.", 140),
            "visual": VisualType.FRAMEWORK,
            "core_content": _compressed_bullets(
                ["Use one main message per slide.", "Keep visible bullets short and few.", "Treat appendix as support, not main-story overflow."],
                fallback_title="Reading rules for the deck",
            ),
            "required_assets": ["Lecture reading rules"],
            "content_tier": ContentTier.LECTURE_CORE,
        },
    ]


def _lecture_appendix_specs(
    sections: list[LectureHeadingBlock],
    children_by_parent: dict[str, list[LectureHeadingBlock]],
    pdf_label: str | None,
    pdf_path: str | None,
    appendix_target: int,
) -> tuple[list[dict[str, Any]], list[str]]:
    specs: list[dict[str, Any]] = []
    decisions: list[str] = []
    for section in sections:
        children = children_by_parent.get(section.number, [])
        for child in children[:2]:
            content = _compressed_bullets(_bullet_candidates_from_body(child.body), fallback_title=child.title)
            source_refs = []
            if pdf_path is not None:
                source_refs.append(
                    SourceMaterialRef(
                        source_id=f"appendix-{child.number.replace('.', '-')}",
                        label=pdf_label or Path(pdf_path).name,
                        path=pdf_path,
                        page=child.page_hint,
                    )
                )
            specs.append(
                {
                    "section": "Appendix",
                    "role": SlideRole.APPENDIX_EVIDENCE if source_refs else SlideRole.REFERENCES,
                    "title": child.title,
                    "takeaway": content[0],
                    "message": content[0],
                    "visual": VisualType.DOCUMENT_CROP if source_refs else VisualType.TABLE,
                    "core_content": content,
                    "required_assets": [child.title],
                    "content_tier": ContentTier.APPENDIX_ONLY,
                    "deck_mode": DeckMode.APPENDIX,
                    "source_material_refs": source_refs,
                    "presenter_notes": _shorten_line("Appendix support slide generated from detailed source material.", 160),
                }
            )
            decisions.append(f"Moved detailed support from {section.number} {section.title} into appendix via {child.number} {child.title}.")
            if len(specs) >= appendix_target - 1:
                break
        if len(specs) >= appendix_target - 1:
            break
    specs.append(
        {
            "section": "Appendix",
            "role": SlideRole.REFERENCES,
            "title": "Source map and backup trail",
            "takeaway": _shorten_line("Use the appendix to recover source formulas, tables, and detailed derivations.", 120),
            "message": _shorten_line("The appendix preserves traceability for material that supports but should not interrupt the lecture arc.", 140),
            "visual": VisualType.TEXT,
            "core_content": _compressed_bullets(
                [
                    "Keep detailed derivations outside the main lecture sequence.",
                    "Use source-linked crops when original notation matters.",
                    "Reserve backup comparisons and tables for post-lecture reference.",
                ]
                if not _contains_hangul(" ".join(section.title for section in sections[:2]))
                else [
                    "상세 유도는 메인 강의 흐름 밖의 appendix로 둔다.",
                    "원래 표기와 도식이 중요하면 source-linked crop을 사용한다.",
                    "보조 비교표와 세부 표는 강의 후반 reference용 appendix로 보낸다.",
                ],
                fallback_title="Source map and backup trail",
            ),
            "required_assets": ["Source map"],
            "content_tier": ContentTier.APPENDIX_ONLY,
            "deck_mode": DeckMode.APPENDIX,
        }
    )
    return specs[:appendix_target], decisions


def _lecture_specs_from_documents(
    workflow_plan: WorkflowPlan,
    materials: list[ProjectMaterial],
    main_story_target: int,
    appendix_target: int,
) -> tuple[list[dict[str, Any]], list[str]]:
    docx_path = _material_path_by_suffix(materials, (".docx",))
    pdf_path = _material_path_by_suffix(materials, (".pdf",))
    blocks = _parse_docx_blocks(docx_path)
    if not blocks:
        return [], []

    pdf_pages = _build_pdf_page_index(pdf_path)
    for block in blocks:
        block.page_hint = _find_heading_page(block, pdf_pages)

    level1_blocks = [block for block in blocks if block.level == 1]
    level2_blocks = [block for block in blocks if block.level == 2]
    if not level2_blocks:
        level2_blocks = level1_blocks
    children_by_parent: dict[str, list[LectureHeadingBlock]] = {}
    for block in blocks:
        if block.level >= 3 and block.parent_key is not None:
            children_by_parent.setdefault(block.parent_key, []).append(block)

    major_titles = {block.chapter_key: block.title for block in level1_blocks}
    pdf_label = pdf_path.name if pdf_path is not None else docx_path.name if docx_path is not None else workflow_plan.deck_title
    pdf_ref_path = str(pdf_path) if pdf_path is not None else str(docx_path) if docx_path is not None else None
    specs: list[dict[str, Any]] = _lecture_orientation_specs(workflow_plan)
    clustering_decisions = [
        "Activated graduate lecture clustering mode for document-based planning.",
        "Collapsed subsection mirroring into section-level concept clusters with 2-4 slides each.",
    ]

    prior_chapter_key: str | None = None
    for section in level2_blocks:
        children = children_by_parent.get(section.number, [])
        child_titles = [child.title for child in children]
        if prior_chapter_key != section.chapter_key:
            chapter_title = major_titles.get(section.chapter_key, section.title)
            specs.append(
                {
                    "section": chapter_title,
                    "role": SlideRole.SECTION_DIVIDER,
                    "title": chapter_title,
                    "takeaway": _lecture_copy(
                        chapter_title,
                        child_titles or [section.title],
                        korean="{section}을 다음 teaching phase로 사용한다.",
                        english="Use {section} as the next teaching phase in the lecture arc.",
                    ),
                    "message": _lecture_copy(
                        chapter_title,
                        child_titles or [section.title],
                        korean="{section}에 들어가기 전에 호흡을 다시 맞추고 다음 개념 cluster를 연다.",
                        english="Reset pacing before the next concept cluster begins in {section}.",
                    ),
                    "visual": VisualType.TEXT,
                    "core_content": [],
                    "required_assets": [],
                    "content_tier": ContentTier.LECTURE_CORE,
                }
            )
            prior_chapter_key = section.chapter_key

        section_language = _lecture_language(section.title, child_titles)
        base_lines = _bullet_candidates_from_body(section.body)
        child_lines = [child.title for child in children] + [line for child in children for line in child.body[:2]]
        concept_bullets = _compressed_bullets(base_lines + child_titles + child_lines, fallback_title=section.title)
        source_refs = []
        if pdf_ref_path is not None:
            source_refs.append(
                SourceMaterialRef(
                    source_id=f"section-{section.number.replace('.', '-')}",
                    label=pdf_label,
                    path=pdf_ref_path,
                    page=section.page_hint,
                )
            )
        section_visual = _section_visual_family(section.title, child_titles)
        cluster_id = f"cluster-{section.number.replace('.', '-')}"
        section_id = _slugify(section.title)
        specs.extend(
            [
                {
                    "section": section.title,
                    "section_id": section_id,
                    "cluster_id": cluster_id,
                    "role": SlideRole.ANALYSIS,
                    "title": section.title,
                    "takeaway": _lecture_copy(
                        section.title,
                        child_titles,
                        korean="{section}은 subsection 목록이 아니라 하나의 concept cluster로 다룬다.",
                        english="Treat {section} as one concept cluster instead of a list of subsections.",
                    ),
                    "message": concept_bullets[0],
                    "visual": VisualType.FRAMEWORK,
                    "core_content": concept_bullets,
                    "required_assets": [section.title, *child_titles[:2]],
                    "content_tier": ContentTier.LECTURE_CORE,
                    "presenter_notes": _shorten_line(
                        " / ".join(child_titles[:4]) if child_titles else "Introduce the concept cluster before moving into detail.",
                        160,
                    ),
                    "source_material_refs": source_refs[:1],
                },
                {
                    "section": section.title,
                    "section_id": section_id,
                    "cluster_id": cluster_id,
                    "role": SlideRole.ANALYSIS if section_visual != VisualType.PROCESS else SlideRole.PROCESS,
                    "title": _lecture_copy(
                        section.title,
                        child_titles,
                        korean="{section}의 핵심 구조",
                        english="Core structure of {section}",
                    ),
                    "takeaway": _lecture_copy(
                        section.title,
                        child_titles,
                        korean="{section}를 설명하는 데 필요한 기준만 남긴다.",
                        english="Keep only the criteria needed to teach {section} cleanly.",
                    ),
                    "message": concept_bullets[min(1, len(concept_bullets) - 1)],
                    "visual": section_visual if section_visual != VisualType.DOCUMENT_CROP else VisualType.PROCESS,
                    "core_content": concept_bullets,
                    "required_assets": [section.title, *child_titles[:3]],
                    "content_tier": ContentTier.LECTURE_CORE,
                    "presenter_notes": _shorten_line(
                        "이 슬라이드를 해당 섹션의 핵심 teaching slide로 사용한다." if section_language == "ko" else "Use this as the main teaching slide for the section.",
                        160,
                    ),
                },
            ]
        )
        example_bullets = _compressed_bullets(child_lines, fallback_title=section.title)
        specs.append(
            {
                "section": section.title,
                "section_id": section_id,
                "cluster_id": cluster_id,
                "role": SlideRole.EVIDENCE if source_refs else SlideRole.COMPARISON,
                "title": _lecture_copy(
                    section.title,
                    child_titles,
                    korean="{section} 예제 적용",
                    english="Worked example for {section}",
                ),
                "takeaway": _lecture_copy(
                    section.title,
                    child_titles,
                    korean="{section}는 예제 하나나 source proof 하나로 고정해 설명한다.",
                    english="Use one worked example or source proof to anchor {section}.",
                ),
                "message": example_bullets[0],
                "visual": VisualType.DOCUMENT_CROP if source_refs else VisualType.COMPARISON,
                "core_content": example_bullets,
                "required_assets": [section.title, *(child_titles[:2] or concept_bullets[:2])],
                "content_tier": ContentTier.SUPPORTING_EXAMPLE,
                "source_material_refs": source_refs,
                "presenter_notes": _shorten_line(
                    "이 source proof는 teaching clarity를 높일 때만 쓰고, 나머지는 appendix로 넘긴다."
                    if section_language == "ko"
                    else "Use the source proof here only if it improves teaching clarity; move the rest to appendix.",
                    160,
                ),
            }
        )
        specs.append(
            {
                "section": section.title,
                "section_id": section_id,
                "cluster_id": cluster_id,
                "role": SlideRole.RECOMMENDATION if section_visual in {VisualType.COMPARISON, VisualType.PROCESS} else SlideRole.COMPARISON,
                "title": _lecture_copy(
                    section.title,
                    child_titles,
                    korean="{section}에서 뽑아야 할 선택 규칙",
                    english="Selection rule from {section}",
                ),
                "takeaway": _lecture_copy(
                    section.title,
                    child_titles,
                    korean="{section} cluster는 청중이 끝까지 가져가야 할 결정 규칙으로 마무리한다.",
                    english="End the cluster with the decision rule the audience should keep from {section}.",
                ),
                "message": concept_bullets[-1],
                "visual": VisualType.COMPARISON if section_visual != VisualType.PROCESS else VisualType.PROCESS,
                "core_content": _compressed_bullets(child_titles + base_lines + child_lines, fallback_title=section.title),
                "required_assets": [section.title, "Selection rule"],
                "content_tier": ContentTier.LECTURE_CORE,
                "presenter_notes": _shorten_line(
                    "다음 섹션으로 이어질 판단 기준을 명확히 남기며 cluster를 마무리한다."
                    if section_language == "ko"
                    else "Close the cluster by stating what should carry into the next section.",
                    160,
                ),
            }
        )
        clustering_decisions.append(
            f"Clustered {section.number} {section.title} and {max(len(children), 1)} source subsection(s) into four lecture slides."
        )

    closing_language = "ko" if any(_contains_hangul(section.title) for section in level2_blocks[:2]) else "en"
    if closing_language == "ko":
        specs.extend(
            [
                {
                    "section": "Closing",
                    "role": SlideRole.RECOMMENDATION,
                    "title": "문제 구조가 방법 선택을 이끈다",
                    "takeaway": _shorten_line("알고리즘 선택은 방법 이름 목록이 아니라 문제 구조 읽기에서 시작한다.", 120),
                    "message": _shorten_line("변수 유형, 미분 가능성, 제약식, 전역성 요구를 먼저 읽고 방법을 고른다.", 140),
                    "visual": VisualType.COMPARISON,
                    "core_content": _compressed_bullets(
                        [
                            "구조가 맞으면 gradient 계열을 쓴다.",
                            "블랙박스나 비매끈 문제는 search 중심 방법을 쓴다.",
                            "전역성 요구와 계산 비용을 함께 본다.",
                        ],
                        fallback_title="문제 구조가 방법 선택을 이끈다",
                    ),
                    "required_assets": ["방법 선택 규칙"],
                    "content_tier": ContentTier.LECTURE_CORE,
                },
                {
                    "section": "Closing",
                    "role": SlideRole.ANALYSIS,
                    "title": "강의 마무리",
                    "takeaway": _shorten_line("정식화, 해석, 선택이 이번 강의의 세 가지 핵심 산출물이다.", 120),
                    "message": _shorten_line("좋은 최적화 실무는 문제를 읽고 방법을 정당화할 수 있을 때 시작된다.", 140),
                    "visual": VisualType.FRAMEWORK,
                    "core_content": _compressed_bullets(
                        ["문제를 정확히 정식화한다.", "최적성 조건을 해석한다.", "구조에 맞는 방법을 선택한다."],
                        fallback_title="강의 마무리",
                    ),
                    "required_assets": ["강의 마무리"],
                    "content_tier": ContentTier.LECTURE_CORE,
                },
            ]
        )
    else:
        specs.extend(
            [
                {
                    "section": "Closing",
                    "role": SlideRole.RECOMMENDATION,
                    "title": "Problem structure drives method choice",
                    "takeaway": _shorten_line("Algorithm choice starts with structure, not with a list of method names.", 120),
                    "message": _shorten_line("Read variable type, differentiability, constraints, and globality requirements before choosing the method.", 140),
                    "visual": VisualType.COMPARISON,
                    "core_content": _compressed_bullets(
                        ["Use gradient methods when the structure supports them.", "Use search-heavy methods when the problem is black-box or non-smooth.", "Balance globality against compute cost."],
                        fallback_title="Problem structure drives method choice",
                    ),
                    "required_assets": ["Method-selection rule"],
                    "content_tier": ContentTier.LECTURE_CORE,
                },
                {
                    "section": "Closing",
                    "role": SlideRole.ANALYSIS,
                    "title": "Lecture close",
                    "takeaway": _shorten_line("Formulation, interpretation, and selection are the three durable outcomes of the lecture.", 120),
                    "message": _shorten_line("Strong optimization practice starts with reading the problem well enough to justify the method.", 140),
                    "visual": VisualType.FRAMEWORK,
                    "core_content": _compressed_bullets(
                        ["Formulate the problem precisely.", "Interpret the optimality conditions.", "Choose the method that matches the structure."],
                        fallback_title="Lecture close",
                    ),
                    "required_assets": ["Lecture close"],
                    "content_tier": ContentTier.LECTURE_CORE,
                },
            ]
        )

    appendix_specs, appendix_decisions = _lecture_appendix_specs(level2_blocks, children_by_parent, pdf_label, pdf_ref_path, appendix_target)
    clustering_decisions.extend(appendix_decisions)
    main_specs = [spec for spec in specs if spec.get("deck_mode", DeckMode.MAIN_STORY) != DeckMode.APPENDIX]
    if len(main_specs) > main_story_target:
        overflow = len(main_specs) - main_story_target
        trimmed: list[dict[str, Any]] = []
        for spec in main_specs:
            if overflow and spec.get("content_tier") == ContentTier.SUPPORTING_EXAMPLE:
                overflow -= 1
                clustering_decisions.append(f"Compressed supporting-example slide `{spec['title']}` to stay within the main-story budget.")
                appendix_specs.insert(0, {**spec, "section": "Appendix", "deck_mode": DeckMode.APPENDIX, "content_tier": ContentTier.APPENDIX_ONLY, "role": SlideRole.APPENDIX_EVIDENCE})
                continue
            trimmed.append(spec)
        main_specs = trimmed[:main_story_target]
    elif len(main_specs) < workflow_plan.main_story_slide_count_range.start:
        clustering_decisions.append("Document outline produced fewer clusters than the lecture budget floor; kept the compressed main story rather than emitting filler slides.")

    return [*main_specs, *appendix_specs[:appendix_target]], clustering_decisions


def _report_decision_specs(main_story_target: int) -> list[dict[str, Any]]:
    specs = [
        {"section": "Executive Summary", "role": SlideRole.EXECUTIVE_SUMMARY, "title": "Decision in one line", "takeaway": "Lead with the decision path before the supporting detail.", "visual": VisualType.TEXT},
        {"section": "Executive Summary", "role": SlideRole.RECOMMENDATION, "title": "Why this path now", "takeaway": "The current evidence supports acting now rather than waiting for more data.", "visual": VisualType.COMPARISON},
        {"section": "Why Now", "role": SlideRole.ANALYSIS, "title": "Performance pressure is concentrated", "takeaway": "The performance gap is concentrated enough to justify a focused response.", "visual": VisualType.CHART},
        {"section": "Why Now", "role": SlideRole.EVIDENCE, "title": "Local evidence already shows the gap", "takeaway": "Existing local material is strong enough to anchor the argument.", "visual": VisualType.DOCUMENT_CROP},
        {"section": "Proof And Options", "role": SlideRole.EVIDENCE, "title": "The strongest evidence belongs in the main story", "takeaway": "One carefully chosen proof object should carry the burden of evidence.", "visual": VisualType.DOCUMENT_CROP},
        {"section": "Proof And Options", "role": SlideRole.COMPARISON, "title": "Options narrow to one practical route", "takeaway": "A direct comparison makes the preferred route hard to miss.", "visual": VisualType.TABLE},
        {"section": "Proof And Options", "role": SlideRole.ANALYSIS, "title": "Tradeoffs stay manageable", "takeaway": "The key tradeoffs are bounded and can be managed transparently.", "visual": VisualType.CHART},
        {"section": "Recommendation And Next Steps", "role": SlideRole.PROCESS, "title": "Next steps stay bounded", "takeaway": "Execution should be framed as a short, controlled sequence.", "visual": VisualType.TIMELINE},
        {"section": "Recommendation And Next Steps", "role": SlideRole.RECOMMENDATION, "title": "Close with the commitment required", "takeaway": "End with the exact approval or operating shift required from the audience.", "visual": VisualType.TEXT},
    ]
    return specs[:main_story_target]


def _explainer_persuasion_specs(main_story_target: int) -> list[dict[str, Any]]:
    specs = [
        {"section": "Problem Framing", "role": SlideRole.EXECUTIVE_SUMMARY, "title": "The problem is concrete and current", "takeaway": "Anchor the audience in one clear problem before selling the answer.", "visual": VisualType.TEXT},
        {"section": "Problem Framing", "role": SlideRole.ANALYSIS, "title": "Why the current state is insufficient", "takeaway": "The current approach fails in a way the audience can recognize quickly.", "visual": VisualType.COMPARISON},
        {"section": "Mechanism And Proof", "role": SlideRole.INFOGRAPHIC, "title": "The mechanism is easy to understand", "takeaway": "Explain the mechanism with one structured visual rather than paragraphs.", "visual": VisualType.FRAMEWORK},
        {"section": "Mechanism And Proof", "role": SlideRole.EVIDENCE, "title": "One proof object does most of the work", "takeaway": "Use one proof object that makes the recommendation credible.", "visual": VisualType.DOCUMENT_CROP},
        {"section": "Recommendation", "role": SlideRole.RECOMMENDATION, "title": "The ask follows naturally", "takeaway": "Turn explanation into a single audience ask without adding a second narrative.", "visual": VisualType.TEXT},
        {"section": "Recommendation", "role": SlideRole.PROCESS, "title": "Next step sequence stays simple", "takeaway": "Show a short sequence for what happens after approval.", "visual": VisualType.PROCESS},
    ]
    return specs[:main_story_target]


def _training_demo_specs(main_story_target: int) -> list[dict[str, Any]]:
    specs = [
        {"section": "Orientation", "role": SlideRole.EXECUTIVE_SUMMARY, "title": "The learning objective is explicit", "takeaway": "Start by defining what the audience must understand and do by the end.", "visual": VisualType.TEXT},
        {"section": "Orientation", "role": SlideRole.ANALYSIS, "title": "Why the workflow changes", "takeaway": "Frame the operational shift before showing screens or steps.", "visual": VisualType.PROCESS},
        {"section": "Workflow Model", "role": SlideRole.INFOGRAPHIC, "title": "The target workflow fits on one slide", "takeaway": "The target workflow should read as one compact end-to-end path.", "visual": VisualType.FRAMEWORK},
        {"section": "Workflow Model", "role": SlideRole.EVIDENCE, "title": "Step one has a concrete screen or crop", "takeaway": "Use one screen reference to ground the first operational step.", "visual": VisualType.DOCUMENT_CROP},
        {"section": "Workflow Model", "role": SlideRole.EVIDENCE, "title": "Step two keeps the same framing", "takeaway": "Repeat the same screenshot framing to preserve continuity.", "visual": VisualType.DOCUMENT_CROP},
        {"section": "Workflow Model", "role": SlideRole.PROCESS, "title": "Control points stay visible", "takeaway": "Show guardrails and decisions inside the process rather than in speaker notes only.", "visual": VisualType.PROCESS},
        {"section": "Workflow Model", "role": SlideRole.ANALYSIS, "title": "The exception path is bounded", "takeaway": "Show what happens when the standard path breaks.", "visual": VisualType.TABLE},
        {"section": "Demo Sequence", "role": SlideRole.SECTION_DIVIDER, "title": "Demo sequence", "takeaway": "Reset the audience before the live demo sequence begins.", "visual": VisualType.TEXT},
        {"section": "Demo Sequence", "role": SlideRole.EXECUTIVE_SUMMARY, "title": "Demo path at a glance", "takeaway": "Summarize the demo path before going step by step.", "visual": VisualType.PROCESS},
        {"section": "Demo Sequence", "role": SlideRole.EVIDENCE, "title": "Demo screen one", "takeaway": "Use a cleaned screen crop rather than a full screenshot dump.", "visual": VisualType.DOCUMENT_CROP},
        {"section": "Demo Sequence", "role": SlideRole.EVIDENCE, "title": "Demo screen two", "takeaway": "Keep annotations consistent across demo slides.", "visual": VisualType.DOCUMENT_CROP},
        {"section": "Demo Sequence", "role": SlideRole.EVIDENCE, "title": "Demo screen three", "takeaway": "Close the live flow with the highest-confidence screen state.", "visual": VisualType.DOCUMENT_CROP},
        {"section": "Rollout And Support", "role": SlideRole.COMPARISON, "title": "Roles and responsibilities stay explicit", "takeaway": "Clarify ownership before rollout so training does not drift into ambiguity.", "visual": VisualType.TABLE},
        {"section": "Rollout And Support", "role": SlideRole.PROCESS, "title": "Rollout pacing is staged", "takeaway": "Show the rollout sequence in one bounded timeline.", "visual": VisualType.TIMELINE},
        {"section": "Rollout And Support", "role": SlideRole.ANALYSIS, "title": "Support channels stay clear", "takeaway": "Clarify how questions and exceptions will be handled after training.", "visual": VisualType.FRAMEWORK},
        {"section": "Rollout And Support", "role": SlideRole.RECOMMENDATION, "title": "Close with the operating commitment", "takeaway": "End with the operating behaviors that must now become standard.", "visual": VisualType.TEXT},
    ]
    return specs[:main_story_target]


def _tight_main_story_specs(main_story_target: int) -> list[dict[str, Any]]:
    specs = [
        {
            "section": "Core Story",
            "role": SlideRole.EXECUTIVE_SUMMARY,
            "title": "State the thesis immediately",
            "takeaway": "Lead with the decision frame and let the rest of the deck support it.",
            "visual": VisualType.TEXT,
            "layout_pattern_id": "summary",
        },
        {
            "section": "Core Story",
            "role": SlideRole.ANALYSIS,
            "title": "Frame the current state",
            "takeaway": "Show the current state in one analytical move without turning the slide into a tradeoff matrix.",
            "visual": VisualType.COMPARISON,
            "layout_pattern_id": "summary",
        },
        {
            "section": "Core Story",
            "role": SlideRole.EVIDENCE,
            "title": "Use one dominant proof object",
            "takeaway": "A single strong proof object is better than several weak fragments.",
            "visual": VisualType.DOCUMENT_CROP,
            "layout_pattern_id": "worked-example",
        },
        {
            "section": "Core Story",
            "role": SlideRole.RECOMMENDATION,
            "title": "Close on the next step",
            "takeaway": "End with one clear action or conclusion.",
            "visual": VisualType.PROCESS,
            "layout_pattern_id": "summary",
        },
    ]
    return specs[:main_story_target]


def _thesis_proof_close_proof_specs(proof_target: int) -> list[dict[str, Any]]:
    base_specs = [
        {
            "section": "Proof Spine",
            "role": SlideRole.EVIDENCE,
            "title": "Lead with the strongest proof of the thesis",
            "takeaway": "Open the proof spine with the strongest direct evidence so the thesis earns credibility immediately.",
            "visual": VisualType.DOCUMENT_CROP,
            "layout_pattern_id": "worked-example",
        },
        {
            "section": "Proof Spine",
            "role": SlideRole.COMPARISON,
            "title": "A comparison slide shows why this thesis wins",
            "takeaway": "Use one comparison move to show the thesis is stronger than the obvious alternative story.",
            "visual": VisualType.TABLE,
            "layout_pattern_id": "comparison",
        },
        {
            "section": "Proof Spine",
            "role": SlideRole.ANALYSIS,
            "title": "The proof spine converges before the ask",
            "takeaway": "Summarize what the combined proof means before the audience reaches the closing ask.",
            "visual": VisualType.FRAMEWORK,
            "layout_pattern_id": "concept-explainer",
        },
    ]
    proof_specs = list(base_specs[:proof_target])
    while len(proof_specs) < proof_target:
        proof_number = len(proof_specs) + 1
        proof_specs.append(
            {
                "section": "Proof Spine",
                "role": SlideRole.ANALYSIS if proof_number % 2 else SlideRole.COMPARISON,
                "title": f"Proof spine unit {proof_number} narrows the close",
                "takeaway": "Additional proof units should sharpen the closing ask rather than reopen the thesis.",
                "visual": VisualType.FRAMEWORK if proof_number % 2 else VisualType.TABLE,
                "layout_pattern_id": "concept-explainer" if proof_number % 2 else "comparison",
            }
        )
    return proof_specs


def _thesis_proof_close_specs(main_story_target: int) -> list[dict[str, Any]]:
    proof_target = max(2, main_story_target - 2)
    return [
        {
            "section": "Thesis",
            "role": SlideRole.EXECUTIVE_SUMMARY,
            "title": "State the thesis before the proof begins",
            "takeaway": "Open on the one thesis the proof spine must substantiate.",
            "visual": VisualType.TEXT,
            "layout_pattern_id": "summary",
        },
        *_thesis_proof_close_proof_specs(proof_target),
        {
            "section": "Close",
            "role": SlideRole.RECOMMENDATION,
            "title": "Close with the ask the proof now supports",
            "takeaway": "Turn the thesis and proof spine into one explicit financing or decision ask.",
            "visual": VisualType.PROCESS,
            "layout_pattern_id": "summary",
        },
    ]


def _evidence_backed_core_proof_specs(proof_target: int) -> list[dict[str, Any]]:
    base_specs = [
        {
            "section": "Proof Modules",
            "role": SlideRole.EVIDENCE,
            "title": "Proof module one anchors the claim",
            "takeaway": "Lead with the strongest proof object so the audience can calibrate confidence quickly.",
            "visual": VisualType.DOCUMENT_CROP,
            "layout_pattern_id": "worked-example",
        },
        {
            "section": "Proof Modules",
            "role": SlideRole.COMPARISON,
            "title": "A second proof path confirms the pattern",
            "takeaway": "Use one structured comparison to show the first proof object is not isolated.",
            "visual": VisualType.TABLE,
            "layout_pattern_id": "comparison",
        },
        {
            "section": "Proof Modules",
            "role": SlideRole.EVIDENCE,
            "title": "The metrics reinforce the same conclusion",
            "takeaway": "A compact quantitative proof should agree with the source-led proof module.",
            "visual": VisualType.CHART,
            "layout_pattern_id": "worked-example",
        },
        {
            "section": "Proof Modules",
            "role": SlideRole.ANALYSIS,
            "title": "Proof modules converge on one interpretation",
            "takeaway": "State the interpretation only after the proof stack is visible on the page.",
            "visual": VisualType.FRAMEWORK,
            "layout_pattern_id": "concept-explainer",
        },
    ]
    proof_specs = list(base_specs[:proof_target])
    while len(proof_specs) < proof_target:
        module_number = len(proof_specs) + 1
        if module_number % 2 == 0:
            proof_specs.append(
                {
                    "section": "Proof Modules",
                    "role": SlideRole.COMPARISON,
                    "title": f"Proof module {module_number} keeps tradeoffs visible",
                    "takeaway": "Keep additional proof modules comparative so the main story still reads as one bounded stack.",
                    "visual": VisualType.TABLE,
                    "layout_pattern_id": "comparison",
                }
            )
        else:
            proof_specs.append(
                {
                    "section": "Proof Modules",
                    "role": SlideRole.EVIDENCE,
                    "title": f"Proof module {module_number} adds one more confidence check",
                    "takeaway": "Add another proof object only when it materially changes confidence in the core claim.",
                    "visual": VisualType.DOCUMENT_CROP,
                    "layout_pattern_id": "worked-example",
                }
            )
    return proof_specs


def _evidence_backed_core_specs(workflow_plan: WorkflowPlan, main_story_target: int) -> list[dict[str, Any]]:
    has_recommendation = bool(workflow_plan.recommendations or workflow_plan.project_snapshot.recommendations)
    if main_story_target <= 4:
        claim_specs = [
            {
                "section": "Core Claim",
                "role": SlideRole.EXECUTIVE_SUMMARY,
                "title": "State the core claim immediately",
                "takeaway": "Open with the one claim the proof stack must substantiate.",
                "visual": VisualType.TEXT,
                "layout_pattern_id": "summary",
            },
        ]
        proof_target = max(2, main_story_target - 2)
    else:
        claim_specs = [
            {
                "section": "Core Claim",
                "role": SlideRole.EXECUTIVE_SUMMARY,
                "title": "State the core claim immediately",
                "takeaway": "Open with the one claim the proof stack must substantiate.",
                "visual": VisualType.TEXT,
                "layout_pattern_id": "summary",
            },
            {
                "section": "Core Claim",
                "role": SlideRole.ANALYSIS,
                "title": "Frame how the proof stack should be read",
                "takeaway": "Explain the evaluation lens before the audience starts reading individual proof modules.",
                "visual": VisualType.CHART,
                "layout_pattern_id": "worked-example",
            },
        ]
        proof_target = max(2, main_story_target - 3)

    closing_spec = {
        "section": "Implications",
        "role": SlideRole.RECOMMENDATION if has_recommendation else SlideRole.ANALYSIS,
        "title": "What the proof set means now" if not has_recommendation else "What the proof set means for the next step",
        "takeaway": (
            "Close by stating the implication the audience should retain from the proof modules."
            if not has_recommendation
            else "Close by turning the proof stack into one bounded recommended move."
        ),
        "visual": VisualType.TEXT if not has_recommendation else VisualType.PROCESS,
        "layout_pattern_id": "summary",
    }
    proof_specs = _evidence_backed_core_proof_specs(proof_target)
    return [*claim_specs, *proof_specs, closing_spec][:main_story_target]


def _generic_specs(main_story_target: int) -> list[dict[str, Any]]:
    specs = [
        {"section": "Core Story", "role": SlideRole.EXECUTIVE_SUMMARY, "title": "State the thesis immediately", "takeaway": "Lead with the thesis and let the rest of the deck support it.", "visual": VisualType.TEXT},
        {"section": "Core Story", "role": SlideRole.ANALYSIS, "title": "Frame the current state", "takeaway": "Show the current state in one analytical move.", "visual": VisualType.COMPARISON},
        {"section": "Core Story", "role": SlideRole.EVIDENCE, "title": "Use one dominant proof object", "takeaway": "A single strong proof object is better than several weak fragments.", "visual": VisualType.DOCUMENT_CROP},
        {"section": "Core Story", "role": SlideRole.RECOMMENDATION, "title": "Close on the next step", "takeaway": "End with one clear action or conclusion.", "visual": VisualType.PROCESS},
    ]
    return specs[:main_story_target]


def _appendix_specs(appendix_target: int) -> list[dict[str, Any]]:
    specs = [
        {"section": "Appendix", "role": SlideRole.APPENDIX_EVIDENCE, "title": "Methods and assumptions", "takeaway": "Methods belong outside the main story unless they change the decision.", "visual": VisualType.TABLE},
        {"section": "Appendix", "role": SlideRole.APPENDIX_EVIDENCE, "title": "Extended evidence", "takeaway": "Store backup evidence in appendix candidates rather than the core narrative.", "visual": VisualType.DOCUMENT_CROP},
        {"section": "Appendix", "role": SlideRole.APPENDIX_EVIDENCE, "title": "Detailed comparisons", "takeaway": "Detailed comparisons should support the main story, not interrupt it.", "visual": VisualType.TABLE},
        {"section": "Appendix", "role": SlideRole.REFERENCES, "title": "References and source map", "takeaway": "Keep a clear source map for later production and QA.", "visual": VisualType.TEXT},
    ]
    return specs[:appendix_target]


def _base_slide_specs(workflow_plan: WorkflowPlan, main_story_target: int, appendix_target: int) -> list[dict[str, Any]]:
    types = {workflow_plan.presentation_type_diagnosis.primary_type, *workflow_plan.presentation_type_diagnosis.secondary_types}
    if workflow_plan.workflow_option == "tight-main-story":
        main_specs = _tight_main_story_specs(main_story_target)
    elif workflow_plan.workflow_option == "evidence-backed-core":
        main_specs = _evidence_backed_core_specs(workflow_plan, main_story_target)
    elif workflow_plan.workflow_option == "thesis-proof-close":
        main_specs = _thesis_proof_close_specs(main_story_target)
    elif PresentationType.TRAINING in types or PresentationType.DEMO in types:
        main_specs = _training_demo_specs(main_story_target)
    elif PresentationType.REPORT in types and PresentationType.DECISION in types:
        main_specs = _report_decision_specs(main_story_target)
    elif PresentationType.EXPLAINER in types and PresentationType.PERSUASION in types:
        main_specs = _explainer_persuasion_specs(main_story_target)
    else:
        main_specs = _generic_specs(main_story_target)
    return main_specs + _appendix_specs(appendix_target)


def _preferred_visual_type(
    requested_visual: VisualType,
    grouped_materials: dict[BriefMaterialType, list[ProjectMaterial]],
) -> VisualType:
    if requested_visual == VisualType.DOCUMENT_CROP:
        if grouped_materials.get(BriefMaterialType.DOCUMENT) or grouped_materials.get(BriefMaterialType.IMAGE) or grouped_materials.get(BriefMaterialType.DECK):
            return VisualType.DOCUMENT_CROP
        return VisualType.CHART
    return requested_visual


def _production_bridge_for_visual(
    visual_type: VisualType,
    grouped_materials: dict[BriefMaterialType, list[ProjectMaterial]],
    fallback_refs: list[SourceMaterialRef],
    slide_title: str,
) -> ProductionBridge:
    if visual_type == VisualType.DOCUMENT_CROP:
        refs = _select_source_refs(
            grouped_materials,
            (BriefMaterialType.DOCUMENT, BriefMaterialType.IMAGE, BriefMaterialType.DECK),
            fallback_refs,
        )
        return ProductionBridge(
            visual_source_preference=VisualSourcePreference.DOCUMENT_CROP,
            source_material_refs=refs,
            crop_subject_hint=slide_title,
            fallback_visual=VisualType.CHART,
            production_mode=ProductionMode.SOURCE_REUSE,
        )
    if visual_type in {VisualType.CHART, VisualType.TABLE, VisualType.COMPARISON}:
        refs = _select_source_refs(
            grouped_materials,
            (BriefMaterialType.SPREADSHEET, BriefMaterialType.DATA, BriefMaterialType.DOCUMENT),
            fallback_refs,
        )
        return ProductionBridge(
            visual_source_preference=VisualSourcePreference.STRUCTURED_VISUAL,
            source_material_refs=refs,
            crop_subject_hint=None,
            fallback_visual=VisualType.TABLE if visual_type == VisualType.CHART else VisualType.CHART,
            production_mode=ProductionMode.STRUCTURED_VISUAL,
        )
    if visual_type in {VisualType.TIMELINE, VisualType.FRAMEWORK, VisualType.PROCESS, VisualType.INFOGRAPHIC}:
        refs = _select_source_refs(
            grouped_materials,
            (BriefMaterialType.NOTES, BriefMaterialType.DOCUMENT, BriefMaterialType.DECK),
            fallback_refs,
        )
        return ProductionBridge(
            visual_source_preference=VisualSourcePreference.STRUCTURED_VISUAL,
            source_material_refs=refs,
            crop_subject_hint=None,
            fallback_visual=VisualType.TABLE,
            production_mode=ProductionMode.STRUCTURED_VISUAL,
        )
    return ProductionBridge(
        visual_source_preference=VisualSourcePreference.STRUCTURED_VISUAL,
        source_material_refs=[],
        crop_subject_hint=None,
        fallback_visual=VisualType.COMPARISON,
        production_mode=ProductionMode.STRUCTURED_VISUAL,
    )


def _required_evidence_assets(visual_type: VisualType, section: str, grouped_materials: dict[BriefMaterialType, list[ProjectMaterial]]) -> list[str]:
    if visual_type == VisualType.DOCUMENT_CROP:
        if grouped_materials.get(BriefMaterialType.IMAGE) or grouped_materials.get(BriefMaterialType.DECK):
            return [f"Clean screenshot crop for {section.lower()}", "Source provenance note"]
        return [f"Document crop for {section.lower()}", "Source provenance note"]
    if visual_type in {VisualType.CHART, VisualType.TABLE, VisualType.COMPARISON}:
        return [f"Structured data values for {section.lower()}", "Checked labels and units"]
    if visual_type in {VisualType.TIMELINE, VisualType.FRAMEWORK, VisualType.PROCESS, VisualType.INFOGRAPHIC}:
        return [f"Step labels or framework nodes for {section.lower()}", "Ordered narrative labels"]
    if visual_type == VisualType.TEXT:
        return [f"Approved thesis wording for {section.lower()}"]
    return [f"Approved content support for {section.lower()}"]


def _layout_pattern_id(role: SlideRole, visual_type: VisualType) -> str:
    if role == SlideRole.TITLE:
        return "cover"
    if role == SlideRole.SECTION_DIVIDER:
        return "section-divider"
    if role in {SlideRole.REFERENCES, SlideRole.APPENDIX_EVIDENCE}:
        return "appendix-reference"
    if role in {SlideRole.RECOMMENDATION, SlideRole.EXECUTIVE_SUMMARY}:
        return "summary"
    if visual_type in {VisualType.TIMELINE, VisualType.PROCESS, VisualType.DECISION_PATH}:
        return "process-flow"
    if visual_type in {VisualType.TABLE, VisualType.COMPARISON}:
        return "comparison"
    if visual_type in {VisualType.DOCUMENT_CROP, VisualType.CHART, VisualType.PHOTO}:
        return "worked-example"
    if visual_type in {VisualType.FRAMEWORK, VisualType.INFOGRAPHIC, VisualType.HIERARCHY, VisualType.METRIC_SUMMARY}:
        return "concept-explainer"
    return "definition-theorem"


def _slide_id(slide_number: int) -> str:
    return f"s{slide_number:03d}"


def _dedupe_visual_types(*visual_types: VisualType | None) -> list[VisualType]:
    ordered: list[VisualType] = []
    seen: set[VisualType] = set()
    for visual_type in visual_types:
        if visual_type is None or visual_type in seen:
            continue
        ordered.append(visual_type)
        seen.add(visual_type)
    return ordered


def _fallback_ladder_for_slide(slide: BlueprintSlide) -> list[VisualType]:
    if slide.production_bridge.visual_source_preference == VisualSourcePreference.DOCUMENT_CROP or slide.visual_type == VisualType.DOCUMENT_CROP:
        return _dedupe_visual_types(
            slide.visual_type,
            VisualType.CHART,
            VisualType.INFOGRAPHIC,
            VisualType.COMPARISON,
        )
    if slide.visual_type == VisualType.CHART:
        return _dedupe_visual_types(
            slide.visual_type,
            VisualType.TABLE,
            VisualType.INFOGRAPHIC,
            VisualType.COMPARISON,
        )
    if slide.visual_type == VisualType.TABLE:
        return _dedupe_visual_types(
            slide.visual_type,
            VisualType.CHART,
            VisualType.INFOGRAPHIC,
            VisualType.COMPARISON,
        )
    if slide.visual_type == VisualType.COMPARISON:
        return _dedupe_visual_types(
            slide.visual_type,
            VisualType.TABLE,
            VisualType.INFOGRAPHIC,
            VisualType.TEXT,
        )
    if slide.visual_type in {VisualType.TIMELINE, VisualType.FRAMEWORK, VisualType.PROCESS}:
        return _dedupe_visual_types(
            slide.visual_type,
            VisualType.TABLE,
            VisualType.INFOGRAPHIC,
            VisualType.COMPARISON,
        )
    if slide.visual_type == VisualType.INFOGRAPHIC:
        return _dedupe_visual_types(
            slide.visual_type,
            VisualType.TABLE,
            VisualType.COMPARISON,
            VisualType.TEXT,
        )
    if slide.visual_type == VisualType.PHOTO:
        return _dedupe_visual_types(
            slide.visual_type,
            VisualType.DOCUMENT_CROP,
            VisualType.COMPARISON,
            VisualType.TEXT,
        )
    if slide.visual_type == VisualType.QUOTE:
        return _dedupe_visual_types(
            slide.visual_type,
            VisualType.TEXT,
            VisualType.COMPARISON,
        )
    return _dedupe_visual_types(slide.visual_type, VisualType.COMPARISON)


def _asset_kind_for_slide(slide: BlueprintSlide) -> AssetKind:
    if slide.production_bridge.visual_source_preference == VisualSourcePreference.DOCUMENT_CROP or slide.visual_type == VisualType.DOCUMENT_CROP:
        return AssetKind.DOCUMENT_CROP
    if slide.production_bridge.visual_source_preference == VisualSourcePreference.EXISTING_ASSET or slide.visual_type == VisualType.PHOTO:
        return AssetKind.IMAGE
    return AssetKind.STRUCTURED_VISUAL


def _asset_priority_for_slide(slide: BlueprintSlide) -> AssetPriority:
    if slide.deck_mode == DeckMode.APPENDIX or slide.slide_role in {SlideRole.APPENDIX_EVIDENCE, SlideRole.REFERENCES}:
        return AssetPriority.LOW
    if slide.slide_role in {SlideRole.EXECUTIVE_SUMMARY, SlideRole.RECOMMENDATION}:
        return AssetPriority.CRITICAL
    if slide.slide_role in {SlideRole.EVIDENCE, SlideRole.ANALYSIS, SlideRole.COMPARISON}:
        return AssetPriority.HIGH
    return AssetPriority.NORMAL


def _asset_quality_requirements_for_slide(slide: BlueprintSlide) -> list[str]:
    requirements = [
        "Keep the one-line takeaway visually dominant.",
        "Use only approved local inputs and design tokens for this deck.",
    ]
    if slide.production_bridge.visual_source_preference == VisualSourcePreference.DOCUMENT_CROP:
        requirements.extend(
            [
                "Crop tightly around the subject hint and preserve presentation-distance readability.",
                "Exclude captions or footnotes unless they are required to support the slide message.",
                "Record provenance for the selected source crop.",
            ]
        )
    elif slide.visual_type in {VisualType.CHART, VisualType.TABLE, VisualType.COMPARISON}:
        requirements.extend(
            [
                "Verify labels, units, and values against the preferred source document.",
                "Prefer the simplest chart or table that supports the slide message.",
                "Maintain legibility from the back of the room.",
            ]
        )
    elif slide.visual_type in {VisualType.TIMELINE, VisualType.FRAMEWORK, VisualType.PROCESS, VisualType.INFOGRAPHIC}:
        requirements.extend(
            [
                "Preserve three-second comprehension by limiting node and label density.",
                "Use the approved spacing and hierarchy system for slide-native visuals.",
                "Fallback to a simpler table or comparison if the structure becomes too dense.",
            ]
        )
    else:
        requirements.extend(
            [
                "Keep support assets simpler than the headline message.",
                "Avoid decorative content that competes with the takeaway.",
            ]
        )
    if slide.deck_mode == DeckMode.APPENDIX:
        requirements.append("Keep appendix assets lighter-weight than main-story visuals unless evidence fidelity requires more detail.")
    return requirements


def _request_brief_for_slide(slide: BlueprintSlide) -> str:
    visual_label = slide.visual_type.value.replace("-", " ")
    if slide.production_bridge.visual_source_preference == VisualSourcePreference.DOCUMENT_CROP:
        return f"Create the source-derived crop for slide {_slide_id(slide.slide_number)} to support: {slide.main_message}"
    if slide.visual_type in {VisualType.CHART, VisualType.TABLE, VisualType.COMPARISON}:
        return f"Prepare the {visual_label} package for slide {_slide_id(slide.slide_number)} to support: {slide.main_message}"
    if slide.visual_type in {VisualType.TIMELINE, VisualType.FRAMEWORK, VisualType.PROCESS, VisualType.INFOGRAPHIC}:
        return f"Prepare the slide-native {visual_label} for slide {_slide_id(slide.slide_number)} to support: {slide.main_message}"
    return f"Prepare the simplest production-safe support for slide {_slide_id(slide.slide_number)} to support: {slide.main_message}"


def _should_request_asset(slide: BlueprintSlide) -> bool:
    if slide.slide_role in {SlideRole.TITLE, SlideRole.SECTION_DIVIDER, SlideRole.REFERENCES}:
        return slide.visual_type != VisualType.TEXT or bool(slide.production_bridge.source_material_refs)
    return True


def derive_asset_requests_from_blueprint(slides: list[BlueprintSlide]) -> list[AssetRequest]:
    requests: list[AssetRequest] = []
    for slide in slides:
        if not _should_request_asset(slide):
            continue
        source_ref = slide.production_bridge.source_material_refs[0] if slide.production_bridge.source_material_refs else None
        fallback_ladder = _fallback_ladder_for_slide(slide)
        fallback_visual = slide.production_bridge.fallback_visual
        if fallback_visual not in fallback_ladder[1:]:
            fallback_visual = fallback_ladder[1] if len(fallback_ladder) > 1 else None
        requests.append(
            AssetRequest(
                request_id=f"asset-req-{_slide_id(slide.slide_number)}",
                slide_number=slide.slide_number,
                slide_id=_slide_id(slide.slide_number),
                slide_message=slide.main_message,
                asset_kind=_asset_kind_for_slide(slide),
                priority=_asset_priority_for_slide(slide),
                brief=_request_brief_for_slide(slide),
                required_visual_type=slide.visual_type,
                visual_type=slide.visual_type,
                visual_source_preference=slide.production_bridge.visual_source_preference,
                source_material_refs=slide.production_bridge.source_material_refs,
                preferred_source_doc=(source_ref.path or source_ref.label) if source_ref is not None else None,
                page_hint=source_ref.page if source_ref is not None else None,
                crop_subject_hint=slide.production_bridge.crop_subject_hint,
                fallback_visual=fallback_visual,
                fallback_ladder=fallback_ladder,
                approval_status=StageStatus.READY,
                production_mode=slide.production_bridge.production_mode,
                asset_quality_requirements=_asset_quality_requirements_for_slide(slide),
                allowed_crop_review_actions=(
                    [
                        CropReviewAction.ACCEPT,
                        CropReviewAction.REVISE,
                        CropReviewAction.FALLBACK_TO_VISUAL,
                    ]
                    if _asset_kind_for_slide(slide) == AssetKind.DOCUMENT_CROP
                    else []
                ),
            )
        )
    return requests


def _core_content_for_slide(
    spec: dict[str, Any],
    visual_type: VisualType,
    bridge: ProductionBridge,
    required_assets: list[str],
) -> list[str]:
    explicit = spec.get("core_content")
    if explicit:
        return _compressed_bullets(explicit, fallback_title=spec.get("title", "Slide"))

    items: list[str] = [spec.get("takeaway", spec.get("message", spec.get("title", "Slide")))]
    if spec.get("message") and spec["message"] != items[0]:
        items.append(spec["message"])

    if visual_type == VisualType.DOCUMENT_CROP and bridge.source_material_refs:
        items.append(
            _shorten_line(
                "Source focus: " + ", ".join(ref.label for ref in bridge.source_material_refs[:2]),
                LECTURE_MAX_BULLET_CHARS,
            )
        )
    elif required_assets:
        items.extend(required_assets[:2])

    return _compressed_bullets(items, fallback_title=spec.get("title", "Slide"))


def _audience_intent_for_slide(workflow_plan: WorkflowPlan, spec: dict[str, Any], deck_mode: DeckMode) -> str:
    audience = ", ".join(workflow_plan.audience[:2]) or "the audience"
    claim = spec.get("message", spec.get("takeaway", spec.get("title", "Slide")))
    if deck_mode == DeckMode.APPENDIX:
        return _shorten_line(f"Let reviewers audit the support for: {claim}", 180)
    return _shorten_line(f"Help {audience} act on: {claim}", 180)


def _evidence_class_for_slide(
    deck_mode: DeckMode,
    visual_type: VisualType,
    bridge: ProductionBridge,
) -> SlideEvidenceClass:
    if deck_mode == DeckMode.APPENDIX:
        return SlideEvidenceClass.APPENDIX_SUPPORT
    if bridge.source_material_refs or bridge.visual_source_preference in {
        VisualSourcePreference.DOCUMENT_CROP,
        VisualSourcePreference.EXISTING_ASSET,
    }:
        return SlideEvidenceClass.SOURCE_BACKED
    if visual_type in {VisualType.CHART, VisualType.TABLE, VisualType.COMPARISON}:
        return SlideEvidenceClass.DATA_BACKED
    if visual_type in {
        VisualType.FRAMEWORK,
        VisualType.HIERARCHY,
        VisualType.INFOGRAPHIC,
        VisualType.METRIC_SUMMARY,
        VisualType.PROCESS,
        VisualType.TIMELINE,
        VisualType.DECISION_PATH,
    }:
        return SlideEvidenceClass.VISUAL_DEMONSTRATION
    if visual_type == VisualType.TEXT and bridge.fallback_visual is not None:
        return SlideEvidenceClass.STRUCTURED_LOGIC
    return SlideEvidenceClass.MESSAGE_ONLY


def _supporting_evidence_for_slide(
    required_assets: list[str],
    bridge: ProductionBridge,
) -> list[str]:
    refs = [ref.label for ref in bridge.source_material_refs[:3]]
    if refs:
        return _dedupe(refs)
    return _dedupe(required_assets[:3])


def _layout_slot_map_for_slide(
    slide_role: SlideRole,
    visual_type: VisualType,
    deck_mode: DeckMode,
    layout_pattern_id: str,
    supporting_evidence: list[str],
    core_content: list[str],
) -> dict[str, str]:
    slot_map: dict[str, str] = {
        "title": "header",
        "claim": "header",
    }
    if slide_role in {SlideRole.TITLE, SlideRole.SECTION_DIVIDER}:
        slot_map["supporting_text"] = "body-1"
        return slot_map
    if visual_type in {VisualType.TEXT, VisualType.QUOTE}:
        slot_map["supporting_text"] = "body-1"
        if supporting_evidence and deck_mode == DeckMode.APPENDIX:
            slot_map["evidence_marker"] = "body-1"
        return slot_map
    single_body_layout = layout_pattern_id in {"evidence-table", "process-flow"}
    slot_map["primary_visual"] = "body-1"
    if single_body_layout and (core_content or supporting_evidence or deck_mode == DeckMode.APPENDIX):
        slot_map["supporting_text"] = "body-1"
    elif core_content or supporting_evidence or deck_mode == DeckMode.APPENDIX:
        slot_map["supporting_text"] = "body-2"
    if supporting_evidence:
        slot_map["evidence_marker"] = "body-1" if single_body_layout else "body-2"
    return slot_map


def _density_budget_for_slide(
    deck_mode: DeckMode,
    visual_type: VisualType,
    layout_slot_map: dict[str, str],
    core_content: list[str],
    supporting_evidence: list[str],
) -> SlideDensityBudget:
    text_char_ceiling = 720 if deck_mode == DeckMode.APPENDIX else 560
    bullet_count_ceiling = 5 if deck_mode == DeckMode.APPENDIX else 3
    visual_node_ceiling = 8 if deck_mode == DeckMode.APPENDIX else 6
    evidence_item_ceiling = 3 if deck_mode == DeckMode.APPENDIX else 2
    if visual_type in {VisualType.TEXT, VisualType.QUOTE}:
        text_char_ceiling += 80
        bullet_count_ceiling += 1
    if visual_type in {VisualType.TABLE, VisualType.CHART, VisualType.COMPARISON}:
        text_char_ceiling += 100
        visual_node_ceiling += 2
    if visual_type in {VisualType.FRAMEWORK, VisualType.PROCESS, VisualType.TIMELINE, VisualType.INFOGRAPHIC}:
        text_char_ceiling -= 40
    if len(core_content) > bullet_count_ceiling:
        bullet_count_ceiling = len(core_content)
    if len(supporting_evidence) > evidence_item_ceiling:
        evidence_item_ceiling = len(supporting_evidence)
    return SlideDensityBudget(
        text_char_ceiling=text_char_ceiling,
        bullet_count_ceiling=bullet_count_ceiling,
        visual_node_ceiling=visual_node_ceiling,
        evidence_item_ceiling=evidence_item_ceiling,
        layout_slot_count=len(set(layout_slot_map.values())),
    )


def _risk_flags_for_slide(
    deck_mode: DeckMode,
    visual_type: VisualType,
    bridge: ProductionBridge,
    supporting_evidence: list[str],
) -> list[str]:
    flags: list[str] = []
    if deck_mode == DeckMode.APPENDIX:
        flags.append("appendix-boundary")
    if bridge.source_material_refs:
        flags.append("source-marker-required")
    if visual_type in {VisualType.FRAMEWORK, VisualType.HIERARCHY, VisualType.INFOGRAPHIC, VisualType.PROCESS, VisualType.TIMELINE}:
        flags.append("node-density")
    if visual_type in {VisualType.TEXT, VisualType.QUOTE} and deck_mode == DeckMode.MAIN_STORY:
        flags.append("weak-visual-anchor")
    if supporting_evidence and not bridge.source_material_refs:
        flags.append("fallback-evidence-route")
    if bridge.fallback_visual is not None and bridge.fallback_visual != visual_type:
        flags.append("fallback-visual-available")
    return _dedupe(flags)


def _qa_acceptance_hints_for_slide(
    deck_mode: DeckMode,
    visual_type: VisualType,
    evidence_class: SlideEvidenceClass,
    layout_slot_map: dict[str, str],
) -> list[str]:
    hints = [
        "one-claim-only",
        "title-visible",
        "layout-id-supported",
        "text-under-density-budget",
    ]
    if "primary_visual" in layout_slot_map or visual_type not in {VisualType.TEXT, VisualType.QUOTE}:
        hints.append("visual-anchor-present")
    if evidence_class != SlideEvidenceClass.MESSAGE_ONLY:
        hints.append("evidence-marker-present")
    if deck_mode == DeckMode.APPENDIX:
        hints.append("appendix-labeled")
    return _dedupe(hints)


def _verification_flags_for_slide(
    deck_mode: DeckMode,
    visual_type: VisualType,
    bridge: ProductionBridge,
) -> list[str]:
    flags = ["Confirm the title, takeaway, and main message still express one claim."]
    if bridge.source_material_refs:
        flags.append("Verify the cited local sources are still available and mapped correctly.")
    if bridge.visual_source_preference == VisualSourcePreference.DOCUMENT_CROP:
        flags.append("Check crop readability, provenance, and caption dependence before production.")
    if visual_type in {VisualType.CHART, VisualType.TABLE, VisualType.COMPARISON}:
        flags.append("Validate labels, units, and numerical selections against the preferred source.")
    if visual_type in {VisualType.FRAMEWORK, VisualType.PROCESS, VisualType.TIMELINE, VisualType.INFOGRAPHIC}:
        flags.append("Check that node and label density still fits the slide frame at presentation distance.")
    if deck_mode == DeckMode.APPENDIX:
        flags.append("Keep the slide support-only so it does not reopen the main-story recommendation.")
    return _dedupe(flags)


def _required_data_for_visual(visual_type: VisualType, required_assets: list[str]) -> list[str]:
    if visual_type in {VisualType.CHART, VisualType.TABLE, VisualType.COMPARISON}:
        return _dedupe(required_assets + ["Validated values", "Checked labels and units"])
    if visual_type in {VisualType.PROCESS, VisualType.TIMELINE, VisualType.FRAMEWORK, VisualType.INFOGRAPHIC}:
        return _dedupe(required_assets + ["Ordered step or node labels"])
    return list(required_assets)


def _required_screens_or_diagrams(visual_type: VisualType, bridge: ProductionBridge) -> list[str]:
    if visual_type in {VisualType.DOCUMENT_CROP, VisualType.PHOTO}:
        return _dedupe([ref.label for ref in bridge.source_material_refs])
    if visual_type in {VisualType.PROCESS, VisualType.TIMELINE, VisualType.FRAMEWORK, VisualType.INFOGRAPHIC}:
        return ["Approved diagram structure or labeled node list"]
    return []


def _infographic_why_visual(visual_type: VisualType) -> str:
    if visual_type == VisualType.TIMELINE:
        return "A timeline makes order, pacing, and sequence easier to grasp than a bullet list."
    if visual_type == VisualType.PROCESS:
        return "A process visual clarifies ownership, flow, and checkpoints faster than prose."
    if visual_type == VisualType.FRAMEWORK:
        return "A framework visual reveals the structure and relationships better than stacked bullets."
    return "A slide-native infographic helps the audience understand the structure faster than narrative bullets."


def _frame_fit_considerations(visual_type: VisualType) -> list[str]:
    considerations = ["Keep the slide readable in three seconds from presentation distance."]
    if visual_type in {VisualType.FRAMEWORK, VisualType.INFOGRAPHIC}:
        considerations.append("Limit node count so the visual still reads as one structure, not a wall of labels.")
    if visual_type in {VisualType.PROCESS, VisualType.TIMELINE}:
        considerations.append("Use one dominant reading direction and avoid multi-branch complexity.")
    return considerations


def _slide_specs_to_blueprint(
    workflow_plan: WorkflowPlan,
    grouped_materials: dict[BriefMaterialType, list[ProjectMaterial]],
    fallback_refs: list[SourceMaterialRef],
) -> tuple[list[BlueprintSlide], list[EvidencePlanItem], list[InfographicPlanItem], dict[str, Any]]:
    materials = list(sum(grouped_materials.values(), []))
    main_story_target, appendix_target = _target_slide_counts(workflow_plan, materials)
    clustering_decisions: list[str] = []
    lecture_artifacts = plan_concept_driven_lecture(
        deck_title=workflow_plan.deck_title,
        materials=materials,
        main_story_target=main_story_target if _is_lecture_mode(workflow_plan) else 0,
        appendix_target=appendix_target if _is_lecture_mode(workflow_plan) else 0,
    )
    if _is_lecture_mode(workflow_plan):
        if lecture_artifacts.concept_graph.lecture_family == LectureFamily.OPTIMIZATION_METHOD:
            specs, clustering_decisions = _lecture_specs_from_documents(
                workflow_plan,
                materials=materials,
                main_story_target=main_story_target,
                appendix_target=appendix_target,
            )
        else:
            specs = lecture_artifacts.specs
            clustering_decisions = list(lecture_artifacts.clustering_decisions)
        if not specs:
            specs = _base_slide_specs(workflow_plan, main_story_target, appendix_target)
            clustering_decisions.append("Fell back to generic blueprint specs because no concept-driven lecture plan could be synthesized.")
    else:
        specs = _base_slide_specs(workflow_plan, main_story_target, appendix_target)
    authored_specs, authoring_preview, authoring_notes = compile_authoring_layer(
        concept_graph=lecture_artifacts.concept_graph,
        teaching_plan=lecture_artifacts.teaching_plan,
        fallback_specs=specs,
    )
    specs = authored_specs
    clustering_decisions.extend(note for note in authoring_notes if note not in clustering_decisions)
    slides: list[BlueprintSlide] = []
    evidence_plan: list[EvidencePlanItem] = []
    infographic_plan: list[InfographicPlanItem] = []

    for slide_number, spec in enumerate(specs, start=1):
        requested_visual = spec.get("visual", VisualType.TEXT)
        visual_type = _preferred_visual_type(requested_visual, grouped_materials)
        bridge = _production_bridge_for_visual(visual_type, grouped_materials, fallback_refs, spec["title"])
        explicit_refs = list(spec.get("source_material_refs", []))
        if explicit_refs:
            bridge = bridge.model_copy(
                update={
                    "source_material_refs": explicit_refs,
                    "crop_subject_hint": bridge.crop_subject_hint or spec["title"],
                }
            )
        required_assets = list(spec.get("required_assets", _required_evidence_assets(visual_type, spec["section"], grouped_materials)))
        deck_mode = spec.get("deck_mode", DeckMode.APPENDIX if spec["section"] == "Appendix" else DeckMode.MAIN_STORY)
        content_tier = spec.get("content_tier")
        verification_flags = _verification_flags_for_slide(deck_mode, visual_type, bridge)
        main_message = spec.get("message", spec["takeaway"])
        core_content = _core_content_for_slide(spec, visual_type, bridge, required_assets)
        evidence_class = _evidence_class_for_slide(deck_mode, visual_type, bridge)
        supporting_evidence = _supporting_evidence_for_slide(required_assets, bridge)
        layout_pattern_id = spec.get("layout_pattern_id", _layout_pattern_id(spec["role"], visual_type))
        layout_slot_map = _layout_slot_map_for_slide(
            spec["role"],
            visual_type,
            deck_mode,
            layout_pattern_id,
            supporting_evidence,
            core_content,
        )
        density_budget = _density_budget_for_slide(
            deck_mode,
            visual_type,
            layout_slot_map,
            core_content,
            supporting_evidence,
        )
        qa_acceptance_hints = _qa_acceptance_hints_for_slide(
            deck_mode,
            visual_type,
            evidence_class,
            layout_slot_map,
        )
        slide = BlueprintSlide(
            slide_number=slide_number,
            section=spec["section"],
            deck_mode=deck_mode,
            content_tier=content_tier,
            slide_role=spec["role"],
            slide_intent=spec.get("slide_intent"),
            title=spec["title"],
            one_line_takeaway=spec["takeaway"],
            main_message=main_message,
            pedagogical_goal=spec.get("pedagogical_goal"),
            concept_ids=list(spec.get("concept_ids", [])),
            visual_type=visual_type,
            layout_pattern_id=layout_pattern_id,
            production_bridge=bridge,
            required_evidence_assets=required_assets,
            audience_intent=spec.get("audience_intent", _audience_intent_for_slide(workflow_plan, spec, deck_mode)),
            primary_claim=spec.get("primary_claim", main_message),
            supporting_evidence=list(spec.get("supporting_evidence", supporting_evidence)),
            evidence_class=spec.get("evidence_class", evidence_class),
            must_keep_text=list(spec.get("must_keep_text", [spec["title"], spec["takeaway"], main_message])),
            optional_text=list(spec.get("optional_text", [item for item in core_content if item not in {spec["title"], spec["takeaway"], main_message}])),
            layout_slot_map=dict(spec.get("layout_slot_map", layout_slot_map)),
            visual_intent=spec.get("visual_intent", visual_type.value),
            density_budget=spec.get("density_budget", density_budget),
            risk_flags=list(spec.get("risk_flags", _risk_flags_for_slide(deck_mode, visual_type, bridge, supporting_evidence))),
            qa_acceptance_hints=list(spec.get("qa_acceptance_hints", qa_acceptance_hints)),
            slide_archetype=spec.get("slide_archetype"),
            chosen_layout_family=spec.get("chosen_layout_family"),
            primary_visual_structure=spec.get("primary_visual_structure"),
            chrome_blocks_used=list(spec.get("chrome_blocks_used", [])),
            content_budget_summary=dict(spec.get("content_budget_summary", {})),
            duplicate_text_flags=list(spec.get("duplicate_text_flags", [])),
            authoring_payload=dict(spec.get("authoring_payload", {})),
            core_content=core_content,
            verification_flags=verification_flags,
            presenter_notes=spec.get("presenter_notes", f"Keep the speaker note focused on the one-line takeaway for slide {slide_number}."),
        )
        slides.append(slide)

        evidence_plan.append(
            EvidencePlanItem(
                slide_number=slide_number,
                evidence_need=spec["takeaway"],
                asset_strategy=(
                    "Reuse local source material first."
                    if bridge.visual_source_preference == VisualSourcePreference.DOCUMENT_CROP
                    else "Build a slide-native visual from approved evidence."
                ),
                required_evidence_assets=required_assets,
                source_material_refs=bridge.source_material_refs,
                required_data=_required_data_for_visual(visual_type, required_assets),
                required_screenshots_or_diagrams=_required_screens_or_diagrams(visual_type, bridge),
                existing_assets=_dedupe([ref.label for ref in bridge.source_material_refs]),
                verification_needs=verification_flags,
                optional_assets=["Secondary proof objects can move to appendix if the main story becomes dense."],
                critical_assets=required_assets[:2] or [spec["takeaway"]],
            )
        )
        if visual_type in {VisualType.FRAMEWORK, VisualType.PROCESS, VisualType.INFOGRAPHIC, VisualType.TIMELINE}:
            infographic_plan.append(
                InfographicPlanItem(
                    slide_number=slide_number,
                    concept=spec["title"],
                    recommended_visual_type=visual_type,
                    purpose=spec["takeaway"],
                    core_message=spec["takeaway"],
                    required_data_labels=_required_data_for_visual(visual_type, required_assets),
                    why_visual=_infographic_why_visual(visual_type),
                    frame_fit_considerations=_frame_fit_considerations(visual_type),
                    fallback_simple_version="Fallback to a simpler table or comparison if density breaks the frame.",
                )
            )
    metadata = {
        "main_story_slide_budget": workflow_plan.main_story_slide_count_range,
        "appendix_slide_budget": workflow_plan.appendix_candidate_slide_count_range,
        "clustering_decisions": clustering_decisions,
        "concept_graph": lecture_artifacts.concept_graph,
        "teaching_plan": lecture_artifacts.teaching_plan,
        "blueprint_preview": lecture_artifacts.blueprint_preview,
        "authoring_preview": authoring_preview,
    }
    return slides, evidence_plan, infographic_plan, metadata


def _section_id_map(story_architecture: list[StorySection]) -> dict[str, str]:
    return {str(section.title): str(section.section_id) for section in story_architecture}


def _proof_coverage_class(slide: BlueprintSlide) -> ProofCoverageClass:
    if slide.slide_role == SlideRole.COMPARISON:
        return ProofCoverageClass.COMPARATIVE_SYNTHESIS
    if slide.slide_role == SlideRole.ANALYSIS:
        return ProofCoverageClass.INTERPRETIVE_SYNTHESIS
    if slide.visual_type in {VisualType.CHART, VisualType.TABLE}:
        return ProofCoverageClass.QUANTITATIVE_EVIDENCE
    return ProofCoverageClass.DIRECT_EVIDENCE


def _proof_evidence_origin(source_refs: list[SourceMaterialRef]) -> ProofEvidenceOrigin:
    if not source_refs:
        return ProofEvidenceOrigin.GENERATED
    origins: set[ProofEvidenceOrigin] = set()
    for ref in source_refs:
        suffix = Path(ref.path).suffix.lower() if ref.path else ""
        if suffix in {".csv", ".tsv", ".xlsx", ".xls"}:
            origins.add(ProofEvidenceOrigin.DATA)
        elif suffix in {".md", ".txt", ".yaml", ".yml", ".json"}:
            origins.add(ProofEvidenceOrigin.NOTES)
        else:
            origins.add(ProofEvidenceOrigin.DOCUMENT)
    if len(origins) == 1:
        return next(iter(origins))
    return ProofEvidenceOrigin.MIXED


def _build_shared_proof_unit_registry(
    workflow_plan: WorkflowPlan,
    blueprint: Blueprint,
) -> ProofUnitRegistry | None:
    selected_option_id = str(getattr(blueprint, "chosen_workflow", workflow_plan.workflow_option))
    policy = shared_proof_consumer_policy(selected_option_id)
    if policy is None:
        return None
    section_ids = _section_id_map(list(getattr(blueprint, "story_architecture", [])))
    main_story_slides = [slide for slide in blueprint.slides if slide.deck_mode == DeckMode.MAIN_STORY]
    claim_slides = [
        slide for slide in main_story_slides if slide.section == policy.claim_anchor_section_title
    ]
    proof_slides = [
        slide for slide in main_story_slides if slide.section == policy.proof_section_title
    ]
    implication_slides = [
        slide for slide in main_story_slides if slide.section == policy.synthesis_anchor_section_title
    ]
    claim_slide_numbers = [slide.slide_number for slide in claim_slides]
    implication_slide_numbers = [slide.slide_number for slide in implication_slides]
    claim_slide_number = claim_slide_numbers[0] if claim_slide_numbers else 1
    implication_slide_number = implication_slide_numbers[-1] if implication_slide_numbers else claim_slide_number
    workflow_provenance = getattr(blueprint, "workflow_option_provenance", None) or getattr(
        workflow_plan, "workflow_option_provenance", None
    )
    workflow_policy_id = str(getattr(workflow_provenance, "policy_id", "") or "")
    workflow_contract_status = (
        getattr(getattr(workflow_provenance, "contract_status", None), "value", getattr(workflow_provenance, "contract_status", None))
        or "absent"
    )

    units: list[ProofUnit] = []
    for module_order_index, slide in enumerate(proof_slides, start=1):
        source_refs = list(slide.production_bridge.source_material_refs)
        provenance_reason_codes: list[str] = []
        status_reason_codes: list[str] = []
        status = ProofModuleStatus.READY
        if slide.deck_mode != DeckMode.MAIN_STORY:
            status = ProofModuleStatus.INVALID
            status_reason_codes.append("appendix-placement-mismatch")
        if not source_refs:
            provenance_reason_codes.append("missing-source-material-refs")
            if status == ProofModuleStatus.READY:
                status = ProofModuleStatus.INCOMPLETE
            status_reason_codes.append("missing-source-material-refs")
        if not slide.required_evidence_assets:
            provenance_reason_codes.append("missing-required-evidence-assets")

        units.append(
            ProofUnit(
                unit_id=f"{policy.unit_id_prefix}-{module_order_index:02d}",
                workflow_option=policy.option_id,
                parent_section_id=section_ids.get(policy.proof_section_title, "proof-section"),
                unit_order_index=module_order_index,
                slide_number=slide.slide_number,
                blueprint_slide_ref=f"slide-{slide.slide_number:03d}",
                slide_title=slide.title,
                slide_role=slide.slide_role,
                layout_pattern_id=slide.layout_pattern_id,
                coverage_class=_proof_coverage_class(slide),
                evidence_origin=_proof_evidence_origin(source_refs),
                placement=slide.deck_mode,
                claim_anchor_section_id=section_ids.get(policy.claim_anchor_section_title, "claim-anchor"),
                claim_anchor_slide_number=claim_slide_number,
                synthesis_anchor_section_id=section_ids.get(
                    policy.synthesis_anchor_section_title,
                    "synthesis-anchor",
                ),
                synthesis_anchor_slide_number=implication_slide_number,
                claim_link_reason=policy.claim_link_reason,
                synthesis_link_reason=policy.synthesis_link_reason,
                required_evidence_assets=list(slide.required_evidence_assets),
                source_material_refs=source_refs,
                workflow_policy_id=workflow_policy_id,
                provenance_reason_codes=provenance_reason_codes,
                status=status,
                status_reason_codes=status_reason_codes,
            )
        )

    direct_evidence_module_count = sum(
        1
        for unit in units
        if unit.coverage_class in {ProofCoverageClass.DIRECT_EVIDENCE, ProofCoverageClass.QUANTITATIVE_EVIDENCE}
    )
    synthesis_module_count = sum(
        1
        for unit in units
        if unit.coverage_class in {ProofCoverageClass.COMPARATIVE_SYNTHESIS, ProofCoverageClass.INTERPRETIVE_SYNTHESIS}
    )
    return ProofUnitRegistry(
        deck_title=blueprint.deck_title,
        workflow_option=policy.option_id,
        requested_workflow_option=getattr(workflow_provenance, "requested_option_id", None),
        workflow_policy_id=workflow_policy_id,
        workflow_contract_status=str(workflow_contract_status),
        claim_anchor_section_id=section_ids.get(policy.claim_anchor_section_title, "claim-anchor"),
        proof_section_id=section_ids.get(policy.proof_section_title, "proof-section"),
        synthesis_anchor_section_id=section_ids.get(
            policy.synthesis_anchor_section_title,
            "synthesis-anchor",
        ),
        claim_anchor_slide_numbers=claim_slide_numbers,
        synthesis_anchor_slide_numbers=implication_slide_numbers,
        unit_minimum=policy.unit_minimum,
        unit_count=len(units),
        direct_evidence_unit_count=direct_evidence_module_count,
        synthesis_unit_count=synthesis_module_count,
        main_story_budget=blueprint.main_story_slide_budget,
        main_story_actual_slide_count=blueprint.main_story_actual_slide_count,
        appendix_actual_slide_count=blueprint.appendix_actual_slide_count,
        units=units,
    )


def _build_proof_module_manifest_from_registry(
    registry: ProofUnitRegistry | None,
    *,
    include_compat_manifest: bool = False,
) -> ProofModuleManifest | None:
    if registry is None:
        return None
    policy = shared_proof_consumer_policy(getattr(registry, "workflow_option", None))
    if (
        not include_compat_manifest
        or policy is None
        or policy.compat_mode != SharedProofCompatMode.LEGACY_MANIFEST_MIGRATION_ONLY
    ):
        return None
    return proof_module_manifest_from_proof_unit_registry(registry)


def _story_architecture_from_slides(slides: list[BlueprintSlide]) -> list[StorySection]:
    sections: list[StorySection] = []
    current_title: str | None = None
    current_group: list[BlueprintSlide] = []

    def flush(group: list[BlueprintSlide]) -> None:
        if not group:
            return
        title = group[0].section
        deck_mode = group[0].deck_mode
        roles: list[SlideRole] = []
        for slide in group:
            if slide.slide_role not in roles:
                roles.append(slide.slide_role)
        count = len(group)
        if deck_mode == DeckMode.APPENDIX:
            purpose = "Hold backup derivations, detailed evidence, and source-heavy support outside the main lecture."
            role_in_story = "Keep support material auditable without reopening the main teaching arc."
        elif title == "Orientation":
            purpose = "Align the lecture objective, roadmap, and reading rules before the core content."
            role_in_story = "Prepare the audience to follow the lecture with consistent teaching questions."
        elif title == "Closing":
            purpose = "Consolidate the teaching arc and restate the durable conceptual integration."
            role_in_story = "Turn the lecture into a stable synthesis rather than a loose recap."
        else:
            purpose = "Teach one concept cluster drawn from the synthesized teaching plan."
            role_in_story = "Advance the lecture through concept explanation, bridges, mechanisms, examples, and limits."
        sections.append(
            StorySection(
                section_id=_slugify(title),
                title=title,
                purpose=purpose,
                role_in_story=role_in_story,
                deck_mode=deck_mode,
                slide_count_range=CountRange(start=count, end=count),
                slide_roles=roles,
            )
        )

    for slide in slides:
        if current_title != slide.section or (current_group and current_group[-1].deck_mode != slide.deck_mode):
            flush(current_group)
            current_group = [slide]
            current_title = slide.section
            continue
        current_group.append(slide)
    flush(current_group)
    return sections


def _gate_status_map(workflow_plan: WorkflowPlan) -> dict[WorkflowGate, StageStatus]:
    return {gate.gate: gate.status for gate in workflow_plan.gates}


def _ensure_gate1_ready(workflow_plan: WorkflowPlan) -> None:
    gate_status = _gate_status_map(workflow_plan)
    workflow_ready = gate_status.get(WorkflowGate.WORKFLOW_DESIGN)
    if workflow_ready not in {StageStatus.READY, StageStatus.APPROVED}:
        raise ValueError("workflow_plan must be ready or approved at Gate 1 before Gate 2 planning")


def _default_primary_color(workflow_plan: WorkflowPlan) -> str:
    types = {workflow_plan.presentation_type_diagnosis.primary_type, *workflow_plan.presentation_type_diagnosis.secondary_types}
    if PresentationType.TRAINING in types or PresentationType.DEMO in types:
        return "#0F766E"
    if PresentationType.DECISION in types or PresentationType.REPORT in types:
        return "#1F2937"
    return "#1D4ED8"


def _default_accent_color(workflow_plan: WorkflowPlan) -> str:
    types = {workflow_plan.presentation_type_diagnosis.primary_type, *workflow_plan.presentation_type_diagnosis.secondary_types}
    if PresentationType.TRAINING in types or PresentationType.DEMO in types:
        return "#2563EB"
    if PresentationType.DECISION in types or PresentationType.PERSUASION in types:
        return "#C2410C"
    return "#0F766E"


def _theme_name(
    workflow_plan: WorkflowPlan,
    reference_dna: ReferenceDNA | None,
    brand_inputs: BrandInputs | None,
) -> str:
    if brand_inputs is not None and brand_inputs.brand_name:
        return _slugify(brand_inputs.brand_name)
    if reference_dna is not None and reference_dna.source_family:
        return _slugify(reference_dna.source_family)
    return _slugify(f"{workflow_plan.presentation_type_diagnosis.primary_type.value}-{workflow_plan.scale_mode.value}")


def _route_visual_families(recommended_route: str) -> list[str]:
    families = {
        "evidence-led-editorial": [
            "editorial thesis slides",
            "source-proof evidence frames",
            "structured comparison tables",
        ],
        "structured-visual-system": [
            "analysis charts",
            "comparison tables",
            "framework and process diagrams",
        ],
        "annotated-screen-and-proof": [
            "annotated crop frames",
            "process walkthrough slides",
            "section-divider resets",
        ],
    }
    return families.get(recommended_route, ["editorial thesis slides", "structured evidence visuals"])


def _title_rules(workflow_plan: WorkflowPlan, reference_dna: ReferenceDNA | None) -> list[str]:
    rules = [
        "Use takeaway-first titles that state the claim, not just the topic.",
        "Keep titles to one line when possible and never more than two short clauses.",
        "Keep one explicit idea per slide title and avoid stacking several findings in one heading.",
        "Use short section-divider labels that reset the audience without adding new evidence.",
    ]
    types = {workflow_plan.presentation_type_diagnosis.primary_type, *workflow_plan.presentation_type_diagnosis.secondary_types}
    if PresentationType.DECISION in types or PresentationType.PERSUASION in types:
        rules.append("Recommendation titles should imply the choice or operating shift directly.")
    if reference_dna is not None:
        rules.extend(reference_dna.pacing_and_title_behavior[:1])
    return _dedupe(rules)


def _terminology_rules(workflow_plan: WorkflowPlan, reference_dna: ReferenceDNA | None) -> list[str]:
    rules = [
        f"Use the approved deck language for the {workflow_plan.presentation_type_diagnosis.diagnosis_label} story consistently across titles, legends, and notes.",
        "Do not rename the same entity, metric, or workstream across sections.",
        "Promote one stable term for the main recommendation and reuse it everywhere.",
    ]
    if reference_dna is not None:
        rules.extend(reference_dna.terminology_guardrails)
    return _dedupe(rules)


def _build_design_system(
    workflow_plan: WorkflowPlan,
    recommended_route: str,
    reference_dna: ReferenceDNA | None,
    brand_inputs: BrandInputs | None,
) -> DesignSystem:
    primary_color = brand_inputs.primary_color if brand_inputs and brand_inputs.primary_color else _default_primary_color(workflow_plan)
    accent_color = brand_inputs.accent_color if brand_inputs and brand_inputs.accent_color else _default_accent_color(workflow_plan)
    font_family = brand_inputs.font_family if brand_inputs and brand_inputs.font_family else "Aptos"
    display_family = f"{font_family} Display" if brand_inputs and brand_inputs.font_family else "Aptos Display"
    title_rules = _title_rules(workflow_plan, reference_dna)
    tone_keywords = _dedupe(
        [
            *(brand_inputs.tone_keywords if brand_inputs is not None else []),
            workflow_plan.presentation_type_diagnosis.primary_type.value,
            workflow_plan.scale_mode.value,
            "evidence-led",
            "slide-native",
        ]
    )
    layout_principles = _dedupe(
        [
            "Keep one dominant thesis or visual zone per slide.",
            "Preserve stable title, evidence, and footer positions across the deck.",
            "Use export-safe spacing and approved layout patterns only.",
            *(reference_dna.layout_logic[:2] if reference_dna is not None else []),
        ]
    )
    layout_rules = _dedupe(
        [
            "Use only approved layout patterns and keep their composition logic stable across the deck.",
            "Let the title band, primary evidence zone, and footer/source zone recur predictably from slide to slide.",
            "Use section-divider layouts only at declared story boundaries, never as decoration between similar analytical slides.",
        ]
    )
    visual_system_rules = _dedupe(
        [
            f"Apply the `{recommended_route}` route consistently across the deck unless an approved appendix exception is recorded.",
            "Keep one dominant visual family per slide and avoid mixing unrelated chart, crop, and icon styles inside the same frame.",
            "Use local reference behavior to guide hierarchy and pacing without copying literal template layouts.",
        ]
    )
    chart_rules = _dedupe(
        [
            "Avoid dual axes and dense legends in the main story.",
            "Use direct labels and end-state emphasis before adding extra series.",
            *(brand_inputs.chart_preferences if brand_inputs is not None else []),
            *(reference_dna.chart_table_treatment[:2] if reference_dna is not None else []),
        ]
    )
    table_rules = _dedupe(
        [
            "Use tables only when exact lookup matters more than shape or trend.",
            "Keep tables narrow enough to preserve presentation-distance readability.",
            "Do not let appendix tables reset the deck's title or spacing hierarchy.",
            *(reference_dna.chart_table_treatment[1:3] if reference_dna is not None else []),
        ]
    )
    highlight_rules = _dedupe(
        [
            "Reserve the signal color for one or two emphasis moves per slide.",
            "Use highlights to clarify the message path, not to decorate every node.",
            "Apply emphasis consistently to titles, key metrics, and end-state comparisons only.",
        ]
    )
    screenshot_rules = _dedupe(
        [
            "Use screenshots or cropped source visuals only when they prove a point that a rebuilt visual cannot carry as credibly.",
            "Frame screenshots inside a clean visual window with one annotation or caption lane rather than scattered callouts.",
            "Fallback to structured diagrams when the screenshot is too dense, too small, or too caption-dependent to read cleanly.",
        ]
    )
    infographic_rules = _dedupe(
        [
            "Use infographic slides only when the structure is clearer than a short bullet sequence.",
            "Fallback to simpler frameworks or tables if node density breaks one-slide readability.",
            "Keep infographic labels short enough to read in three seconds.",
        ]
    )
    section_divider_style = (
        brand_inputs.section_divider_style
        if brand_inputs is not None and brand_inputs.section_divider_style
        else (
            reference_dna.section_divider_style
            if reference_dna is not None
            else "Muted divider band with one short title and a clear visual reset."
        )
    )
    icon_style = (
        brand_inputs.icon_style
        if brand_inputs is not None and brand_inputs.icon_style
        else (
            reference_dna.icon_illustration_treatment[0]
            if reference_dna is not None and reference_dna.icon_illustration_treatment
            else "Minimal icon family with restrained stroke weight."
        )
    )
    return DesignSystem(
        deck_title=workflow_plan.deck_title,
        slide_ratio=workflow_plan.slide_ratio,
        brand_name=brand_inputs.brand_name if brand_inputs is not None else None,
        theme_name=_theme_name(workflow_plan, reference_dna, brand_inputs),
        visual_route_id=recommended_route,
        reference_source_family=reference_dna.source_family if reference_dna is not None else None,
        tone_keywords=tone_keywords,
        color_tokens=[
            ColorToken(token="ink", hex=primary_color, usage="Titles, key numbers, and primary thesis text"),
            ColorToken(token="signal", hex=accent_color, usage="Highlights, recommendation cues, and emphasis"),
            ColorToken(token="canvas", hex="#F8FAFC", usage="Background panels and quiet supporting zones"),
        ],
        typography_tokens=[
            TypographyToken(token="title", font_family=display_family, size_pt=24, weight="bold", usage="Slide titles"),
            TypographyToken(token="body", font_family=font_family, size_pt=11, weight="regular", usage="Body text"),
            TypographyToken(token="caption", font_family=font_family, size_pt=9, weight="regular", usage="Captions and source notes"),
        ],
        spacing_scale=[4, 8, 16, 24, 32],
        layout_principles=layout_principles,
        chart_rules=chart_rules,
        table_rules=table_rules,
        screenshot_rules=screenshot_rules,
        highlight_rules=highlight_rules,
        callout_method="Use one signal-color callout chip or underline per slide, paired with direct labels rather than floating annotation clutter.",
        title_rules=title_rules,
        section_divider_style=section_divider_style,
        visual_families=_route_visual_families(recommended_route),
        infographic_rules=infographic_rules,
        layout_rules=layout_rules,
        visual_system_rules=visual_system_rules,
        icon_style=icon_style,
    )


def _appendix_start(slides: list[BlueprintSlide]) -> int | None:
    for slide in slides:
        if slide.deck_mode == DeckMode.APPENDIX:
            return slide.slide_number
    return None


def _build_workflow_delta(
    workflow_plan: WorkflowPlan,
    slides: list[BlueprintSlide],
    recommended_route: str,
    reference_dna: ReferenceDNA | None,
) -> list[str]:
    main_story_count = len([slide for slide in slides if slide.deck_mode == DeckMode.MAIN_STORY])
    appendix_count = len([slide for slide in slides if slide.deck_mode == DeckMode.APPENDIX])
    crop_count = len([slide for slide in slides if slide.production_bridge.visual_source_preference == VisualSourcePreference.DOCUMENT_CROP])
    deltas = [
        f"Locked workflow `{workflow_plan.workflow_option}` into {main_story_count} main-story slides.",
        f"Reserved {appendix_count} appendix candidate slides." if appendix_count else "No appendix slides were required in the current blueprint pass.",
        f"Recommended visual route is `{recommended_route}` with {crop_count} crop- or source-reuse-biased slides.",
        "Converted workflow ranges into numbered slides, named sections, and production-facing bridge fields.",
    ]
    if reference_dna is not None and reference_dna.source_family:
        deltas.append(f"Adapted local reference behavior from `{reference_dna.source_family}` without copying literal template layouts.")
    return _dedupe(deltas)


def _workflow_delta_details(
    workflow_plan: WorkflowPlan,
    slides: list[BlueprintSlide],
    recommended_route: str,
    reference_dna: ReferenceDNA | None,
) -> list[WorkflowDeltaDetail]:
    main_story_count = len([slide for slide in slides if slide.deck_mode == DeckMode.MAIN_STORY])
    appendix_count = len([slide for slide in slides if slide.deck_mode == DeckMode.APPENDIX])
    crop_count = len([slide for slide in slides if slide.production_bridge.visual_source_preference == VisualSourcePreference.DOCUMENT_CROP])
    details = [
        WorkflowDeltaDetail(
            change=f"Finalized `{workflow_plan.workflow_option}` as a numbered blueprint with {main_story_count} main-story slides.",
            reason="Gate 2 needs fixed slide records, not Gate 1 ranges, before any production worker can proceed safely.",
            deck_impact="The deck now has stable slide ids, sections, and layout assignments for every approved main-story slide.",
        ),
        WorkflowDeltaDetail(
            change=(
                f"Reserved {appendix_count} appendix slides."
                if appendix_count
                else "Collapsed appendix usage to zero slides in the current pass."
            ),
            reason="Support detail should stay outside the main story unless it changes the audience decision or learning shift.",
            deck_impact="Main-story slides stay shorter and the appendix boundary becomes explicit for downstream production and QA.",
        ),
        WorkflowDeltaDetail(
            change=f"Selected `{recommended_route}` as the deck-wide visual route with {crop_count} source-reuse-biased slides.",
            reason="Gate 2 must lock one dominant visual system before crops, visuals, compile, or QA work begins.",
            deck_impact="The design system, layout library, and slide ledger now point to one coherent route instead of several mixed styles.",
        ),
    ]
    if reference_dna is not None and reference_dna.source_family:
        details.append(
            WorkflowDeltaDetail(
                change=f"Adapted direction from the local `{reference_dna.source_family}` reference family.",
                reason="The planner should borrow local reference hierarchy and pacing without copying any single template.",
                deck_impact="Reference-derived spacing, proof framing, and divider behavior are now reflected in the approved Gate 2 package.",
            )
        )
    return details


def _build_assumptions(
    workflow_plan: WorkflowPlan,
    recommended_route: str,
    reference_dna: ReferenceDNA | None,
    brand_inputs: BrandInputs | None,
) -> list[str]:
    assumptions = list(workflow_plan.assumptions)
    assumptions.append(f"Assume the `{recommended_route}` route can remain stable across the approved deck.")
    if reference_dna is None:
        assumptions.append("Assume local reference scanning can be added later without breaking downstream Gate 2 contracts.")
    if brand_inputs is None:
        assumptions.append("Assume a neutral executive visual system is acceptable until explicit brand inputs arrive.")
    return _dedupe(assumptions)


def _build_risks(
    workflow_plan: WorkflowPlan,
    slides: list[BlueprintSlide],
    reference_dna: ReferenceDNA | None,
    brand_inputs: BrandInputs | None,
) -> list[str]:
    risks = list(workflow_plan.risks)
    crop_count = len([slide for slide in slides if slide.production_bridge.visual_source_preference == VisualSourcePreference.DOCUMENT_CROP])
    if crop_count >= 3:
        risks.append("Several slides depend on clean source-derived crops, so crop review discipline remains critical.")
    if reference_dna is None:
        risks.append("No local reference_dna was supplied, so the visual system relies on deterministic defaults only.")
    if brand_inputs is None:
        risks.append("Brand inputs are absent, so theme approval may still change before production.")
    if _supports_continuity_controls(workflow_plan.scale_mode):
        risks.append("Large-deck continuity depends on the ledger and constitution remaining authoritative across batches.")
    return _dedupe(risks)


def _build_verification_points(
    workflow_plan: WorkflowPlan,
    slides: list[BlueprintSlide],
    reference_dna: ReferenceDNA | None,
) -> list[VerificationPoint]:
    appendix_start = _appendix_start(slides)
    points = [
        VerificationPoint(
            checkpoint="Story fit",
            rationale=f"Confirm the numbered blueprint still fits the approved `{workflow_plan.workflow_option}` workflow and the {workflow_plan.scale_mode.value} scale target.",
        ),
        VerificationPoint(
            checkpoint="One-slide-one-message",
            rationale="Every slide title, takeaway, and main message should still collapse to one claim before production begins.",
        ),
        VerificationPoint(
            checkpoint="Evidence readiness",
            rationale="Slides that prefer document crops or structured evidence must have usable local source references before production handoff.",
        ),
        VerificationPoint(
            checkpoint="Terminology lock",
            rationale="Approve terminology, numbering, and section labels before any batch production starts.",
        ),
    ]
    if appendix_start is not None:
        points.append(
            VerificationPoint(
                checkpoint="Appendix boundary",
                rationale=f"Confirm that slides {appendix_start} onward are support-only and do not introduce new main-story recommendations.",
            )
        )
    if reference_dna is not None:
        points.append(
            VerificationPoint(
                checkpoint="Reference fit",
                rationale="Confirm the blueprint borrows the reference pack's hierarchy and pacing behavior without copying literal layouts.",
            )
        )
    return points


def _build_deck_constitution(
    workflow_plan: WorkflowPlan,
    design_system: DesignSystem,
    layout_library: LayoutLibrary,
    slides: list[BlueprintSlide],
    reference_dna: ReferenceDNA | None,
) -> DeckConstitution:
    appendix_start = _appendix_start(slides)
    section_titles = _dedupe([slide.section for slide in slides if slide.slide_role != SlideRole.SECTION_DIVIDER])
    message_spine = [slide.main_message for slide in slides if slide.deck_mode == DeckMode.MAIN_STORY][:4]
    design_token_refs = [token.token for token in design_system.color_tokens] + [token.token for token in design_system.typography_tokens]
    appendix_policy = (
        f"Slides {appendix_start} onward are appendix-only and may contain methods, backup evidence, or references but no new main-story recommendation."
        if appendix_start is not None
        else "No appendix boundary is currently approved; keep all slides inside the main story."
    )
    lecture_mode = _is_lecture_mode(workflow_plan)
    story_rules = [
        "Each main-story slide must carry one explicit takeaway and one dominant proof or structure.",
        "Lead with the recommendation or thesis before methods, references, or backup detail.",
        f"Keep the current approved section order stable: {', '.join(section_titles)}.",
    ]
    qa_stop_conditions = [
        f"Stop after {workflow_plan.bounded_qa_rounds} failed remediation rounds on the same issue cluster.",
        "Escalate immediately if numbering, terminology, section hierarchy, or appendix boundary drift recurs.",
    ]
    section_divider_rules = [
        "Only use divider slides at declared story architecture boundaries.",
        "Divider slides must not introduce new evidence or a second narrative.",
        f"Use the approved divider style: {design_system.section_divider_style}",
    ]
    if lecture_mode:
        story_rules.extend(
            [
                "Cluster aligned subsections into teaching modules instead of mirroring document headings one by one.",
                "Keep supporting examples and derivations out of the main lecture unless they directly unlock comprehension of the next concept.",
                "Do not emit repetitive section-preview rhythm slides unless a major lecture phase genuinely changes.",
            ]
        )
        qa_stop_conditions.extend(
            [
                "Block ship if internal, placeholder, or fallback-debug text becomes visible on any slide.",
                "Block ship if lecture slide counts drift outside the approved main-story or appendix budget without a recorded exception.",
                "Block ship if a slide uses a layout outside the approved lecture layout library or violates its supported content types.",
            ]
        )
        section_divider_rules.insert(1, "Divider slides are allowed only for orientation, major chapter boundaries, closing, or appendix start.")
    return DeckConstitution(
        deck_title=workflow_plan.deck_title,
        deck_mode=workflow_plan.deck_mode,
        locked_workflow=workflow_plan.workflow_option,
        deck_objective=workflow_plan.objective,
        audience_definition=list(workflow_plan.audience),
        delivery_mode=workflow_plan.project_snapshot.delivery_mode.value,
        message_spine=message_spine,
        narrative_promise=workflow_plan.objective,
        section_logic=[
            f"{section_title}: keep the declared section order stable and do not move slides across the appendix boundary without a state update."
            for section_title in section_titles
        ],
        story_rules=story_rules,
        title_rules=design_system.title_rules,
        terminology_rules=_terminology_rules(workflow_plan, reference_dna),
        tone_voice=list(design_system.tone_keywords),
        approved_visual_route=design_system.visual_route_id,
        design_token_refs=_dedupe(design_token_refs),
        layout_pattern_ids=[pattern.pattern_id for pattern in layout_library.patterns],
        infographic_rules=list(design_system.infographic_rules),
        chart_rules=list(design_system.chart_rules),
        table_rules=list(design_system.table_rules),
        screenshot_rules=list(design_system.screenshot_rules),
        icon_style=design_system.icon_style,
        appendix_policy=appendix_policy,
        appendix_boundary_rule=(
            f"Appendix begins at slide {appendix_start}; appendix slides must support the main story and may not reset the recommendation."
            if appendix_start is not None
            else "No appendix boundary is active in the current blueprint pass."
        ),
        methods_policy="Keep methods detail in appendix slides, support notes, or later QA artifacts unless it changes the decision or operating shift.",
        references_policy="Keep references compact, off the title line, and traceable through source material refs or asset manifests.",
        numbering_rules=[
            "Keep slide numbers sequential across blueprint, ledger, and compiled output.",
            "Do not renumber approved slides inside a batch without updating the ledger and any downstream manifests.",
            "If appendix slides exist, continue numbering rather than restarting.",
        ],
        section_divider_rules=section_divider_rules,
        navigation_rules=[
            "Use section-divider slides and stable numbering to show the audience where they are in the story.",
            "Do not introduce ad hoc navigation labels that conflict with the approved section names or appendix boundary.",
        ],
        source_handling_rules=[
            "Use only approved local source material refs as the basis for crops, screenshots, charts, or tables.",
            "Do not let source-derived visuals appear without traceable provenance or a clear fallback path.",
            "Use reference behaviors as guidance only; never copy a local reference template literally.",
        ],
        recurring_motifs=[
            "Short thesis-first titles",
            "Restrained emphasis using the signal color",
            "One dominant evidence or structure zone per slide",
        ],
        continuity_rules=[
            "Keep approved terminology consistent across titles, callouts, legends, and speaker notes.",
            "Keep section ordering, visual route, and layout patterns stable across batches.",
            "Main story and appendix must stay visibly distinct in numbering and section labels.",
        ],
        visual_consistency_rules=[
            f"Use the `{design_system.theme_name}` design system and `{design_system.visual_route_id}` route throughout the deck.",
            "Only approved layout patterns may be used during production and compile.",
            "Prefer export-safe slide-native visuals and disciplined crops over ad hoc composition tricks.",
        ],
        qa_stop_conditions=qa_stop_conditions,
    )


def _build_layout_library() -> LayoutLibrary:
    return LayoutLibrary(
        deck_title="placeholder",
        patterns=[
            LayoutPattern(
                pattern_id="cover",
                name="Cover",
                slide_roles=[SlideRole.TITLE],
                supported_visual_types=[VisualType.TEXT, VisualType.QUOTE],
                body_slots=1,
                safe_area_notes="Keep the cover sparse, with strong whitespace and one dominant title block.",
                use_cases=["Deck opener", "Title slide", "Opening thesis"],
                composition_logic="Use a large title zone, one short subtitle or context line, and minimal supporting ornament.",
                density_guidance="Extremely low density. Avoid secondary evidence objects on the cover.",
                fit_risks=["Can feel empty if too many metadata elements are added.", "Should not be reused for evidence slides."],
            ),
            LayoutPattern(
                pattern_id="concept-explainer",
                name="Concept explainer",
                slide_roles=[SlideRole.ANALYSIS, SlideRole.INFOGRAPHIC, SlideRole.EXECUTIVE_SUMMARY],
                supported_visual_types=[VisualType.FRAMEWORK, VisualType.INFOGRAPHIC, VisualType.TEXT, VisualType.HIERARCHY, VisualType.METRIC_SUMMARY],
                body_slots=2,
                safe_area_notes="Reserve equal breathing room for the concept visual and the interpretation layer.",
                use_cases=["Concept explainer", "Model introduction", "Mechanism slide"],
                composition_logic="Pair one structured concept visual with a short interpretation rail that explains why the structure matters.",
                density_guidance="Medium density. Cap concept labels before the slide turns into a worksheet.",
                fit_risks=["Can become abstract if no business implication is stated.", "Frameworks with too many nodes will not survive presentation distance."],
            ),
            LayoutPattern(
                pattern_id="definition-theorem",
                name="Definition or theorem",
                slide_roles=[SlideRole.ANALYSIS, SlideRole.EVIDENCE, SlideRole.EXECUTIVE_SUMMARY],
                supported_visual_types=[VisualType.TEXT, VisualType.QUOTE, VisualType.FRAMEWORK],
                body_slots=2,
                safe_area_notes="Keep the formal statement short and leave room for one interpretation block.",
                use_cases=["Definitions", "Optimality conditions", "Short theorem statements"],
                composition_logic="Use a stable title band, one short formal statement, and one interpretation or implication block.",
                density_guidance="Low to medium density. Avoid multi-paragraph derivations on the visible slide.",
                fit_risks=["Breaks if proof text turns into paragraph blocks.", "Dense equations should move to appendix support."],
            ),
            LayoutPattern(
                pattern_id="comparison",
                name="Comparison",
                slide_roles=[SlideRole.COMPARISON, SlideRole.EVIDENCE, SlideRole.APPENDIX_EVIDENCE],
                supported_visual_types=[VisualType.COMPARISON, VisualType.TABLE],
                body_slots=2,
                safe_area_notes="Keep the comparison object dominant and the teaching takeaway short.",
                use_cases=["Method comparison", "Tradeoff table", "Constraint vs. method mapping"],
                composition_logic="Use one comparison object with one short takeaway block and minimal secondary annotation.",
                density_guidance="Medium density at most. Cut rows and columns aggressively in the main story.",
                fit_risks=["Large matrices overwhelm the teaching flow.", "Tables with lookup-heavy detail belong in appendix."],
            ),
            LayoutPattern(
                pattern_id="process-flow",
                name="Process or algorithm flow",
                slide_roles=[SlideRole.PROCESS, SlideRole.RECOMMENDATION, SlideRole.EXECUTIVE_SUMMARY, SlideRole.ANALYSIS],
                supported_visual_types=[VisualType.TIMELINE, VisualType.PROCESS, VisualType.DECISION_PATH],
                body_slots=1,
                safe_area_notes="Show one left-to-right or top-to-bottom flow with stable step spacing and minimal ornament.",
                use_cases=["Algorithm steps", "Method selection flow", "Lecture roadmap"],
                composition_logic="Use one linear flow with short step labels and one highlighted decision or handoff moment.",
                density_guidance="Low to medium density. Cap step count before the flow becomes a poster.",
                fit_risks=["Branching complexity weakens the reading order.", "Step labels that are too long will collapse spacing."],
            ),
            LayoutPattern(
                pattern_id="worked-example",
                name="Worked example",
                slide_roles=[SlideRole.EVIDENCE, SlideRole.ANALYSIS, SlideRole.APPENDIX_EVIDENCE],
                supported_visual_types=[VisualType.DOCUMENT_CROP, VisualType.PHOTO, VisualType.CHART],
                body_slots=2,
                safe_area_notes="Preserve a clean crop frame and a separate caption or provenance lane below the visual.",
                use_cases=["Source proof", "Worked visual", "Chart-led example"],
                composition_logic="Let one worked visual dominate the slide, then place a short interpretation or provenance lane beside or below it.",
                density_guidance="Medium density. One dominant crop only; keep annotation count low.",
                fit_risks=["Fails quickly if the crop is too dense or too small.", "Needs a clear fallback route when the source visual is weak."],
            ),
            LayoutPattern(
                pattern_id="summary",
                name="Summary",
                slide_roles=[SlideRole.RECOMMENDATION, SlideRole.EXECUTIVE_SUMMARY, SlideRole.ANALYSIS],
                supported_visual_types=[VisualType.TEXT, VisualType.COMPARISON, VisualType.CHART, VisualType.QUOTE, VisualType.PROCESS],
                body_slots=2,
                safe_area_notes="Reserve a stable thesis zone and one compact support object without collapsing margins.",
                use_cases=["Lecture summary", "Method selection", "Close-out slide"],
                composition_logic="Use a thesis-first title, one short claim card, and one compact reinforcement object or bullet stack.",
                density_guidance="Low to medium density. The takeaway must remain visually dominant.",
                fit_risks=["Loses impact if the support block becomes a mini-report page.", "Should not carry several unrelated proof objects."],
            ),
            LayoutPattern(
                pattern_id="section-divider",
                name="Section divider",
                slide_roles=[SlideRole.SECTION_DIVIDER],
                supported_visual_types=[VisualType.TEXT],
                body_slots=1,
                safe_area_notes="Keep divider slides visually quiet with a single reset label and strong whitespace.",
                use_cases=["Section dividers", "Pacing resets"],
                composition_logic="Use one short section label, a restrained band or block, and strong whitespace to reset pacing.",
                density_guidance="Very low density. No new analytical content.",
                fit_risks=["Becomes decorative noise if overused.", "Should never carry a second message or evidence object."],
            ),
            LayoutPattern(
                pattern_id="appendix-reference",
                name="Appendix reference",
                slide_roles=[SlideRole.APPENDIX_EVIDENCE, SlideRole.REFERENCES, SlideRole.COMPARISON],
                supported_visual_types=[VisualType.TABLE, VisualType.TEXT, VisualType.DOCUMENT_CROP, VisualType.CHART, VisualType.COMPARISON],
                body_slots=2,
                safe_area_notes="Preserve a compact source or method lane and keep appendix labeling obvious.",
                use_cases=["Reference slides", "Methods appendix", "Backup evidence"],
                composition_logic="Combine one appendix-only proof or method block with a compact source or note lane that makes the support role explicit.",
                density_guidance="Medium density in appendix only. Still preserve title and source hierarchy.",
                fit_risks=["Can bleed back into the main story if numbering or labeling is unclear.", "Appendix slides should not introduce a new recommendation."],
            ),
        ],
    )


def _with_deck_title(layout_library: LayoutLibrary, deck_title: str, slide_ratio: str) -> LayoutLibrary:
    payload = layout_library.model_dump(mode="json", exclude_none=True)
    payload["deck_title"] = deck_title
    payload["slide_ratio"] = slide_ratio
    return LayoutLibrary.model_validate(payload)


def _build_slide_ledger(
    workflow_plan: WorkflowPlan,
    slides: list[BlueprintSlide],
    asset_requests: list[AssetRequest],
) -> SlideLedger:
    request_ids_by_slide: dict[int, list[str]] = {}
    dependency_kinds_by_slide: dict[int, list[AssetKind]] = {}
    for request in asset_requests:
        request_ids_by_slide.setdefault(request.slide_number, []).append(request.request_id)
        dependency_kinds_by_slide.setdefault(request.slide_number, [])
        if request.asset_kind not in dependency_kinds_by_slide[request.slide_number]:
            dependency_kinds_by_slide[request.slide_number].append(request.asset_kind)

    entries: list[SlideLedgerEntry] = []
    for slide in slides:
        section_id = _slugify(slide.section)
        part_id = "appendix" if slide.deck_mode == DeckMode.APPENDIX else "main-story"
        cluster_id = f"cluster-{section_id}"
        if _supports_continuity_controls(workflow_plan.scale_mode):
            batch_id = f"batch-{section_id}"
        elif slide.deck_mode == DeckMode.APPENDIX:
            batch_id = "batch-appendix-01"
        else:
            batch_id = None
        asset_request_ids = request_ids_by_slide.get(slide.slide_number, [])
        asset_dependency_kinds = dependency_kinds_by_slide.get(slide.slide_number, [])
        entries.append(
            SlideLedgerEntry(
                slide_number=slide.slide_number,
                slide_id=_slide_id(slide.slide_number),
                slide_role=slide.slide_role,
                slide_intent=slide.slide_intent,
                title=slide.title,
                final_title=slide.title,
                title_status=StageStatus.READY,
                one_line_takeaway=slide.one_line_takeaway,
                main_message=slide.main_message,
                pedagogical_goal=slide.pedagogical_goal,
                concept_ids=slide.concept_ids,
                section=slide.section,
                lineage_id=_slide_id(slide.slide_number),
                part_id=part_id,
                section_id=section_id,
                cluster_id=cluster_id,
                deck_mode=slide.deck_mode,
                content_tier=slide.content_tier,
                visual_type=slide.visual_type,
                visual_source_preference=slide.production_bridge.visual_source_preference,
                production_mode=slide.production_bridge.production_mode,
                layout_pattern_id=slide.layout_pattern_id,
                required_evidence_assets=slide.required_evidence_assets,
                asset_request_ids=asset_request_ids,
                asset_dependency_kinds=asset_dependency_kinds,
                batch_id=batch_id,
                blueprint_status=StageStatus.READY,
                asset_status=StageStatus.READY if asset_request_ids else StageStatus.COMPLETE,
                visual_status=StageStatus.DRAFT,
                compile_status=StageStatus.DRAFT,
                depends_on=[slide.slide_number - 1] if slide.slide_number > 1 and slide.slide_role != SlideRole.SECTION_DIVIDER else [],
                change_note="Derived from the approved Gate 2 blueprint package.",
            )
        )

    appendix_start = _appendix_start(slides)
    continuity_notes = [
        "Treat the slide ledger as the authoritative index for numbering, section order, and artifact status.",
        "Keep layout pattern ids, terminology, and visual route consistent with the approved design system.",
        "Keep lecture-core content in the main story and route support-heavy detail into appendix-only slides.",
    ]
    if appendix_start is not None:
        continuity_notes.append(f"Appendix begins at slide {appendix_start} and must remain outside the main story.")
    if _supports_continuity_controls(workflow_plan.scale_mode):
        continuity_notes.append("Large-deck production should honor section-based batches to reduce continuity drift.")

    return SlideLedger(
        deck_title=workflow_plan.deck_title,
        entries=entries,
        continuity_notes=continuity_notes,
    )


def plan_gate2(
    workflow_plan: WorkflowPlan,
    brief: WorkflowBriefInput | None = None,
    brand_inputs: BrandInputs | None = None,
    reference_dna: ReferenceDNA | None = None,
) -> Gate2Outputs:
    _ensure_gate1_ready(workflow_plan)
    materials = _merge_materials(workflow_plan, brief)
    grouped_materials = _material_groups(materials)
    fallback_refs = _material_refs(materials)
    story_architecture = _story_architecture(workflow_plan)
    selected_option = _selected_workflow_option(workflow_plan)
    chosen_workflow_phases = _chosen_workflow_phases(workflow_plan, story_architecture)
    story_structure = _story_structure(workflow_plan, story_architecture)
    visual_routes = _visual_routes(workflow_plan, reference_dna, grouped_materials)
    recommended_route = _recommended_visual_route(workflow_plan, reference_dna, grouped_materials)
    recommended_route_reason = _recommended_route_reason(workflow_plan, recommended_route, reference_dna, grouped_materials)
    visual_reference_summary = _visual_reference_summary(reference_dna, visual_routes, recommended_route)
    presentation_brief = workflow_plan.presentation_brief
    slide_function_outline = workflow_plan.slide_function_outline
    canonical_generation_profile = build_canonical_generation_profile(
        deck_title=workflow_plan.deck_title,
        materials=materials,
        presentation_brief=presentation_brief,
        slide_function_outline=slide_function_outline,
        reference_dna=reference_dna,
        brand_context=brand_inputs.model_dump(mode="json", exclude_none=True) if brand_inputs is not None else None,
    )
    slides, evidence_plan, infographic_plan, blueprint_meta = _slide_specs_to_blueprint(
        workflow_plan,
        grouped_materials,
        fallback_refs,
    )
    slides = _apply_slide_function_outline(slides, slide_function_outline)
    concept_graph = blueprint_meta["concept_graph"]
    teaching_plan = blueprint_meta["teaching_plan"]
    blueprint_preview = blueprint_meta["blueprint_preview"]
    authoring_preview = blueprint_meta["authoring_preview"]
    if _is_lecture_mode(workflow_plan):
        story_architecture = _story_architecture_from_slides(slides)
        chosen_workflow_phases = _chosen_workflow_phases(workflow_plan, story_architecture)
        story_structure = _story_structure(workflow_plan, story_architecture)
    asset_requests = derive_asset_requests_from_blueprint(slides)
    design_system = _build_design_system(workflow_plan, recommended_route, reference_dna, brand_inputs)
    layout_library = _with_deck_title(_build_layout_library(), workflow_plan.deck_title, workflow_plan.slide_ratio)
    deck_constitution = _build_deck_constitution(workflow_plan, design_system, layout_library, slides, reference_dna)
    slide_ledger = _build_slide_ledger(workflow_plan, slides, asset_requests)
    workflow_delta = _build_workflow_delta(workflow_plan, slides, recommended_route, reference_dna)
    workflow_delta_details = _workflow_delta_details(workflow_plan, slides, recommended_route, reference_dna)
    assumptions = _build_assumptions(workflow_plan, recommended_route, reference_dna, brand_inputs)
    risks = _build_risks(workflow_plan, slides, reference_dna, brand_inputs)
    verification_points = _build_verification_points(workflow_plan, slides, reference_dna)
    blueprint = Blueprint(
        deck_title=workflow_plan.deck_title,
        chosen_workflow=workflow_plan.workflow_option,
        lecture_family=concept_graph.lecture_family if concept_graph.nodes else None,
        lecture_family_evidence=concept_graph.lecture_family_evidence,
        rejected_families=concept_graph.rejected_families,
        central_concepts=blueprint_preview.central_concepts,
        flagged_drift_risks=teaching_plan.flagged_drift_risks,
        repetition_stats=teaching_plan.repetition_stats,
        chosen_workflow_label=selected_option.label,
        chosen_workflow_summary=selected_option.summary,
        chosen_workflow_phases=chosen_workflow_phases,
        workflow_option_provenance=workflow_plan.workflow_option_provenance,
        workflow_delta=workflow_delta,
        workflow_delta_details=workflow_delta_details,
        presentation_brief=presentation_brief,
        canonical_generation_profile=canonical_generation_profile,
        slide_function_outline=slide_function_outline,
        communication_core=_communication_core(workflow_plan),
        story_architecture=story_architecture,
        story_structure=story_structure,
        deck_mode=workflow_plan.deck_mode,
        slide_ratio=workflow_plan.slide_ratio,
        approval_status=StageStatus.READY,
        visual_reference_summary=visual_reference_summary,
        visual_routes=visual_routes,
        recommended_route=recommended_route,
        recommended_route_reason=recommended_route_reason,
        infographic_plan=infographic_plan,
        evidence_asset_plan=evidence_plan,
        assumptions=assumptions,
        risks=risks,
        verification_points=verification_points,
        slides=slides,
        appendix_start=_appendix_start(slides),
        main_story_slide_budget=blueprint_meta.get("main_story_slide_budget"),
        appendix_slide_budget=blueprint_meta.get("appendix_slide_budget"),
        clustering_decisions=blueprint_meta.get("clustering_decisions", []),
    )
    proof_unit_registry = _build_shared_proof_unit_registry(workflow_plan, blueprint)
    proof_module_manifest = _build_proof_module_manifest_from_registry(proof_unit_registry)
    return Gate2Outputs(
        blueprint=blueprint,
        proof_unit_registry=proof_unit_registry,
        proof_module_manifest=proof_module_manifest,
        presentation_brief=presentation_brief,
        canonical_generation_profile=canonical_generation_profile,
        slide_function_outline=slide_function_outline,
        concept_graph=concept_graph,
        teaching_plan=teaching_plan,
        blueprint_preview=blueprint_preview,
        authoring_preview=authoring_preview,
        design_system=design_system,
        deck_constitution=deck_constitution,
        layout_library=layout_library,
        slide_ledger=slide_ledger,
        asset_requests=AssetRequests(deck_title=workflow_plan.deck_title, requests=asset_requests),
    )


def plan_gate2_from_files(
    workflow_plan_path: str | Path,
    brief_path: str | Path | None = None,
    reference_dna_path: str | Path | None = None,
    brand_inputs_path: str | Path | None = None,
) -> Gate2Outputs:
    context = load_gate2_context(workflow_plan_path, reference_dna_path)
    brief = load_workflow_brief(brief_path) if brief_path is not None else None
    brand_inputs = load_brand_inputs(brand_inputs_path) if brand_inputs_path is not None else None
    return plan_gate2(
        context.workflow_plan,
        brief=brief,
        brand_inputs=brand_inputs,
        reference_dna=context.reference_dna,
    )


def write_gate2_outputs(outputs: Gate2Outputs, output_dir: str | Path) -> dict[str, Path]:
    resolved_output_dir = Path(output_dir)
    resolved_output_dir.mkdir(parents=True, exist_ok=True)
    models = {
        "blueprint": outputs.blueprint,
        "presentation_brief": outputs.presentation_brief,
        "canonical_generation_profile": outputs.canonical_generation_profile,
        "slide_function_outline": outputs.slide_function_outline,
        "concept_graph": outputs.concept_graph,
        "teaching_plan": outputs.teaching_plan,
        "blueprint_preview": outputs.blueprint_preview,
        "authoring_preview": outputs.authoring_preview,
        "design_system": outputs.design_system,
        "deck_constitution": outputs.deck_constitution,
        "layout_library": outputs.layout_library,
        "slide_ledger": outputs.slide_ledger,
        "asset_requests": outputs.asset_requests,
    }
    if outputs.proof_unit_registry is not None:
        models["proof_unit_registry"] = outputs.proof_unit_registry
    else:
        stale_proof_unit_registry_path = resolved_output_dir / DEFAULT_STATE_FILENAMES["proof_unit_registry"]
        if stale_proof_unit_registry_path.is_file():
            stale_proof_unit_registry_path.unlink()
    if outputs.proof_module_manifest is not None:
        models["proof_module_manifest"] = outputs.proof_module_manifest
    else:
        stale_proof_module_manifest_path = resolved_output_dir / DEFAULT_STATE_FILENAMES["proof_module_manifest"]
        if stale_proof_module_manifest_path.is_file():
            stale_proof_module_manifest_path.unlink()
    written: dict[str, Path] = {}
    for schema_name, model in models.items():
        if model is None:
            continue
        path = resolved_output_dir / DEFAULT_STATE_FILENAMES[schema_name]
        save_state_file(model, path)
        written[schema_name] = path
    return written

