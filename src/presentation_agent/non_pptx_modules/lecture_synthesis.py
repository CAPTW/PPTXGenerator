"""Concept-driven lecture planning for Gate 2 blueprint synthesis."""

from __future__ import annotations

import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable

from .state_schemas import (
    BlueprintPreview,
    ConceptEdge,
    ConceptEdgeType,
    ConceptGraph,
    ConceptNode,
    ConceptType,
    ContentTier,
    DeckMode,
    LectureFamily,
    ProjectMaterial,
    SlideIntent,
    SlideRole,
    SourceMaterialRef,
    TeachingPlan,
    TeachingPlanSlide,
    VisualType,
)

HEADING_RE = re.compile(r"^(?P<number>\d+(?:\.\d+)+)\s+(?P<title>.+)$")
TITLE_MAX_CHARS = 60
MAIN_STORY_MESSAGE_MAX_CHARS = 118
APPENDIX_MESSAGE_MAX_CHARS = 112
DIRECT_APPENDIX_SOURCE_SUFFIXES = (".png", ".jpg", ".jpeg", ".gif", ".webp", ".svg")


@dataclass(slots=True)
class SourceChunk:
    source_path: str
    source_label: str
    page: int | None
    heading: str | None
    text: str

    @property
    def lowered(self) -> str:
        return self.text.lower()


@dataclass(slots=True)
class ConceptTemplate:
    concept_id: str
    label: str
    section: str
    concept_type: ConceptType
    definition: str
    aliases: tuple[str, ...]
    primary_intent: SlideIntent
    secondary_intent: SlideIntent | None = None
    visual_hint: VisualType | None = None
    importance: float = 1.0
    prerequisites: tuple[str, ...] = ()
    downstream: tuple[str, ...] = ()


@dataclass(slots=True)
class TeachingSeed:
    slide_key: str
    section: str
    deck_mode: DeckMode
    intent: SlideIntent
    title: str
    pedagogical_goal: str
    main_message: str
    concept_ids: list[str] = field(default_factory=list)
    evidence: list[SourceMaterialRef] = field(default_factory=list)
    visual_hint: VisualType | None = None
    priority: int = 1
    importance: float = 0.0
    secondary: bool = False


@dataclass(slots=True)
class LecturePlanningArtifacts:
    concept_graph: ConceptGraph
    teaching_plan: TeachingPlan
    blueprint_preview: BlueprintPreview
    specs: list[dict[str, object]]
    clustering_decisions: list[str]


FAMILY_TERMS: dict[LectureFamily, dict[str, int]] = {
    LectureFamily.OPTIMIZATION_METHOD: {
        "optimization": 3,
        "objective": 2,
        "constraint": 2,
        "convex": 2,
        "optimality": 3,
        "gradient": 2,
        "hessian": 2,
        "lagrange": 2,
        "kkt": 3,
        "line search": 2,
        "newton": 2,
        "problem formulation": 3,
        "method selection": 3,
    },
    LectureFamily.CONCEPT_TO_ALGORITHM_MAPPING: {
        "gene": 2,
        "allele": 2,
        "genotype": 3,
        "phenotype": 3,
        "natural selection": 3,
        "genetic algorithm": 4,
        "solution string": 3,
        "encoding": 3,
        "fitness": 3,
        "selection operator": 3,
        "crossover": 3,
        "mutation": 3,
        "schema": 2,
        "building block": 2,
        "bridge": 3,
        "correspondence": 3,
        "analog": 2,
        "maps to": 3,
    },
    LectureFamily.MECHANISM_PROCESS: {
        "mechanism": 3,
        "process": 2,
        "cycle": 2,
        "generation": 2,
        "step": 1,
        "operator": 2,
        "workflow": 1,
    },
    LectureFamily.APPLICATION_COMPARISON: {
        "application": 3,
        "compare": 2,
        "comparison": 2,
        "trade-off": 2,
        "tradeoff": 2,
        "limitation": 3,
        "criticism": 2,
        "case": 1,
        "survey": 1,
    },
}

SUSPICIOUS_OPTIMIZATION_FRAMING = (
    "problem formulation",
    "optimality conditions",
    "method selection",
    "structurally appropriate",
    "graduate optimization",
    "convexity",
    "kkt",
    "line search",
)


def _normalize_text(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def _contains_hangul(text: str) -> bool:
    return any("\uac00" <= character <= "\ud7a3" for character in text)


def _shorten_line(value: str, limit: int = 140) -> str:
    cleaned = _normalize_text(value)
    if len(cleaned) <= limit:
        return cleaned
    trimmed = cleaned[: limit - 1].rstrip(" ,;:/-")
    return f"{trimmed}..."


def _compress_claim(value: str, limit: int) -> str:
    cleaned = _normalize_text(value)
    replacements = (
        (" rather than ", " instead of "),
        (" is useful because ", " helps because "),
        ("the encoded candidate is evaluated through the task environment", "the encoded candidate is tested in the task environment"),
        ("the role split between", "the role split across"),
        ("generic source summary", "source summary"),
        ("one point estimate", "one estimate"),
    )
    for before, after in replacements:
        cleaned = cleaned.replace(before, after)
    return _shorten_line(cleaned, limit)


def _normalize_title_style(value: str, limit: int = TITLE_MAX_CHARS) -> str:
    cleaned = _normalize_text(value)
    cleaned = re.sub(r"^How the lecture moves from\s+", "From ", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"^Why use the analogy at all\??$", "Why the analogy helps", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"^Worked example\s*:\s*", "Worked example ", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"^Appendix support\s*:\s*", "Evidence summary for ", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"^Deepening\s+", "", cleaned, flags=re.IGNORECASE)
    cleaned = cleaned.rstrip(" .?!:;")
    return _shorten_line(cleaned, limit)


def _message_limit(deck_mode: DeckMode) -> int:
    return APPENDIX_MESSAGE_MAX_CHARS if deck_mode == DeckMode.APPENDIX else MAIN_STORY_MESSAGE_MAX_CHARS


def _slugify(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-") or "item"


def _dedupe(items: list[str]) -> list[str]:
    seen: set[str] = set()
    ordered: list[str] = []
    for item in items:
        cleaned = item.strip()
        if cleaned and cleaned not in seen:
            seen.add(cleaned)
            ordered.append(cleaned)
    return ordered


def _resolve_material_path(path_text: str | None) -> Path | None:
    if not path_text:
        return None
    candidate = Path(path_text)
    if candidate.is_absolute():
        return candidate if candidate.exists() else None
    resolved = candidate.resolve()
    return resolved if resolved.exists() else None


def _material_path_by_suffix(materials: list[ProjectMaterial], suffixes: tuple[str, ...]) -> Path | None:
    for material in materials:
        if not material.path:
            continue
        if Path(material.path).suffix.lower() not in suffixes:
            continue
        resolved = _resolve_material_path(material.path)
        if resolved is not None:
            return resolved
    return None


def _pdf_chunks(pdf_path: Path | None) -> list[SourceChunk]:
    if pdf_path is None:
        return []
    try:
        import fitz
    except ImportError:
        return []
    chunks: list[SourceChunk] = []
    with fitz.open(pdf_path) as pdf:
        for page_number, page in enumerate(pdf, start=1):
            text = _normalize_text(page.get_text("text"))
            if not text:
                continue
            heading = None
            match = re.search(r"(4\.\d+\s+[^\n]+|1\.\d+\s+[^\n]+)", text)
            if match:
                heading = _normalize_text(match.group(1))
            chunks.append(
                SourceChunk(
                    source_path=str(pdf_path),
                    source_label=pdf_path.name,
                    page=page_number,
                    heading=heading,
                    text=text,
                )
            )
    return chunks


def _docx_chunks(docx_path: Path | None) -> list[SourceChunk]:
    if docx_path is None:
        return []
    try:
        from docx import Document
    except ImportError:
        return []
    document = Document(str(docx_path))
    chunks: list[SourceChunk] = []
    heading: str | None = None
    buffer: list[str] = []

    def flush() -> None:
        nonlocal buffer
        text = _normalize_text(" ".join(buffer))
        if text:
            chunks.append(
                SourceChunk(
                    source_path=str(docx_path),
                    source_label=docx_path.name,
                    page=None,
                    heading=heading,
                    text=text,
                )
            )
        buffer = []

    for paragraph in document.paragraphs:
        text = _normalize_text(paragraph.text)
        if not text:
            continue
        style_name = _normalize_text(getattr(getattr(paragraph, "style", None), "name", "")).lower()
        is_heading = style_name.startswith("heading") or HEADING_RE.match(text) is not None
        if is_heading:
            flush()
            heading = HEADING_RE.sub(r"\g<title>", text)
            buffer = [heading]
            continue
        buffer.append(text)
        if len(" ".join(buffer)) > 1200:
            flush()
    flush()
    return chunks


def _collect_source_chunks(materials: list[ProjectMaterial]) -> list[SourceChunk]:
    docx_path = _material_path_by_suffix(materials, (".docx",))
    pdf_path = _material_path_by_suffix(materials, (".pdf",))
    chunks = [*_docx_chunks(docx_path), *_pdf_chunks(pdf_path)]
    deduped: dict[tuple[str, int | None, str], SourceChunk] = {}
    for chunk in chunks:
        key = (chunk.source_path, chunk.page, _shorten_line(chunk.text, 240))
        deduped.setdefault(key, chunk)
    return list(deduped.values())


def _alias_in_text(alias: str, text: str) -> bool:
    lowered = text.lower()
    if all(character.isascii() and (character.isalnum() or character in {" ", "-", "_"}) for character in alias):
        pattern = r"\b" + re.escape(alias.lower()).replace(r"\ ", r"\s+") + r"\b"
        return re.search(pattern, lowered) is not None
    return alias.lower() in lowered


def _match_snippet(alias: str, chunk: SourceChunk) -> str:
    lowered = chunk.lowered
    probe = alias.lower()
    index = lowered.find(probe)
    if index < 0:
        return _shorten_line(chunk.text, 180)
    start = max(0, index - 70)
    end = min(len(chunk.text), index + max(len(alias) + 110, 160))
    return _shorten_line(chunk.text[start:end], 180)


def _source_ref(
    *,
    concept_id: str,
    chunk: SourceChunk,
    note: str,
) -> SourceMaterialRef:
    return SourceMaterialRef(
        source_id=f"{concept_id}-{_slugify(chunk.source_label)}-{chunk.page or 'doc'}",
        label=chunk.source_label,
        path=chunk.source_path,
        page=chunk.page,
        notes=note,
    )


def _family_score(
    chunks: list[SourceChunk],
    family: LectureFamily,
) -> tuple[int, Counter[str]]:
    hits: Counter[str] = Counter()
    score = 0
    for chunk in chunks:
        lowered = chunk.lowered
        for token, weight in FAMILY_TERMS[family].items():
            if token in lowered:
                hits[token] += 1
                score += weight
    if family == LectureFamily.CONCEPT_TO_ALGORITHM_MAPPING:
        biology_hits = sum(hits[token] for token in ("gene", "allele", "genotype", "phenotype", "natural selection"))
        ga_hits = sum(hits[token] for token in ("genetic algorithm", "encoding", "fitness", "selection operator", "crossover", "mutation"))
        bridge_hits = sum(hits[token] for token in ("bridge", "correspondence", "maps to", "analog"))
        if biology_hits >= 3 and ga_hits >= 3:
            score += 16
        elif biology_hits and ga_hits:
            score += 8
        if bridge_hits:
            score += 6
    if family == LectureFamily.OPTIMIZATION_METHOD:
        if hits["optimization"] and (hits["optimality"] or hits["kkt"] or hits["gradient"]):
            score += 8
        if hits["problem formulation"] or hits["method selection"]:
            score += 6
    if family == LectureFamily.MECHANISM_PROCESS and (hits["mechanism"] or hits["cycle"] or hits["generation"]):
        score += 4
    if family == LectureFamily.APPLICATION_COMPARISON and (hits["application"] or hits["limitation"] or hits["comparison"]):
        score += 4
    return score, hits


def _decide_lecture_family(chunks: list[SourceChunk]) -> tuple[LectureFamily, list[str], list[str]]:
    scored: list[tuple[int, LectureFamily, Counter[str]]] = []
    for family in LectureFamily:
        score, hits = _family_score(chunks, family)
        scored.append((score, family, hits))
    scored.sort(key=lambda item: (-item[0], item[1].value))
    selected_score, selected_family, selected_hits = scored[0]
    evidence = [
        f"Selected `{selected_family.value}` with score {selected_score} from source signals {', '.join(token for token, _ in selected_hits.most_common(6)) or 'none'}."
    ]
    for chunk in chunks:
        chunk_text = chunk.lowered
        matched_tokens = [token for token, _ in selected_hits.most_common() if token in chunk_text][:3]
        if matched_tokens:
            location = f"{chunk.source_label} p.{chunk.page}" if chunk.page is not None else chunk.source_label
            evidence.append(f"{location}: evidence for {', '.join(matched_tokens)}.")
            if len(evidence) >= 4:
                break
    rejected = []
    for score, family, hits in scored[1:]:
        rejected.append(
            f"{family.value}: weaker score {score} from {', '.join(token for token, _ in hits.most_common(4)) or 'no strong family terms'}."
        )
    return selected_family, evidence, rejected


INTENT_VISUAL_DEFAULTS: dict[SlideIntent, VisualType] = {
    SlideIntent.ORIENTATION: VisualType.TEXT,
    SlideIntent.ANCHOR_CONCEPT: VisualType.FRAMEWORK,
    SlideIntent.CONCEPT_UNPACKING: VisualType.HIERARCHY,
    SlideIntent.MAPPING_BRIDGE: VisualType.COMPARISON,
    SlideIntent.MECHANISM_WALKTHROUGH: VisualType.PROCESS,
    SlideIntent.WORKED_EXAMPLE: VisualType.DOCUMENT_CROP,
    SlideIntent.COMPARISON_TRADEOFF: VisualType.COMPARISON,
    SlideIntent.MISCONCEPTION_PITFALL: VisualType.COMPARISON,
    SlideIntent.APPLICATION_VIGNETTE: VisualType.DOCUMENT_CROP,
    SlideIntent.SUMMARY_INTEGRATION: VisualType.FRAMEWORK,
    SlideIntent.APPENDIX_EVIDENCE_SUPPORT: VisualType.TEXT,
}

EVOLUTIONARY_MAPPING_TEMPLATES: tuple[ConceptTemplate, ...] = (
    ConceptTemplate(
        concept_id="genes",
        label="Genes",
        section="Genes as units",
        concept_type=ConceptType.BIOLOGICAL_CONCEPT,
        definition="Genes act as the information-bearing units whose values can vary across candidates.",
        aliases=("gene", "genes"),
        primary_intent=SlideIntent.ANCHOR_CONCEPT,
        secondary_intent=SlideIntent.MAPPING_BRIDGE,
        visual_hint=VisualType.HIERARCHY,
        importance=1.5,
        downstream=("representation-encoding",),
    ),
    ConceptTemplate(
        concept_id="alleles",
        label="Alleles",
        section="Alleles and variants",
        concept_type=ConceptType.BIOLOGICAL_CONCEPT,
        definition="Alleles are alternate gene values and motivate discrete encoding choices in search.",
        aliases=("allele", "alleles"),
        primary_intent=SlideIntent.CONCEPT_UNPACKING,
        secondary_intent=SlideIntent.MAPPING_BRIDGE,
        visual_hint=VisualType.COMPARISON,
        importance=1.1,
        prerequisites=("genes",),
        downstream=("representation-encoding",),
    ),
    ConceptTemplate(
        concept_id="genotype",
        label="Genotype",
        section="Encoded traits",
        concept_type=ConceptType.BIOLOGICAL_CONCEPT,
        definition="Genotype names the encoded internal description before evaluation through the environment.",
        aliases=("genotype", "genotypes"),
        primary_intent=SlideIntent.ANCHOR_CONCEPT,
        secondary_intent=SlideIntent.MAPPING_BRIDGE,
        visual_hint=VisualType.FRAMEWORK,
        importance=1.6,
        prerequisites=("genes",),
        downstream=("fitness-bridge", "representation-encoding"),
    ),
    ConceptTemplate(
        concept_id="phenotype",
        label="Phenotype",
        section="Encoded traits",
        concept_type=ConceptType.BIOLOGICAL_CONCEPT,
        definition="Phenotype names the expressed behavior that is actually judged by the environment.",
        aliases=("phenotype", "phenotypes"),
        primary_intent=SlideIntent.ANCHOR_CONCEPT,
        secondary_intent=SlideIntent.MAPPING_BRIDGE,
        visual_hint=VisualType.FRAMEWORK,
        importance=1.6,
        prerequisites=("genotype",),
        downstream=("fitness-bridge", "fitness"),
    ),
    ConceptTemplate(
        concept_id="variation",
        label="Variation",
        section="Selection and variation",
        concept_type=ConceptType.BIOLOGICAL_CONCEPT,
        definition="Variation keeps populations diverse enough for selection to work on different candidates.",
        aliases=("variation", "diversity"),
        primary_intent=SlideIntent.CONCEPT_UNPACKING,
        secondary_intent=SlideIntent.COMPARISON_TRADEOFF,
        visual_hint=VisualType.COMPARISON,
        importance=1.2,
        downstream=("selection-operator", "mutation-operator"),
    ),
    ConceptTemplate(
        concept_id="natural-selection",
        label="Natural selection",
        section="Selection and variation",
        concept_type=ConceptType.BIOLOGICAL_CONCEPT,
        definition="Natural selection changes population composition by reusing fitter variants more often.",
        aliases=("natural selection",),
        primary_intent=SlideIntent.ANCHOR_CONCEPT,
        secondary_intent=SlideIntent.MAPPING_BRIDGE,
        visual_hint=VisualType.PROCESS,
        importance=1.7,
        prerequisites=("variation",),
        downstream=("selection-operator", "generation-cycle"),
    ),
    ConceptTemplate(
        concept_id="genetic-algorithm",
        label="Genetic algorithm",
        section="GA overview",
        concept_type=ConceptType.GA_CONCEPT,
        definition="A genetic algorithm searches by evolving a population of encoded candidates over repeated generations.",
        aliases=("genetic algorithm", "genetic algorithms", "ga"),
        primary_intent=SlideIntent.ANCHOR_CONCEPT,
        secondary_intent=SlideIntent.MECHANISM_WALKTHROUGH,
        visual_hint=VisualType.PROCESS,
        importance=2.0,
        downstream=("generation-cycle", "applications", "limitations"),
    ),
    ConceptTemplate(
        concept_id="representation-encoding",
        label="Representation and encoding",
        section="Representation bridge",
        concept_type=ConceptType.MAPPING_CONCEPT,
        definition="Encoding decides how candidate solutions are written so the operators can act on them reliably.",
        aliases=("encoding", "representation", "solution string"),
        primary_intent=SlideIntent.MAPPING_BRIDGE,
        secondary_intent=SlideIntent.CONCEPT_UNPACKING,
        visual_hint=VisualType.COMPARISON,
        importance=1.8,
        prerequisites=("genes", "genotype"),
        downstream=("fitness", "generation-cycle"),
    ),
    ConceptTemplate(
        concept_id="fitness",
        label="Fitness",
        section="Evaluation bridge",
        concept_type=ConceptType.GA_CONCEPT,
        definition="Fitness scores how well a candidate performs so selection can prefer better variants.",
        aliases=("fitness", "fitness function"),
        primary_intent=SlideIntent.ANCHOR_CONCEPT,
        secondary_intent=SlideIntent.MAPPING_BRIDGE,
        visual_hint=VisualType.CHART,
        importance=1.9,
        prerequisites=("representation-encoding", "phenotype"),
        downstream=("selection-operator",),
    ),
    ConceptTemplate(
        concept_id="selection-operator",
        label="Selection operator",
        section="Selection operator",
        concept_type=ConceptType.OPERATOR_MECHANISM,
        definition="Selection operators bias reproduction toward fitter candidates without guaranteeing a single winner.",
        aliases=("selection operator", "selection", "tournament selection", "roulette"),
        primary_intent=SlideIntent.MECHANISM_WALKTHROUGH,
        secondary_intent=SlideIntent.MAPPING_BRIDGE,
        visual_hint=VisualType.PROCESS,
        importance=1.8,
        prerequisites=("fitness", "natural-selection"),
        downstream=("crossover-operator", "mutation-operator"),
    ),
    ConceptTemplate(
        concept_id="crossover-operator",
        label="Crossover",
        section="Crossover and building blocks",
        concept_type=ConceptType.OPERATOR_MECHANISM,
        definition="Crossover recombines useful partial structures from parent solutions into new offspring.",
        aliases=("crossover", "cross-over"),
        primary_intent=SlideIntent.MECHANISM_WALKTHROUGH,
        secondary_intent=SlideIntent.COMPARISON_TRADEOFF,
        visual_hint=VisualType.PROCESS,
        importance=1.8,
        prerequisites=("selection-operator",),
        downstream=("schema-intuition", "generation-cycle"),
    ),
    ConceptTemplate(
        concept_id="mutation-operator",
        label="Mutation operator",
        section="Mutation and diversity",
        concept_type=ConceptType.OPERATOR_MECHANISM,
        definition="Mutation perturbs solutions so the search can escape brittle local patterns and recover diversity.",
        aliases=("mutation operator", "mutation"),
        primary_intent=SlideIntent.MECHANISM_WALKTHROUGH,
        secondary_intent=SlideIntent.COMPARISON_TRADEOFF,
        visual_hint=VisualType.PROCESS,
        importance=1.8,
        prerequisites=("selection-operator", "variation"),
        downstream=("generation-cycle", "exploration-exploitation"),
    ),
    ConceptTemplate(
        concept_id="generation-cycle",
        label="One generation cycle",
        section="Generation cycle",
        concept_type=ConceptType.OPERATOR_MECHANISM,
        definition="A full generation cycles through evaluation, selection, crossover, mutation, and replacement.",
        aliases=("generation", "generation cycle", "one generation"),
        primary_intent=SlideIntent.MECHANISM_WALKTHROUGH,
        secondary_intent=SlideIntent.WORKED_EXAMPLE,
        visual_hint=VisualType.PROCESS,
        importance=2.0,
        prerequisites=("fitness", "selection-operator", "crossover-operator", "mutation-operator"),
        downstream=("applications", "limitations"),
    ),
    ConceptTemplate(
        concept_id="biology-ga-correspondence",
        label="Biology to GA correspondence",
        section="Representation bridge",
        concept_type=ConceptType.MAPPING_CONCEPT,
        definition="The lecture should make the biology to algorithm mapping explicit before diving into operators.",
        aliases=("correspondence", "maps to", "analogy", "bridge"),
        primary_intent=SlideIntent.MAPPING_BRIDGE,
        secondary_intent=SlideIntent.SUMMARY_INTEGRATION,
        visual_hint=VisualType.TABLE,
        importance=2.0,
        prerequisites=("genes", "genotype", "phenotype", "genetic-algorithm"),
        downstream=("representation-encoding", "fitness", "selection-operator"),
    ),
    ConceptTemplate(
        concept_id="fitness-bridge",
        label="Phenotype to fitness bridge",
        section="Evaluation bridge",
        concept_type=ConceptType.MAPPING_CONCEPT,
        definition="Phenotype explains why encoded candidates still need an external evaluation environment.",
        aliases=("phenotype", "fitness"),
        primary_intent=SlideIntent.MAPPING_BRIDGE,
        secondary_intent=SlideIntent.WORKED_EXAMPLE,
        visual_hint=VisualType.COMPARISON,
        importance=1.4,
        prerequisites=("phenotype", "fitness"),
        downstream=("selection-operator",),
    ),
    ConceptTemplate(
        concept_id="schema-intuition",
        label="Schema or building-block intuition",
        section="Crossover and building blocks",
        concept_type=ConceptType.GA_CONCEPT,
        definition="Schema intuition explains why crossover can preserve and recombine short useful patterns.",
        aliases=("schema", "building block", "building-block"),
        primary_intent=SlideIntent.CONCEPT_UNPACKING,
        secondary_intent=SlideIntent.MAPPING_BRIDGE,
        visual_hint=VisualType.FRAMEWORK,
        importance=1.4,
        prerequisites=("crossover-operator",),
        downstream=("applications", "limitations"),
    ),
    ConceptTemplate(
        concept_id="exploration-exploitation",
        label="Exploration and exploitation tradeoff",
        section="Tradeoffs and applications",
        concept_type=ConceptType.PARAMETER_TRADEOFF,
        definition="Genetic algorithms work only when diversity and selection pressure are balanced deliberately.",
        aliases=("exploration", "exploitation", "selection pressure"),
        primary_intent=SlideIntent.COMPARISON_TRADEOFF,
        secondary_intent=SlideIntent.MISCONCEPTION_PITFALL,
        visual_hint=VisualType.COMPARISON,
        importance=1.5,
        prerequisites=("selection-operator", "mutation-operator"),
        downstream=("limitations",),
    ),
    ConceptTemplate(
        concept_id="applications",
        label="Applications",
        section="Tradeoffs and applications",
        concept_type=ConceptType.APPLICATION,
        definition="Applications show where global search over rugged or poorly structured spaces can justify GA overhead.",
        aliases=("application", "applications"),
        primary_intent=SlideIntent.APPLICATION_VIGNETTE,
        secondary_intent=SlideIntent.COMPARISON_TRADEOFF,
        visual_hint=VisualType.CHART,
        importance=1.3,
        prerequisites=("generation-cycle",),
    ),
    ConceptTemplate(
        concept_id="limitations",
        label="Limitations",
        section="Tradeoffs and applications",
        concept_type=ConceptType.LIMITATION,
        definition="Limitations include premature convergence, expensive fitness evaluation, and representation mismatch.",
        aliases=("limitation", "limitations", "premature convergence"),
        primary_intent=SlideIntent.MISCONCEPTION_PITFALL,
        secondary_intent=SlideIntent.COMPARISON_TRADEOFF,
        visual_hint=VisualType.COMPARISON,
        importance=1.5,
        prerequisites=("generation-cycle", "exploration-exploitation"),
    ),
)

OPTIMIZATION_TEMPLATES: tuple[ConceptTemplate, ...] = (
    ConceptTemplate(
        concept_id="problem-formulation",
        label="Problem formulation",
        section="Problem structure",
        concept_type=ConceptType.OPTIMIZATION_CONCEPT,
        definition="Problem formulation defines variables, objective, and constraints before method choice becomes meaningful.",
        aliases=("problem formulation", "design variables"),
        primary_intent=SlideIntent.ANCHOR_CONCEPT,
        secondary_intent=SlideIntent.CONCEPT_UNPACKING,
        visual_hint=VisualType.FRAMEWORK,
        importance=1.9,
        downstream=("optimality-conditions", "method-selection"),
    ),
    ConceptTemplate(
        concept_id="objective-function",
        label="Objective function",
        section="Problem structure",
        concept_type=ConceptType.MATHEMATICAL_CONCEPT,
        definition="The objective function states what the search is trying to minimize or maximize.",
        aliases=("objective function", "objective"),
        primary_intent=SlideIntent.CONCEPT_UNPACKING,
        secondary_intent=SlideIntent.WORKED_EXAMPLE,
        visual_hint=VisualType.CHART,
        importance=1.3,
        prerequisites=("problem-formulation",),
    ),
    ConceptTemplate(
        concept_id="constraint-set",
        label="Constraint set",
        section="Problem structure",
        concept_type=ConceptType.MATHEMATICAL_CONCEPT,
        definition="Constraints define the feasible region and change both geometry and method choice.",
        aliases=("constraint", "constraints", "feasible region"),
        primary_intent=SlideIntent.CONCEPT_UNPACKING,
        secondary_intent=SlideIntent.COMPARISON_TRADEOFF,
        visual_hint=VisualType.FRAMEWORK,
        importance=1.4,
        prerequisites=("problem-formulation",),
        downstream=("lagrange-kkt", "method-selection"),
    ),
    ConceptTemplate(
        concept_id="optimality-conditions",
        label="Optimality conditions",
        section="Optimality and geometry",
        concept_type=ConceptType.MATHEMATICAL_CONCEPT,
        definition="Optimality conditions describe the local tests that separate promising stationary points from the wrong ones.",
        aliases=("optimality conditions", "stationary"),
        primary_intent=SlideIntent.ANCHOR_CONCEPT,
        secondary_intent=SlideIntent.CONCEPT_UNPACKING,
        visual_hint=VisualType.FRAMEWORK,
        importance=1.8,
        prerequisites=("problem-formulation",),
        downstream=("gradient-hessian", "lagrange-kkt"),
    ),
    ConceptTemplate(
        concept_id="convexity",
        label="Convexity",
        section="Optimality and geometry",
        concept_type=ConceptType.MATHEMATICAL_CONCEPT,
        definition="Convexity upgrades local guarantees and explains when first-order checks are globally meaningful.",
        aliases=("convex", "convexity"),
        primary_intent=SlideIntent.CONCEPT_UNPACKING,
        secondary_intent=SlideIntent.COMPARISON_TRADEOFF,
        visual_hint=VisualType.COMPARISON,
        importance=1.4,
        prerequisites=("optimality-conditions",),
        downstream=("method-selection",),
    ),
    ConceptTemplate(
        concept_id="gradient-hessian",
        label="Gradients and Hessians",
        section="Optimality and geometry",
        concept_type=ConceptType.MATHEMATICAL_CONCEPT,
        definition="Gradient and Hessian information explain direction, curvature, and second-order classification.",
        aliases=("gradient", "hessian"),
        primary_intent=SlideIntent.CONCEPT_UNPACKING,
        secondary_intent=SlideIntent.WORKED_EXAMPLE,
        visual_hint=VisualType.FRAMEWORK,
        importance=1.5,
        prerequisites=("optimality-conditions",),
        downstream=("newton-method",),
    ),
    ConceptTemplate(
        concept_id="lagrange-kkt",
        label="Lagrange multipliers and KKT",
        section="Optimality and geometry",
        concept_type=ConceptType.MATHEMATICAL_CONCEPT,
        definition="Lagrange multipliers and KKT conditions explain constrained optimality and active constraints.",
        aliases=("lagrange", "kkt"),
        primary_intent=SlideIntent.CONCEPT_UNPACKING,
        secondary_intent=SlideIntent.WORKED_EXAMPLE,
        visual_hint=VisualType.FRAMEWORK,
        importance=1.6,
        prerequisites=("constraint-set", "optimality-conditions"),
        downstream=("constrained-methods", "method-selection"),
    ),
    ConceptTemplate(
        concept_id="line-search",
        label="Line search and descent logic",
        section="Method families",
        concept_type=ConceptType.OPERATOR_MECHANISM,
        definition="Line search explains how step-size control stabilizes iterative first-order methods.",
        aliases=("line search", "steepest descent", "descent"),
        primary_intent=SlideIntent.MECHANISM_WALKTHROUGH,
        secondary_intent=SlideIntent.WORKED_EXAMPLE,
        visual_hint=VisualType.PROCESS,
        importance=1.3,
        prerequisites=("gradient-hessian",),
        downstream=("method-selection",),
    ),
    ConceptTemplate(
        concept_id="newton-method",
        label="Newton and quasi-Newton methods",
        section="Method families",
        concept_type=ConceptType.OPERATOR_MECHANISM,
        definition="Newton-style methods use curvature information to accelerate convergence when the local model is trustworthy.",
        aliases=("newton", "quasi-newton"),
        primary_intent=SlideIntent.MECHANISM_WALKTHROUGH,
        secondary_intent=SlideIntent.COMPARISON_TRADEOFF,
        visual_hint=VisualType.PROCESS,
        importance=1.5,
        prerequisites=("gradient-hessian",),
        downstream=("method-selection",),
    ),
    ConceptTemplate(
        concept_id="constrained-methods",
        label="Penalty, barrier, and SQP methods",
        section="Method families",
        concept_type=ConceptType.OPERATOR_MECHANISM,
        definition="Constrained methods restructure the problem so constraints remain visible during search.",
        aliases=("penalty", "barrier", "sqp"),
        primary_intent=SlideIntent.MECHANISM_WALKTHROUGH,
        secondary_intent=SlideIntent.COMPARISON_TRADEOFF,
        visual_hint=VisualType.PROCESS,
        importance=1.5,
        prerequisites=("lagrange-kkt",),
        downstream=("method-selection",),
    ),
    ConceptTemplate(
        concept_id="method-selection",
        label="Method selection",
        section="Comparison and fit",
        concept_type=ConceptType.PARAMETER_TRADEOFF,
        definition="Method selection follows from structure, derivatives, and constraints rather than from a generic preference order.",
        aliases=("method selection", "which method"),
        primary_intent=SlideIntent.COMPARISON_TRADEOFF,
        secondary_intent=SlideIntent.SUMMARY_INTEGRATION,
        visual_hint=VisualType.COMPARISON,
        importance=1.8,
        prerequisites=("problem-formulation", "convexity", "constrained-methods"),
        downstream=("applications", "limitations"),
    ),
    ConceptTemplate(
        concept_id="applications",
        label="Applications",
        section="Applications and limits",
        concept_type=ConceptType.APPLICATION,
        definition="Applications ground abstract method families in real design, fitting, and control problems.",
        aliases=("application", "applications"),
        primary_intent=SlideIntent.APPLICATION_VIGNETTE,
        secondary_intent=SlideIntent.WORKED_EXAMPLE,
        visual_hint=VisualType.CHART,
        importance=1.2,
        prerequisites=("method-selection",),
    ),
    ConceptTemplate(
        concept_id="limitations",
        label="Limitations and failure modes",
        section="Applications and limits",
        concept_type=ConceptType.LIMITATION,
        definition="Failure modes include nonconvexity, poor conditioning, wrong derivatives, and brittle constraint handling.",
        aliases=("limitation", "limitations", "ill-conditioned"),
        primary_intent=SlideIntent.MISCONCEPTION_PITFALL,
        secondary_intent=SlideIntent.COMPARISON_TRADEOFF,
        visual_hint=VisualType.COMPARISON,
        importance=1.4,
        prerequisites=("method-selection",),
    ),
)

GENERIC_FAMILY_TEMPLATES: dict[LectureFamily, tuple[ConceptTemplate, ...]] = {
    LectureFamily.MECHANISM_PROCESS: (
        ConceptTemplate(
            concept_id="process-cycle",
            label="Process cycle",
            section="Mechanism",
            concept_type=ConceptType.OPERATOR_MECHANISM,
            definition="A mechanism lecture should show the step sequence and what changes after each step.",
            aliases=("process", "cycle", "step", "workflow"),
            primary_intent=SlideIntent.MECHANISM_WALKTHROUGH,
            secondary_intent=SlideIntent.CONCEPT_UNPACKING,
            visual_hint=VisualType.PROCESS,
            importance=1.2,
        ),
        ConceptTemplate(
            concept_id="failure-mode",
            label="Failure mode",
            section="Applications and limits",
            concept_type=ConceptType.LIMITATION,
            definition="Mechanism lectures should make the main failure mode visible before closing.",
            aliases=("failure", "limitation", "pitfall"),
            primary_intent=SlideIntent.MISCONCEPTION_PITFALL,
            secondary_intent=SlideIntent.COMPARISON_TRADEOFF,
            visual_hint=VisualType.COMPARISON,
            importance=1.0,
        ),
    ),
    LectureFamily.APPLICATION_COMPARISON: (
        ConceptTemplate(
            concept_id="use-cases",
            label="Use cases",
            section="Applications",
            concept_type=ConceptType.APPLICATION,
            definition="Application lectures need concrete use cases rather than generic summary slides.",
            aliases=("application", "use case", "case"),
            primary_intent=SlideIntent.APPLICATION_VIGNETTE,
            secondary_intent=SlideIntent.COMPARISON_TRADEOFF,
            visual_hint=VisualType.CHART,
            importance=1.2,
        ),
        ConceptTemplate(
            concept_id="fit-criteria",
            label="Fit criteria",
            section="Comparison and fit",
            concept_type=ConceptType.PARAMETER_TRADEOFF,
            definition="Comparison lectures should explain what criteria separate a good fit from a bad one.",
            aliases=("criteria", "tradeoff", "compare", "comparison"),
            primary_intent=SlideIntent.COMPARISON_TRADEOFF,
            secondary_intent=SlideIntent.SUMMARY_INTEGRATION,
            visual_hint=VisualType.COMPARISON,
            importance=1.1,
        ),
    ),
}

CONCEPT_SECTION_LOOKUP: dict[str, str] = {
    template.concept_id: template.section
    for template in (
        *EVOLUTIONARY_MAPPING_TEMPLATES,
        *OPTIMIZATION_TEMPLATES,
        *GENERIC_FAMILY_TEMPLATES[LectureFamily.MECHANISM_PROCESS],
        *GENERIC_FAMILY_TEMPLATES[LectureFamily.APPLICATION_COMPARISON],
    )
}
FAMILY_CONCEPT_SECTION_LOOKUP: dict[LectureFamily, dict[str, str]] = {
    LectureFamily.CONCEPT_TO_ALGORITHM_MAPPING: {template.concept_id: template.section for template in EVOLUTIONARY_MAPPING_TEMPLATES},
    LectureFamily.OPTIMIZATION_METHOD: {template.concept_id: template.section for template in OPTIMIZATION_TEMPLATES},
    LectureFamily.MECHANISM_PROCESS: {
        template.concept_id: template.section for template in GENERIC_FAMILY_TEMPLATES[LectureFamily.MECHANISM_PROCESS]
    },
    LectureFamily.APPLICATION_COMPARISON: {
        template.concept_id: template.section for template in GENERIC_FAMILY_TEMPLATES[LectureFamily.APPLICATION_COMPARISON]
    },
}

CONCEPT_PRIMARY_INTENT_LOOKUP: dict[str, SlideIntent] = {
    template.concept_id: template.primary_intent
    for template in (
        *EVOLUTIONARY_MAPPING_TEMPLATES,
        *OPTIMIZATION_TEMPLATES,
        *GENERIC_FAMILY_TEMPLATES[LectureFamily.MECHANISM_PROCESS],
        *GENERIC_FAMILY_TEMPLATES[LectureFamily.APPLICATION_COMPARISON],
    )
}

CONCEPT_SECONDARY_INTENT_LOOKUP: dict[str, SlideIntent | None] = {
    template.concept_id: template.secondary_intent
    for template in (
        *EVOLUTIONARY_MAPPING_TEMPLATES,
        *OPTIMIZATION_TEMPLATES,
        *GENERIC_FAMILY_TEMPLATES[LectureFamily.MECHANISM_PROCESS],
        *GENERIC_FAMILY_TEMPLATES[LectureFamily.APPLICATION_COMPARISON],
    )
}

CONCEPT_VISUAL_LOOKUP: dict[str, VisualType | None] = {
    template.concept_id: template.visual_hint
    for template in (
        *EVOLUTIONARY_MAPPING_TEMPLATES,
        *OPTIMIZATION_TEMPLATES,
        *GENERIC_FAMILY_TEMPLATES[LectureFamily.MECHANISM_PROCESS],
        *GENERIC_FAMILY_TEMPLATES[LectureFamily.APPLICATION_COMPARISON],
    )
}


def _concept_section(
    concept_id: str,
    fallback: str = "Core concepts",
    lecture_family: LectureFamily | None = None,
) -> str:
    if lecture_family is not None:
        family_lookup = FAMILY_CONCEPT_SECTION_LOOKUP.get(lecture_family, {})
        if concept_id in family_lookup:
            return family_lookup[concept_id]
    return CONCEPT_SECTION_LOOKUP.get(concept_id, fallback)


def _concept_primary_intent(concept_id: str, fallback: SlideIntent = SlideIntent.CONCEPT_UNPACKING) -> SlideIntent:
    return CONCEPT_PRIMARY_INTENT_LOOKUP.get(concept_id, fallback)


def _concept_visual(concept_id: str, fallback: VisualType | None = None) -> VisualType | None:
    return CONCEPT_VISUAL_LOOKUP.get(concept_id, fallback)

FAMILY_EDGE_TEMPLATES: dict[LectureFamily, tuple[tuple[str, str, ConceptEdgeType, str], ...]] = {
    LectureFamily.CONCEPT_TO_ALGORITHM_MAPPING: (
        ("genes", "representation-encoding", ConceptEdgeType.ANALOGY_MAPPING, "Genes motivate encoded units in the representation."),
        ("alleles", "representation-encoding", ConceptEdgeType.ANALOGY_MAPPING, "Alleles motivate alternate encoded values."),
        ("genotype", "representation-encoding", ConceptEdgeType.ANALOGY_MAPPING, "Genotype aligns with the encoded internal string."),
        ("phenotype", "fitness", ConceptEdgeType.ANALOGY_MAPPING, "Phenotype explains why evaluation is external to the encoding."),
        ("natural-selection", "selection-operator", ConceptEdgeType.ANALOGY_MAPPING, "Selection operator mirrors selective retention."),
        ("representation-encoding", "fitness", ConceptEdgeType.PREREQUISITE, "Fitness requires an encoding to evaluate."),
        ("fitness", "selection-operator", ConceptEdgeType.PREREQUISITE, "Selection uses fitness information."),
        ("selection-operator", "crossover-operator", ConceptEdgeType.MECHANISM, "Selected parents feed crossover."),
        ("selection-operator", "mutation-operator", ConceptEdgeType.MECHANISM, "Selected offspring can still mutate."),
        ("crossover-operator", "schema-intuition", ConceptEdgeType.MECHANISM, "Schema intuition depends on recombination preserving short patterns."),
        ("mutation-operator", "exploration-exploitation", ConceptEdgeType.MECHANISM, "Mutation affects exploration pressure."),
        ("generation-cycle", "applications", ConceptEdgeType.APPLICATION_OF, "Applications follow from the generation cycle operating end to end."),
        ("exploration-exploitation", "limitations", ConceptEdgeType.LIMITATION_OF, "Poor balance causes GA failure modes."),
    ),
    LectureFamily.OPTIMIZATION_METHOD: (
        ("problem-formulation", "optimality-conditions", ConceptEdgeType.PREREQUISITE, "Problem formulation comes before optimality tests."),
        ("problem-formulation", "method-selection", ConceptEdgeType.PREREQUISITE, "Method selection depends on structure."),
        ("constraint-set", "lagrange-kkt", ConceptEdgeType.PREREQUISITE, "Constrained optimality builds on constraints."),
        ("gradient-hessian", "newton-method", ConceptEdgeType.PREREQUISITE, "Newton-style methods use curvature."),
        ("lagrange-kkt", "constrained-methods", ConceptEdgeType.PREREQUISITE, "Constrained methods interpret active constraints."),
        ("line-search", "method-selection", ConceptEdgeType.CONTRAST, "Line-search methods differ from second-order methods."),
        ("newton-method", "method-selection", ConceptEdgeType.CONTRAST, "Newton-style methods trade cost for curvature information."),
        ("constrained-methods", "method-selection", ConceptEdgeType.CONTRAST, "Constraint handling changes method fit."),
        ("method-selection", "applications", ConceptEdgeType.APPLICATION_OF, "Applications depend on the method fit criteria."),
        ("limitations", "method-selection", ConceptEdgeType.LIMITATION_OF, "Limitations narrow the method choice."),
    ),
}


def _templates_for_family(chunks: list[SourceChunk], family: LectureFamily) -> tuple[ConceptTemplate, ...]:
    lowered = " ".join(chunk.lowered for chunk in chunks)
    if family == LectureFamily.CONCEPT_TO_ALGORITHM_MAPPING:
        if any(token in lowered for token in ("gene", "genotype", "phenotype", "genetic algorithm")):
            return EVOLUTIONARY_MAPPING_TEMPLATES
        return EVOLUTIONARY_MAPPING_TEMPLATES
    if family == LectureFamily.OPTIMIZATION_METHOD:
        return OPTIMIZATION_TEMPLATES
    return GENERIC_FAMILY_TEMPLATES.get(family, ())


def _matching_aliases(template: ConceptTemplate, chunk: SourceChunk) -> list[str]:
    return [alias for alias in template.aliases if _alias_in_text(alias, chunk.text)]


def _build_concept_graph(
    *,
    deck_title: str,
    chunks: list[SourceChunk],
    lecture_family: LectureFamily,
    lecture_family_evidence: list[str],
    rejected_families: list[str],
) -> ConceptGraph:
    templates = _templates_for_family(chunks, lecture_family)
    nodes: list[ConceptNode] = []
    for template in templates:
        matched_chunks: list[SourceChunk] = []
        matched_aliases: list[str] = []
        for chunk in chunks:
            aliases = _matching_aliases(template, chunk)
            if aliases:
                matched_chunks.append(chunk)
                matched_aliases.append(aliases[0])
        if not matched_chunks:
            continue
        refs: list[SourceMaterialRef] = []
        for alias, chunk in zip(matched_aliases, matched_chunks, strict=False):
            refs.append(_source_ref(concept_id=template.concept_id, chunk=chunk, note=_match_snippet(alias, chunk)))
        importance = round(template.importance + min(len(matched_chunks), 4) * 0.25 + len(template.downstream) * 0.05, 2)
        nodes.append(
            ConceptNode(
                concept_id=template.concept_id,
                label=template.label,
                concept_type=template.concept_type,
                short_teaching_definition=template.definition,
                evidence=refs[:3],
                importance=importance,
                prerequisite_concepts=list(template.prerequisites),
                downstream_concepts=list(template.downstream),
            )
        )
    if not nodes:
        return ConceptGraph(
            deck_title=deck_title,
            lecture_family=lecture_family,
            lecture_family_evidence=lecture_family_evidence,
            rejected_families=rejected_families,
            nodes=[],
            edges=[],
            central_concept_ids=[],
        )
    node_ids = {node.concept_id for node in nodes}
    normalized_nodes = [
        node.model_copy(
            update={
                "prerequisite_concepts": [concept_id for concept_id in node.prerequisite_concepts if concept_id in node_ids],
                "downstream_concepts": [concept_id for concept_id in node.downstream_concepts if concept_id in node_ids],
            }
        )
        for node in nodes
    ]
    edges: list[ConceptEdge] = []
    seen: set[tuple[str, str, ConceptEdgeType]] = set()
    for source_id, target_id, edge_type, rationale in FAMILY_EDGE_TEMPLATES.get(lecture_family, ()):
        if source_id not in node_ids or target_id not in node_ids:
            continue
        key = (source_id, target_id, edge_type)
        if key in seen:
            continue
        seen.add(key)
        edges.append(
            ConceptEdge(
                source_concept_id=source_id,
                target_concept_id=target_id,
                edge_type=edge_type,
                rationale=rationale,
            )
        )
    central_concept_ids = [
        node.concept_id for node in sorted(normalized_nodes, key=lambda item: (-item.importance, item.label.lower()))[:8]
    ]
    return ConceptGraph(
        deck_title=deck_title,
        lecture_family=lecture_family,
        lecture_family_evidence=lecture_family_evidence,
        rejected_families=rejected_families,
        nodes=normalized_nodes,
        edges=edges,
        central_concept_ids=central_concept_ids,
    )


def _node_refs(concept_ids: Iterable[str], nodes_by_id: dict[str, ConceptNode], limit: int = 2) -> list[SourceMaterialRef]:
    refs: list[SourceMaterialRef] = []
    seen: set[str] = set()
    for concept_id in concept_ids:
        node = nodes_by_id.get(concept_id)
        if node is None:
            continue
        for ref in node.evidence:
            if ref.source_id not in seen:
                seen.add(ref.source_id)
                refs.append(ref)
                if len(refs) >= limit:
                    return refs
    return refs


def _available(ids: Iterable[str], nodes_by_id: dict[str, ConceptNode]) -> list[str]:
    return [concept_id for concept_id in ids if concept_id in nodes_by_id]


def _seed(
    *,
    slide_key: str,
    section: str,
    intent: SlideIntent,
    title: str,
    pedagogical_goal: str,
    main_message: str,
    concept_ids: Iterable[str],
    nodes_by_id: dict[str, ConceptNode],
    deck_mode: DeckMode = DeckMode.MAIN_STORY,
    visual_hint: VisualType | None = None,
    priority: int = 1,
    importance: float = 0.0,
    secondary: bool = False,
) -> TeachingSeed:
    resolved_ids = [concept_id for concept_id in concept_ids if concept_id in nodes_by_id]
    return TeachingSeed(
        slide_key=slide_key,
        section=section,
        deck_mode=deck_mode,
        intent=intent,
        title=_normalize_title_style(title),
        pedagogical_goal=_shorten_line(pedagogical_goal, 180),
        main_message=_compress_claim(main_message, _message_limit(deck_mode)),
        concept_ids=resolved_ids,
        evidence=_node_refs(resolved_ids, nodes_by_id),
        visual_hint=visual_hint,
        priority=priority,
        importance=importance,
        secondary=secondary,
    )


def _extension_title(node: ConceptNode, intent: SlideIntent) -> str:
    if intent == SlideIntent.MAPPING_BRIDGE:
        return f"{node.label} as a bridge"
    if intent == SlideIntent.MECHANISM_WALKTHROUGH:
        return f"{node.label} in the cycle"
    if intent == SlideIntent.COMPARISON_TRADEOFF:
        return f"{node.label} tradeoffs"
    if intent == SlideIntent.MISCONCEPTION_PITFALL:
        return f"{node.label} pitfalls"
    if intent == SlideIntent.APPLICATION_VIGNETTE:
        return f"{node.label} in practice"
    return node.label


def _is_direct_appendix_source(ref: SourceMaterialRef) -> bool:
    path = _resolve_material_path(ref.path)
    return path is not None and path.suffix.lower() in DIRECT_APPENDIX_SOURCE_SUFFIXES


def _appendix_support_mode(node: ConceptNode) -> str:
    if any(_is_direct_appendix_source(ref) for ref in node.evidence):
        return "source-linked"
    if len(node.evidence) >= 2:
        return "source-map"
    if node.evidence:
        return "evidence-summary"
    return "reference-only"


def _appendix_support_title(node: ConceptNode, mode: str) -> str:
    if mode == "source-linked":
        return f"Source-linked evidence for {node.label}"
    if mode == "source-map":
        return f"Source map for {node.label}"
    if mode == "reference-only":
        return f"Reference notes for {node.label}"
    return f"Evidence summary for {node.label}"


def _appendix_support_message(node: ConceptNode, mode: str) -> str:
    if mode == "source-linked":
        return f"A direct source excerpt supports {node.label} without reopening the main story."
    if mode == "source-map":
        return f"This appendix maps where {node.label} appears across the source set."
    if mode == "reference-only":
        return f"This appendix records the references behind {node.label} even when no excerpt is required."
    return f"This appendix condenses the source evidence for {node.label} into a traceable support summary."


def _appendix_support_visual(mode: str) -> VisualType:
    if mode == "source-linked":
        return VisualType.DOCUMENT_CROP
    if mode in {"evidence-summary", "source-map"}:
        return VisualType.COMPARISON
    return VisualType.TEXT


def _mapping_family_seeds(deck_title: str, graph: ConceptGraph) -> list[TeachingSeed]:
    nodes_by_id = {node.concept_id: node for node in graph.nodes}
    biology = _available(("genes", "alleles", "genotype", "phenotype", "variation", "natural-selection"), nodes_by_id)
    seeds: list[TeachingSeed] = [
        _seed(
            slide_key="orientation-title",
            section="Orientation",
            intent=SlideIntent.ORIENTATION,
            title="From genetics concepts to genetic algorithms" if biology else deck_title,
            pedagogical_goal="Open with the conceptual teaching question, not with generic optimization framing.",
            main_message="The lecture starts from biological intuition and only then translates it into GA design choices.",
            concept_ids=biology[:3] + _available(("genetic-algorithm",), nodes_by_id),
            nodes_by_id=nodes_by_id,
            visual_hint=VisualType.TEXT,
            priority=10,
        ),
        _seed(
            slide_key="orientation-roadmap",
            section="Orientation",
            intent=SlideIntent.ORIENTATION,
            title="Biology to GA teaching arc",
            pedagogical_goal="Show the pedagogical arc before entering detail.",
            main_message="We move from genes and selection to mapping bridges, operators, examples, and limits.",
            concept_ids=biology[:2] + _available(("biology-ga-correspondence", "generation-cycle", "limitations"), nodes_by_id),
            nodes_by_id=nodes_by_id,
            visual_hint=VisualType.PROCESS,
            priority=10,
        ),
        _seed(
            slide_key="orientation-analogy",
            section="Orientation",
            intent=SlideIntent.ANCHOR_CONCEPT,
            title="Why the analogy helps",
            pedagogical_goal="Frame the analogy as a teaching device that clarifies representation and operator roles.",
            main_message="The biology analogy organizes encoding, evaluation, and variation into one coherent story.",
            concept_ids=_available(("biology-ga-correspondence", "representation-encoding", "fitness", "selection-operator"), nodes_by_id),
            nodes_by_id=nodes_by_id,
            visual_hint=VisualType.FRAMEWORK,
            priority=9,
        ),
    ]
    for concept_id in biology:
        node = nodes_by_id[concept_id]
        seeds.append(
            _seed(
                slide_key=f"bio-{concept_id}",
                section=_concept_section(concept_id, lecture_family=graph.lecture_family),
                intent=_concept_primary_intent(concept_id),
                title=node.label,
                pedagogical_goal=f"Teach `{node.label}` as a prerequisite idea before algorithmic translation.",
                main_message=node.short_teaching_definition,
                concept_ids=[concept_id],
                nodes_by_id=nodes_by_id,
                visual_hint=_concept_visual(concept_id),
                priority=8,
                importance=node.importance,
            )
        )
    for slide_key, title, goal, message, concept_ids, section in (
        (
            "bridge-correspondence",
            "Biology to GA map",
            "Make the source-domain to algorithm mapping explicit before the operator walkthrough begins.",
            "Genes, genotype, phenotype, and selection map to specific GA parts, not a vague optimization metaphor.",
            ("biology-ga-correspondence", "genes", "genotype", "phenotype", "genetic-algorithm"),
            "Representation bridge",
        ),
        (
            "bridge-encoding",
            "Genes and alleles become an encoding",
            "Bridge biological units to the GA representation layer.",
            "The encoding is the algorithmic answer to how genes and alleles are written into a candidate solution.",
            ("genes", "alleles", "representation-encoding"),
            "Representation bridge",
        ),
        (
            "bridge-fitness",
            "Phenotype shows why fitness is external",
            "Teach why evaluation sits outside the encoded string.",
            "Fitness appears when the encoded candidate is tested in the task environment.",
            ("phenotype", "fitness", "fitness-bridge"),
            "Evaluation bridge",
        ),
        (
            "bridge-selection",
            "Natural selection becomes a selection operator",
            "Translate selective pressure into algorithmic sampling rules.",
            "Selection operators approximate survival pressure without guaranteeing that one best candidate dominates forever.",
            ("natural-selection", "selection-operator"),
            "Selection operator",
        ),
        (
            "bridge-operator-split",
            "Operators play different roles",
            "Prevent the lecture from collapsing all GA components into one repeated scaffold.",
            "Encoding sets the search space, fitness scores candidates, and operators create or filter variation.",
            ("representation-encoding", "fitness", "selection-operator", "crossover-operator", "mutation-operator"),
            "Selection operator",
        ),
        (
            "bridge-schema",
            "Why building blocks matter",
            "Connect the analogy to schema intuition only if the source supports it.",
            "Schema intuition explains why recombination can preserve short useful patterns while still exploring new combinations.",
            ("schema-intuition", "crossover-operator"),
            "Crossover and building blocks",
        ),
    ):
        available_ids = _available(concept_ids, nodes_by_id)
        if available_ids:
            seeds.append(
                _seed(
                    slide_key=slide_key,
                    section=section,
                    intent=SlideIntent.MAPPING_BRIDGE,
                    title=title,
                    pedagogical_goal=goal,
                    main_message=message,
                    concept_ids=available_ids,
                    nodes_by_id=nodes_by_id,
                    visual_hint=VisualType.COMPARISON,
                    priority=8,
                )
            )
    for slide_key, intent, title, goal, message, concept_ids, section in (
        (
            "ga-anchor",
            SlideIntent.ANCHOR_CONCEPT,
            "A GA searches with populations",
            "Anchor the lecture in the algorithm before expanding the operator sequence.",
            "A genetic algorithm updates a population of encoded candidates instead of refining one estimate.",
            ("genetic-algorithm", "generation-cycle"),
            "GA overview",
        ),
        (
            "ga-selection",
            SlideIntent.MECHANISM_WALKTHROUGH,
            "Selection pressure changes the candidate pool",
            "Make the selection mechanism visible before discussing the whole cycle.",
            "Selection increases the reuse of fitter candidates but still leaves room for diversity and stochasticity.",
            ("selection-operator", "exploration-exploitation"),
            "Selection operator",
        ),
        (
            "ga-crossover",
            SlideIntent.MECHANISM_WALKTHROUGH,
            "Crossover recombines partial building blocks",
            "Teach what crossover does and what it cannot do by itself.",
            "Crossover is strongest when useful partial structures are already present and encoded compatibly.",
            ("crossover-operator", "schema-intuition"),
            "Crossover and building blocks",
        ),
        (
            "ga-mutation",
            SlideIntent.MECHANISM_WALKTHROUGH,
            "Mutation protects exploration",
            "Explain why mutation matters even when selection and crossover already exist.",
            "Mutation keeps the search from freezing around brittle structures by reintroducing novelty at controlled frequency.",
            ("mutation-operator", "exploration-exploitation"),
            "Mutation and diversity",
        ),
        (
            "ga-cycle",
            SlideIntent.MECHANISM_WALKTHROUGH,
            "One generation cycle",
            "Walk the audience through the end-to-end generational update.",
            "A generation evaluates candidates, selects parents, creates offspring, mutates them, and forms the next population.",
            ("generation-cycle", "fitness", "selection-operator", "crossover-operator", "mutation-operator"),
            "Generation cycle",
        ),
    ):
        available_ids = _available(concept_ids, nodes_by_id)
        if available_ids:
            seeds.append(
                _seed(
                    slide_key=slide_key,
                    section=section,
                    intent=intent,
                    title=title,
                    pedagogical_goal=goal,
                    main_message=message,
                    concept_ids=available_ids,
                    nodes_by_id=nodes_by_id,
                    visual_hint=INTENT_VISUAL_DEFAULTS[intent],
                    priority=8,
                )
            )
    for slide_key, title, goal, message, concept_ids in (
        (
            "example-setup",
            "Worked example with a toy population",
            "Introduce one small example that is simple enough to compute live.",
            "A toy population makes encoding, fitness, and operator roles easy to see.",
            ("representation-encoding", "fitness", "selection-operator"),
        ),
        (
            "example-generation",
            "Run one generation on the toy example",
            "Walk through the generation cycle with visible intermediate states.",
            "One generation is enough to show how selection, crossover, and mutation reshape the population.",
            ("generation-cycle", "selection-operator", "crossover-operator", "mutation-operator"),
        ),
    ):
        available_ids = _available(concept_ids, nodes_by_id)
        if available_ids:
            seeds.append(
                _seed(
                    slide_key=slide_key,
                    section="Worked example",
                    intent=SlideIntent.WORKED_EXAMPLE,
                    title=title,
                    pedagogical_goal=goal,
                    main_message=message,
                    concept_ids=available_ids,
                    nodes_by_id=nodes_by_id,
                    visual_hint=VisualType.PROCESS,
                    priority=8,
                )
            )
    return seeds


def _optimization_family_seeds(deck_title: str, graph: ConceptGraph) -> list[TeachingSeed]:
    nodes_by_id = {node.concept_id: node for node in graph.nodes}
    seeds: list[TeachingSeed] = [
        _seed(
            slide_key="opt-title",
            section="Orientation",
            intent=SlideIntent.ORIENTATION,
            title=deck_title,
            pedagogical_goal="Open with the optimization teaching arc rather than with detached section summaries.",
            main_message="The lecture starts with problem structure, then builds optimality intuition, method families, and fit criteria.",
            concept_ids=_available(("problem-formulation", "optimality-conditions", "method-selection"), nodes_by_id),
            nodes_by_id=nodes_by_id,
            visual_hint=VisualType.TEXT,
            priority=10,
        ),
        _seed(
            slide_key="opt-roadmap",
            section="Orientation",
            intent=SlideIntent.ORIENTATION,
            title="How problem structure controls method choice",
            pedagogical_goal="Frame the central teaching dependency for the whole lecture.",
            main_message="Method choice only makes sense after the audience understands the objective, constraints, geometry, and derivative information.",
            concept_ids=_available(("problem-formulation", "constraint-set", "convexity", "method-selection"), nodes_by_id),
            nodes_by_id=nodes_by_id,
            visual_hint=VisualType.PROCESS,
            priority=10,
        ),
    ]
    for concept_id in _available(
        (
            "problem-formulation",
            "objective-function",
            "constraint-set",
            "optimality-conditions",
            "convexity",
            "gradient-hessian",
            "lagrange-kkt",
            "line-search",
            "newton-method",
            "constrained-methods",
            "method-selection",
            "applications",
            "limitations",
        ),
        nodes_by_id,
    ):
        node = nodes_by_id[concept_id]
        seeds.append(
            _seed(
                slide_key=f"opt-{concept_id}",
                section=_concept_section(concept_id, lecture_family=graph.lecture_family),
                intent=_concept_primary_intent(concept_id),
                title=node.label,
                pedagogical_goal=f"Teach `{node.label}` as a concept dependency inside the optimization story.",
                main_message=node.short_teaching_definition,
                concept_ids=[concept_id, *node.prerequisite_concepts[:1], *node.downstream_concepts[:1]],
                nodes_by_id=nodes_by_id,
                visual_hint=_concept_visual(concept_id),
                priority=8,
                importance=node.importance,
            )
        )
    return seeds


def _generic_family_seeds(deck_title: str, graph: ConceptGraph) -> list[TeachingSeed]:
    nodes_by_id = {node.concept_id: node for node in graph.nodes}
    seeds: list[TeachingSeed] = [
        _seed(
            slide_key="generic-title",
            section="Orientation",
            intent=SlideIntent.ORIENTATION,
            title=deck_title,
            pedagogical_goal="Open by stating the teaching path before source detail.",
            main_message="The lecture is organized as a teaching sequence rather than as a heading-by-heading summary.",
            concept_ids=graph.central_concept_ids[:3],
            nodes_by_id=nodes_by_id,
            visual_hint=VisualType.TEXT,
            priority=10,
        )
    ]
    for node in graph.nodes:
        seeds.append(
            _seed(
                slide_key=f"generic-{node.concept_id}",
                section=_concept_section(node.concept_id, lecture_family=graph.lecture_family),
                intent=_concept_primary_intent(node.concept_id),
                title=node.label,
                pedagogical_goal=f"Teach `{node.label}` in the main story.",
                main_message=node.short_teaching_definition,
                concept_ids=[node.concept_id],
                nodes_by_id=nodes_by_id,
                visual_hint=_concept_visual(node.concept_id),
                priority=7,
                importance=node.importance,
            )
        )
    return seeds


def _initial_main_story_seeds(graph: ConceptGraph) -> list[TeachingSeed]:
    if graph.lecture_family == LectureFamily.CONCEPT_TO_ALGORITHM_MAPPING:
        return _mapping_family_seeds(graph.deck_title, graph)
    if graph.lecture_family == LectureFamily.OPTIMIZATION_METHOD:
        return _optimization_family_seeds(graph.deck_title, graph)
    return _generic_family_seeds(graph.deck_title, graph)


def _eligible_extension_intent(node: ConceptNode, family: LectureFamily) -> SlideIntent:
    if family == LectureFamily.CONCEPT_TO_ALGORITHM_MAPPING:
        if node.concept_type == ConceptType.MAPPING_CONCEPT:
            return SlideIntent.MAPPING_BRIDGE
        if node.concept_type == ConceptType.OPERATOR_MECHANISM:
            return SlideIntent.MECHANISM_WALKTHROUGH
        if node.concept_type == ConceptType.PARAMETER_TRADEOFF:
            return SlideIntent.COMPARISON_TRADEOFF
        if node.concept_type == ConceptType.APPLICATION:
            return SlideIntent.APPLICATION_VIGNETTE
        if node.concept_type == ConceptType.LIMITATION:
            return SlideIntent.MISCONCEPTION_PITFALL
    if family == LectureFamily.OPTIMIZATION_METHOD:
        if node.concept_type == ConceptType.OPERATOR_MECHANISM:
            return SlideIntent.MECHANISM_WALKTHROUGH
        if node.concept_type in {ConceptType.PARAMETER_TRADEOFF, ConceptType.LIMITATION}:
            return SlideIntent.COMPARISON_TRADEOFF
        if node.concept_type == ConceptType.APPLICATION:
            return SlideIntent.APPLICATION_VIGNETTE
    return CONCEPT_SECONDARY_INTENT_LOOKUP.get(node.concept_id) or SlideIntent.CONCEPT_UNPACKING


def _reorder_to_limit_repetition(seeds: list[TeachingSeed]) -> list[TeachingSeed]:
    if not seeds:
        return []
    ordered: list[TeachingSeed] = []
    queue = list(seeds)
    while queue:
        if len(ordered) < 2 or ordered[-1].intent != ordered[-2].intent or queue[0].intent != ordered[-1].intent:
            ordered.append(queue.pop(0))
            continue
        replacement_index = next(
            (index for index, candidate in enumerate(queue[1:], start=1) if candidate.intent != ordered[-1].intent),
            None,
        )
        if replacement_index is None:
            ordered.append(queue.pop(0))
        else:
            ordered.append(queue.pop(replacement_index))
    return ordered


def _select_main_story_seeds(graph: ConceptGraph, target: int) -> list[TeachingSeed]:
    seeds = _reorder_to_limit_repetition(_initial_main_story_seeds(graph))
    nodes_by_id = {node.concept_id: node for node in graph.nodes}
    concept_order = [node for node in sorted(graph.nodes, key=lambda item: (-item.importance, item.label.lower()))]
    extension_index = 0
    while len(seeds) < target and concept_order:
        node = concept_order[extension_index % len(concept_order)]
        intent = _eligible_extension_intent(node, graph.lecture_family)
        seeds.append(
            _seed(
                slide_key=f"extend-{node.concept_id}-{len(seeds) + 1:03d}",
                section=_concept_section(node.concept_id, lecture_family=graph.lecture_family),
                intent=intent,
                title=_extension_title(node, intent),
                pedagogical_goal=f"Use another pedagogical move so `{node.label}` stays in the teaching plan instead of disappearing into summary.",
                main_message=node.short_teaching_definition,
                concept_ids=[node.concept_id, *node.prerequisite_concepts[:1], *node.downstream_concepts[:1]],
                nodes_by_id=nodes_by_id,
                visual_hint=_concept_visual(node.concept_id) or INTENT_VISUAL_DEFAULTS[intent],
                priority=4,
                importance=node.importance,
                secondary=True,
            )
        )
        seeds = _reorder_to_limit_repetition(seeds)
        extension_index += 1
    return seeds[:target]


def _build_appendix_seeds(graph: ConceptGraph, target: int) -> list[TeachingSeed]:
    nodes_by_id = {node.concept_id: node for node in graph.nodes}
    appendix: list[TeachingSeed] = []
    ordered_nodes = sorted(graph.nodes, key=lambda item: (-len(item.evidence), -item.importance, item.label.lower()))
    while len(appendix) < target and ordered_nodes:
        node = ordered_nodes[len(appendix) % len(ordered_nodes)]
        support_mode = _appendix_support_mode(node)
        appendix.append(
            _seed(
                slide_key=f"appendix-{node.concept_id}-{len(appendix) + 1:03d}",
                section="Appendix",
                intent=SlideIntent.APPENDIX_EVIDENCE_SUPPORT,
                title=_appendix_support_title(node, support_mode),
                pedagogical_goal=f"Keep source-backed support for `{node.label}` outside the main teaching arc.",
                main_message=_appendix_support_message(node, support_mode),
                concept_ids=[node.concept_id],
                nodes_by_id=nodes_by_id,
                deck_mode=DeckMode.APPENDIX,
                visual_hint=_appendix_support_visual(support_mode),
            )
        )
    return appendix


def _repetition_stats(slides: list[TeachingPlanSlide]) -> dict[str, int]:
    main_story = [slide for slide in slides if slide.deck_mode == DeckMode.MAIN_STORY]
    counts = Counter(slide.intent.value for slide in main_story)
    max_run = 0
    current_run = 0
    previous: str | None = None
    section_counts = Counter(slide.section for slide in main_story)
    max_section_run = 0
    current_section_run = 0
    previous_section: str | None = None
    section_transitions = 0
    for slide in main_story:
        if slide.intent.value == previous:
            current_run += 1
        else:
            previous = slide.intent.value
            current_run = 1
        max_run = max(max_run, current_run)
        if slide.section == previous_section:
            current_section_run += 1
        else:
            if previous_section is not None:
                section_transitions += 1
            previous_section = slide.section
            current_section_run = 1
        max_section_run = max(max_section_run, current_section_run)
    stats = dict(sorted(counts.items()))
    for section_name, count in sorted(section_counts.items()):
        stats[f"section-{_slugify(section_name)}"] = count
    stats["max-consecutive-same-intent"] = max_run
    stats["max-section-run"] = max_section_run
    stats["section-transitions"] = section_transitions
    return stats


def _flagged_drift_risks(lecture_family: LectureFamily, slides: list[TeachingPlanSlide], repetition_stats: dict[str, int]) -> list[str]:
    risks: list[str] = []
    main_story = [slide for slide in slides if slide.deck_mode == DeckMode.MAIN_STORY]
    early_text = " ".join(f"{slide.title} {slide.main_message}" for slide in main_story[:5]).lower()
    if lecture_family != LectureFamily.OPTIMIZATION_METHOD:
        suspicious = [term for term in SUSPICIOUS_OPTIMIZATION_FRAMING if term in early_text]
        if suspicious:
            risks.append(f"Early main-story slides still contain suspicious optimization framing: {', '.join(suspicious)}.")
        else:
            risks.append("Avoid optimization-family framing unless the selected lecture family and source evidence explicitly support it.")
    if repetition_stats.get("max-consecutive-same-intent", 0) > 2:
        risks.append(
            f"Main story currently reaches {repetition_stats['max-consecutive-same-intent']} consecutive slides with the same intent."
        )
    return _dedupe(risks)


def _teaching_plan(graph: ConceptGraph, main_story_target: int, appendix_target: int) -> TeachingPlan:
    main_story = _select_main_story_seeds(graph, main_story_target)
    appendix = _build_appendix_seeds(graph, appendix_target)
    plan_slides = [
        TeachingPlanSlide(
            slide_key=seed.slide_key,
            section=seed.section,
            deck_mode=seed.deck_mode,
            intent=seed.intent,
            title=seed.title,
            pedagogical_goal=seed.pedagogical_goal,
            main_message=seed.main_message,
            concept_ids=seed.concept_ids,
            evidence=seed.evidence,
            visual_hint=seed.visual_hint,
        )
        for seed in [*main_story, *appendix]
    ]
    repetition_stats = _repetition_stats(plan_slides)
    flagged_drift_risks = _flagged_drift_risks(graph.lecture_family, plan_slides, repetition_stats)
    appendix_plan = [
        f"{slide.title}: {slide.main_message}"
        for slide in plan_slides
        if slide.deck_mode == DeckMode.APPENDIX
    ][:10]
    return TeachingPlan(
        deck_title=graph.deck_title,
        lecture_family=graph.lecture_family,
        lecture_family_evidence=graph.lecture_family_evidence,
        rejected_families=graph.rejected_families,
        central_concept_ids=graph.central_concept_ids,
        slides=plan_slides,
        appendix_plan=appendix_plan,
        flagged_drift_risks=flagged_drift_risks,
        repetition_stats=repetition_stats,
    )


def _preview(graph: ConceptGraph, teaching_plan: TeachingPlan) -> BlueprintPreview:
    nodes_by_id = {node.concept_id: node for node in graph.nodes}
    main_story = [slide for slide in teaching_plan.slides if slide.deck_mode == DeckMode.MAIN_STORY]
    appendix = [slide for slide in teaching_plan.slides if slide.deck_mode == DeckMode.APPENDIX]
    central_labels = [nodes_by_id[concept_id].label for concept_id in graph.central_concept_ids if concept_id in nodes_by_id]
    section_counts = Counter(slide.section for slide in main_story)
    return BlueprintPreview(
        deck_title=graph.deck_title,
        lecture_family=graph.lecture_family,
        lecture_family_evidence=graph.lecture_family_evidence,
        rejected_families=graph.rejected_families,
        concept_graph_summary=[
            f"Concept graph contains {len(graph.nodes)} nodes and {len(graph.edges)} edges.",
            f"Central concepts: {', '.join(central_labels[:6]) or 'none'}.",
        ],
        teaching_plan_summary=[
            f"Teaching plan uses {len(main_story)} main-story slides and {len(appendix)} appendix support slides.",
            "Main-story intent counts: "
            + ", ".join(
                f"{intent}={count}"
                for intent, count in sorted(
                    (
                        item
                        for item in teaching_plan.repetition_stats.items()
                        if not item[0].startswith("max-") and not item[0].startswith("section-")
                    )
                )
            ),
            "Section pacing: "
            + ", ".join(
                f"{section}={count}"
                for section, count in sorted(section_counts.items())
            ),
        ],
        ordered_main_story_slide_intents=[slide.intent.value for slide in main_story],
        appendix_plan=teaching_plan.appendix_plan,
        central_concepts=central_labels,
        flagged_drift_risks=teaching_plan.flagged_drift_risks,
        repetition_stats=teaching_plan.repetition_stats,
    )


def _bullet_lines(slide: TeachingPlanSlide, nodes_by_id: dict[str, ConceptNode]) -> list[str]:
    bullets: list[str] = [slide.main_message]
    for concept_id in slide.concept_ids[:2]:
        node = nodes_by_id.get(concept_id)
        if node is not None:
            bullets.append(node.short_teaching_definition)
    if slide.evidence:
        ref = slide.evidence[0]
        location = f"{ref.label} p.{ref.page}" if ref.page is not None else ref.label
        bullets.append(f"Evidence anchor: {location}")
    else:
        bullets.append(slide.pedagogical_goal)
    return _dedupe(_shorten_line(item, 110) for item in bullets)[:3]


def _resolved_visual(slide: TeachingPlanSlide) -> VisualType:
    return slide.visual_hint or INTENT_VISUAL_DEFAULTS[slide.intent]


def _spec_role(index: int, slide: TeachingPlanSlide) -> SlideRole:
    if slide.deck_mode == DeckMode.APPENDIX:
        return SlideRole.APPENDIX_EVIDENCE
    if index == 1 and slide.intent == SlideIntent.ORIENTATION:
        return SlideRole.TITLE
    if slide.intent == SlideIntent.ORIENTATION:
        return SlideRole.EXECUTIVE_SUMMARY
    if slide.intent == SlideIntent.SUMMARY_INTEGRATION:
        return SlideRole.RECOMMENDATION
    visual = _resolved_visual(slide)
    if visual in {VisualType.COMPARISON, VisualType.TABLE}:
        return SlideRole.COMPARISON
    if visual in {VisualType.PROCESS, VisualType.TIMELINE, VisualType.DECISION_PATH}:
        return SlideRole.PROCESS
    if visual in {VisualType.DOCUMENT_CROP, VisualType.CHART, VisualType.PHOTO}:
        if slide.intent in {SlideIntent.WORKED_EXAMPLE, SlideIntent.APPLICATION_VIGNETTE}:
            return SlideRole.EVIDENCE
        return SlideRole.ANALYSIS
    return SlideRole.ANALYSIS


def _specs_from_teaching_plan(graph: ConceptGraph, teaching_plan: TeachingPlan) -> list[dict[str, object]]:
    nodes_by_id = {node.concept_id: node for node in graph.nodes}
    specs: list[dict[str, object]] = []
    for index, slide in enumerate(teaching_plan.slides, start=1):
        specs.append(
            {
                "section": slide.section,
                "deck_mode": slide.deck_mode,
                "role": _spec_role(index, slide),
                "slide_intent": slide.intent,
                "title": slide.title,
                "takeaway": slide.main_message,
                "message": slide.main_message,
                "pedagogical_goal": slide.pedagogical_goal,
                "concept_ids": slide.concept_ids,
                  "visual": _resolved_visual(slide),
                "core_content": _bullet_lines(slide, nodes_by_id),
                "required_assets": [nodes_by_id[concept_id].label for concept_id in slide.concept_ids[:3] if concept_id in nodes_by_id] or [slide.title],
                "source_material_refs": slide.evidence,
                "content_tier": ContentTier.APPENDIX_ONLY if slide.deck_mode == DeckMode.APPENDIX else (
                    ContentTier.SUPPORTING_EXAMPLE if slide.intent in {SlideIntent.WORKED_EXAMPLE, SlideIntent.APPLICATION_VIGNETTE} else ContentTier.LECTURE_CORE
                ),
                "presenter_notes": slide.pedagogical_goal,
            }
        )
    return specs


def plan_concept_driven_lecture(
    *,
    deck_title: str,
    materials: list[ProjectMaterial],
    main_story_target: int,
    appendix_target: int,
) -> LecturePlanningArtifacts:
    chunks = _collect_source_chunks(materials)
    if not chunks:
        empty_graph = ConceptGraph(
            deck_title=deck_title,
            lecture_family=LectureFamily.APPLICATION_COMPARISON,
            lecture_family_evidence=["No document text was available for concept extraction."],
            rejected_families=[],
            nodes=[],
            edges=[],
            central_concept_ids=[],
        )
        empty_plan = _teaching_plan(empty_graph, 0, 0)
        empty_preview = _preview(empty_graph, empty_plan)
        return LecturePlanningArtifacts(
            concept_graph=empty_graph,
            teaching_plan=empty_plan,
            blueprint_preview=empty_preview,
            specs=[],
            clustering_decisions=["Attempted concept-driven lecture planning but no source text was available."],
        )
    lecture_family, lecture_family_evidence, rejected_families = _decide_lecture_family(chunks)
    concept_graph = _build_concept_graph(
        deck_title=deck_title,
        chunks=chunks,
        lecture_family=lecture_family,
        lecture_family_evidence=lecture_family_evidence,
        rejected_families=rejected_families,
    )
    teaching_plan = _teaching_plan(concept_graph, main_story_target, appendix_target)
    blueprint_preview = _preview(concept_graph, teaching_plan)
    return LecturePlanningArtifacts(
        concept_graph=concept_graph,
        teaching_plan=teaching_plan,
        blueprint_preview=blueprint_preview,
        specs=_specs_from_teaching_plan(concept_graph, teaching_plan),
        clustering_decisions=[
            f"Selected lecture family `{concept_graph.lecture_family.value}` from source evidence before blueprint construction.",
            f"Built a concept graph with {len(concept_graph.nodes)} nodes and {len(concept_graph.edges)} edges instead of mirroring document headings.",
            "Synthesized a teaching plan from concept dependencies and pedagogical moves before emitting slide specs.",
            "Derived appendix support after the main teaching plan so appendix slides do not become a cleanup dump.",
        ],
    )
    return seeds
