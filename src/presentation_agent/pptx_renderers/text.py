from __future__ import annotations

from typing import Any

from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from ..slide_scene import TextBox, TextRun
from .common import (
    append_warning,
    stable_scene_shape_name,
)
from .style import SceneStyleContext, resolve_fill_style, resolve_text_style


def render_text_box(
    slide: Any,
    text_box: TextBox,
    *,
    style_context: SceneStyleContext,
    used_shape_names: set[str],
    slide_number: int,
    slide_id: str,
    warnings: list[Any],
) -> str:
    shape = slide.shapes.add_textbox(
        Inches(text_box.bounds.x),
        Inches(text_box.bounds.y),
        Inches(text_box.bounds.width),
        Inches(text_box.bounds.height),
    )
    shape.name = stable_scene_shape_name("text_box", text_box.object_id, used_shape_names)
    if text_box.fill is not None:
        shape.fill.solid()
        shape.fill.fore_color.rgb = resolve_fill_style(
            style_context,
            text_box.fill,
            "#FFFFFF",
            slide_number=slide_number,
            slide_id=slide_id,
            object_id=text_box.object_id,
        ).color_rgb
    else:
        shape.fill.background()
    shape.line.fill.background()

    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = text_box.fit.mode != "none"
    frame.vertical_anchor = MSO_ANCHOR.TOP

    if text_box.fit.mode not in {"none", "wrap", "fail"}:
        append_warning(
            warnings,
            code="fit_policy_not_enforced",
            severity="warning",
            slide_number=slide_number,
            slide_id=slide_id,
            object_id=text_box.object_id,
            message=f"TextBox fit policy {text_box.fit.mode!r} is not enforced exactly in the scene compile path.",
            details={"fit_mode": text_box.fit.mode},
        )

    if text_box.bullet_list:
        for index, item in enumerate(text_box.bullet_list):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.alignment = PP_ALIGN.LEFT
            paragraph.level = item.level
            _apply_bullet_style(
                paragraph,
                item.bullet_style,
                slide_number=slide_number,
                slide_id=slide_id,
                object_id=text_box.object_id,
                warnings=warnings,
            )
            for fragment_index, fragment in enumerate(item.runs):
                run = paragraph.add_run() if fragment_index > 0 else paragraph.runs[0] if paragraph.runs else paragraph.add_run()
                _apply_run_style(run, fragment, text_box.role, style_context, slide_number, slide_id, text_box.object_id)
        return shape.name

    paragraph_groups = _paragraph_groups(text_box.runs)
    for index, group in enumerate(paragraph_groups):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = PP_ALIGN.LEFT
        _clear_bullet_style(paragraph)
        for fragment_index, (source_run, text) in enumerate(group):
            run = paragraph.add_run() if fragment_index > 0 else paragraph.runs[0] if paragraph.runs else paragraph.add_run()
            _apply_run_style(
                run,
                source_run.model_copy(update={"text": text}),
                text_box.role,
                style_context,
                slide_number,
                slide_id,
                text_box.object_id,
            )
    return shape.name


def _paragraph_groups(runs: list[TextRun]) -> list[list[tuple[TextRun, str]]]:
    paragraphs: list[list[tuple[TextRun, str]]] = [[]]
    for source_run in runs:
        parts = source_run.text.split("\n")
        for part_index, part in enumerate(parts):
            paragraphs[-1].append((source_run, part))
            if part_index < len(parts) - 1:
                paragraphs.append([])
    non_empty = [group for group in paragraphs if group]
    return non_empty or [[(TextRun(text=""), "")]]


def _apply_run_style(
    run: Any,
    source_run: TextRun,
    role: str,
    style_context: SceneStyleContext,
    slide_number: int,
    slide_id: str,
    object_id: str | None,
) -> None:
    text_style = resolve_text_style(
        style_context,
        source_run,
        role,
        slide_number=slide_number,
        slide_id=slide_id,
        object_id=object_id,
    )
    run.text = source_run.text
    run.font.name = text_style.font_name
    run.font.size = Pt(text_style.size_pt)
    run.font.bold = text_style.bold
    run.font.italic = text_style.italic
    run.font.color.rgb = text_style.color_rgb


def _apply_bullet_style(
    paragraph: Any,
    bullet_style: str,
    *,
    slide_number: int,
    slide_id: str,
    object_id: str,
    warnings: list[Any],
) -> None:
    _clear_bullet_style(paragraph)
    if bullet_style == "none":
        _set_no_bullet(paragraph)
        return
    p_pr = paragraph._p.get_or_add_pPr()
    if bullet_style == "bullet":
        bullet = OxmlElement("a:buChar")
        bullet.set("char", "\u2022")
        p_pr.insert(0, bullet)
        return
    if bullet_style == "number":
        bullet = OxmlElement("a:buAutoNum")
        bullet.set("type", "arabicPeriod")
        bullet.set("startAt", "1")
        p_pr.insert(0, bullet)
        return
    append_warning(
        warnings,
        code="unsupported_bullet_style",
        severity="warning",
        slide_number=slide_number,
        slide_id=slide_id,
        object_id=object_id,
        message=f"Unsupported bullet style {bullet_style!r}; rendered paragraph without an explicit bullet style.",
        details={"bullet_style": bullet_style},
    )
    _set_no_bullet(paragraph)


def _clear_bullet_style(paragraph: Any) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    for child in list(p_pr):
        if child.tag in {qn("a:buNone"), qn("a:buChar"), qn("a:buAutoNum")}:
            p_pr.remove(child)


def _set_no_bullet(paragraph: Any) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    bullet = OxmlElement("a:buNone")
    p_pr.insert(0, bullet)
